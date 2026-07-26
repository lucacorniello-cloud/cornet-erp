from __future__ import annotations

from contextlib import asynccontextmanager
from collections import defaultdict
from copy import copy
import csv
from datetime import date, datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import re
import uuid
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from jose import jwt
from openpyxl import load_workbook
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from passlib.context import CryptContext
from pydantic import BaseModel, Field
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import Boolean, Date, DateTime, ForeignKey, Integer, String, Text, UniqueConstraint, create_engine, delete, func, or_, select
from sqlalchemy.dialects.postgresql import JSONB, UUID
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker

DATABASE_URL = os.getenv(
    "DATABASE_URL",
    "postgresql+psycopg://cornet:cornet_dev_password@db:5432/cornet_erp",
)
SECRET_KEY = os.getenv("SECRET_KEY", "local-development-secret")


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    email: Mapped[str] = mapped_column(String(255), unique=True, index=True)
    full_name: Mapped[str] = mapped_column(String(255))
    hashed_password: Mapped[str] = mapped_column(String(255))
    is_active: Mapped[bool] = mapped_column(Boolean, default=True)
    role: Mapped[str] = mapped_column(String(50), default="admin")


class Customer(Base):
    __tablename__ = "customers"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    segment: Mapped[str] = mapped_column(String(40), default="BUSINESS_SME")
    business_name: Mapped[str] = mapped_column(String(255), index=True)
    first_name: Mapped[str | None] = mapped_column(String(120))
    last_name: Mapped[str | None] = mapped_column(String(120))
    tax_id: Mapped[str | None] = mapped_column(String(32), unique=True)
    fiscal_code: Mapped[str | None] = mapped_column(String(32))
    windtre_customer_code: Mapped[str | None] = mapped_column(String(80), index=True)
    email: Mapped[str | None] = mapped_column(String(255))
    phone: Mapped[str | None] = mapped_column(String(80))
    address: Mapped[str | None] = mapped_column(String(255))
    postal_code: Mapped[str | None] = mapped_column(String(12))
    city: Mapped[str | None] = mapped_column(String(120))
    province: Mapped[str | None] = mapped_column(String(8))
    birth_date: Mapped[date | None] = mapped_column(Date)
    birth_place: Mapped[str | None] = mapped_column(String(120))
    birth_province: Mapped[str | None] = mapped_column(String(8))
    gender: Mapped[str | None] = mapped_column(String(10))
    document_type: Mapped[str | None] = mapped_column(String(80))
    document_number: Mapped[str | None] = mapped_column(String(80))
    document_issue_date: Mapped[date | None] = mapped_column(Date)
    document_expiry_date: Mapped[date | None] = mapped_column(Date)
    document_issuer: Mapped[str | None] = mapped_column(String(160))
    portfolio_status: Mapped[str] = mapped_column(String(40), default="ACTIVE")
    first_seen_month: Mapped[str | None] = mapped_column(String(7))
    last_seen_month: Mapped[str | None] = mapped_column(String(7))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class CustomerAccountCode(Base):
    __tablename__ = "customer_account_codes"
    __table_args__ = (
        UniqueConstraint("operator", "market", "customer_code", name="uq_customer_account_code_market"),
    )
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    customer_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("customers.id", ondelete="CASCADE"), index=True)
    operator: Mapped[str] = mapped_column(String(80), index=True)
    market: Mapped[str] = mapped_column(String(40), index=True)
    customer_code: Mapped[str] = mapped_column(String(80), index=True)
    is_primary: Mapped[bool] = mapped_column(Boolean, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class IncentiveCompetition(Base):
    __tablename__ = "incentive_competitions"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    name: Mapped[str] = mapped_column(String(255), index=True)
    operator: Mapped[str] = mapped_column(String(80), default="WINDTRE", index=True)
    market: Mapped[str] = mapped_column(String(80), default="MIXED")
    dealer_code: Mapped[str | None] = mapped_column(String(80))
    start_date: Mapped[date] = mapped_column(Date, index=True)
    end_date: Mapped[date] = mapped_column(Date, index=True)
    status: Mapped[str] = mapped_column(String(30), default="DRAFT", index=True)
    source_document: Mapped[str | None] = mapped_column(String(500))
    configuration: Mapped[dict[str, Any]] = mapped_column(JSONB, default=dict)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class IncentiveActivation(Base):
    __tablename__ = "incentive_activations"
    __table_args__ = (
        UniqueConstraint("competition_id", "source_key", name="uq_incentive_activation_source"),
    )
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    competition_id: Mapped[uuid.UUID] = mapped_column(
        ForeignKey("incentive_competitions.id", ondelete="CASCADE"), index=True
    )
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"), index=True)
    activation_date: Mapped[date] = mapped_column(Date, index=True)
    source_type: Mapped[str] = mapped_column(String(40), default="MANUAL")
    source_key: Mapped[str] = mapped_column(String(255))
    seller_name: Mapped[str | None] = mapped_column(String(255))
    track: Mapped[str] = mapped_column(String(40), index=True)
    offer: Mapped[str | None] = mapped_column(String(255))
    asset_number: Mapped[str | None] = mapped_column(String(120))
    customer_code: Mapped[str | None] = mapped_column(String(120), index=True)
    contract_code: Mapped[str | None] = mapped_column(String(120), index=True)
    monthly_fee_cents: Mapped[int] = mapped_column(Integer, default=0)
    direct_bonus_cents: Mapped[int] = mapped_column(Integer, default=0)
    attributes: Mapped[dict[str, Any]] = mapped_column(JSONB, default=dict)
    status: Mapped[str] = mapped_column(String(30), default="VALID", index=True)
    notes: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class IncentivePdcImport(Base):
    __tablename__ = "incentive_pdc_imports"
    __table_args__ = (
        UniqueConstraint("competition_id", "file_sha256", name="uq_incentive_pdc_import_file"),
    )
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    competition_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("incentive_competitions.id", ondelete="CASCADE"), index=True)
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"), index=True)
    file_name: Mapped[str] = mapped_column(String(255))
    file_sha256: Mapped[str] = mapped_column(String(64), index=True)
    stored_path: Mapped[str] = mapped_column(String(500))
    document_type: Mapped[str] = mapped_column(String(80), default="WINDTRE_PDC")
    status: Mapped[str] = mapped_column(String(30), default="IMPORTED")
    extracted_data: Mapped[dict[str, Any]] = mapped_column(JSONB, default=dict)
    activation_ids: Mapped[list[Any]] = mapped_column(JSONB, default=list)
    imported_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class PostActivationTask(Base):
    __tablename__ = "post_activation_tasks"
    __table_args__ = (
        UniqueConstraint("pdc_import_id", "item_key", name="uq_post_activation_task_item"),
    )
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    pdc_import_id: Mapped[uuid.UUID] = mapped_column(
        ForeignKey("incentive_pdc_imports.id", ondelete="CASCADE"), index=True
    )
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"), index=True)
    item_key: Mapped[str] = mapped_column(String(255))
    item_type: Mapped[str] = mapped_column(String(30), default="OPTION")
    item_name: Mapped[str] = mapped_column(String(500))
    customer_code: Mapped[str | None] = mapped_column(String(120), index=True)
    contract_code: Mapped[str | None] = mapped_column(String(120), index=True)
    asset_number: Mapped[str | None] = mapped_column(String(120))
    activation_date: Mapped[date] = mapped_column(Date, index=True)
    action_required: Mapped[bool] = mapped_column(Boolean, default=False, index=True)
    status: Mapped[str] = mapped_column(String(30), default="REVIEW", index=True)
    due_date: Mapped[date | None] = mapped_column(Date, index=True)
    completed_date: Mapped[date | None] = mapped_column(Date)
    notes: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True), default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class PostActivationRule(Base):
    __tablename__ = "post_activation_rules"
    __table_args__ = (
        UniqueConstraint("operator", "item_type", "item_key", name="uq_post_activation_rule_item"),
    )
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    operator: Mapped[str] = mapped_column(String(80), default="WINDTRE", index=True)
    item_type: Mapped[str] = mapped_column(String(30), index=True)
    item_key: Mapped[str] = mapped_column(String(255))
    item_name: Mapped[str] = mapped_column(String(500))
    can_deactivate: Mapped[bool] = mapped_column(Boolean, default=False, index=True)
    default_action_required: Mapped[bool] = mapped_column(Boolean, default=False)
    is_active: Mapped[bool] = mapped_column(Boolean, default=True, index=True)
    notes: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True), default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class WindTreImport(Base):
    __tablename__ = "windtre_imports"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    competence_month: Mapped[str] = mapped_column(String(7), index=True)
    file_name: Mapped[str] = mapped_column(String(255))
    file_sha256: Mapped[str] = mapped_column(String(64))
    status: Mapped[str] = mapped_column(String(30), default="COMPLETED")
    row_count: Mapped[int] = mapped_column(Integer, default=0)
    customer_count: Mapped[int] = mapped_column(Integer, default=0)
    new_customers: Mapped[int] = mapped_column(Integer, default=0)
    missing_customers: Mapped[int] = mapped_column(Integer, default=0)
    new_assets: Mapped[int] = mapped_column(Integer, default=0)
    removed_assets: Mapped[int] = mapped_column(Integer, default=0)
    field_changes: Mapped[int] = mapped_column(Integer, default=0)
    campaign_changes: Mapped[int] = mapped_column(Integer, default=0)
    uploaded_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    rows: Mapped[list["WindTreImportRow"]] = relationship(cascade="all, delete-orphan")
    changes: Mapped[list["WindTreChange"]] = relationship(cascade="all, delete-orphan")


class WindTreImportRow(Base):
    __tablename__ = "windtre_import_rows"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    import_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("windtre_imports.id", ondelete="CASCADE"), index=True)
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"))
    row_number: Mapped[int] = mapped_column(Integer)
    customer_key: Mapped[str] = mapped_column(String(255), index=True)
    asset_key: Mapped[str] = mapped_column(String(255), index=True)
    business_name: Mapped[str] = mapped_column(String(255))
    asset_type: Mapped[str | None] = mapped_column(String(80))
    asset_number: Mapped[str | None] = mapped_column(String(120))
    current_plan: Mapped[str | None] = mapped_column(String(255))
    current_status: Mapped[str | None] = mapped_column(String(80))
    monthly_fee: Mapped[str | None] = mapped_column(String(80))
    raw_data: Mapped[dict[str, Any]] = mapped_column(JSONB)
    campaigns: Mapped[dict[str, Any]] = mapped_column(JSONB)


class WindTreChange(Base):
    __tablename__ = "windtre_changes"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    import_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("windtre_imports.id", ondelete="CASCADE"), index=True)
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"))
    change_type: Mapped[str] = mapped_column(String(50), index=True)
    customer_key: Mapped[str] = mapped_column(String(255))
    asset_key: Mapped[str | None] = mapped_column(String(255))
    field_name: Mapped[str | None] = mapped_column(String(255))
    old_value: Mapped[str | None] = mapped_column(Text)
    new_value: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class StoreSettings(Base):
    __tablename__ = "store_settings"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    store_name: Mapped[str] = mapped_column(String(255), default="Cornet Solutions")
    legal_name: Mapped[str | None] = mapped_column(String(255))
    tax_id: Mapped[str | None] = mapped_column(String(32))
    fiscal_code: Mapped[str | None] = mapped_column(String(32))
    dealer_code: Mapped[str | None] = mapped_column(String(80))
    address: Mapped[str | None] = mapped_column(String(255))
    city: Mapped[str | None] = mapped_column(String(120))
    postal_code: Mapped[str | None] = mapped_column(String(12))
    province: Mapped[str | None] = mapped_column(String(8))
    phone: Mapped[str | None] = mapped_column(String(80))
    whatsapp: Mapped[str | None] = mapped_column(String(80))
    email: Mapped[str | None] = mapped_column(String(255))
    website: Mapped[str | None] = mapped_column(String(255))
    logo_path: Mapped[str | None] = mapped_column(String(500))
    partner_logo_path: Mapped[str | None] = mapped_column(String(500))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class OperatorBrand(Base):
    __tablename__ = "operator_brands"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    operator: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    display_name: Mapped[str] = mapped_column(String(120))
    logo_path: Mapped[str | None] = mapped_column(String(500))
    is_active: Mapped[bool] = mapped_column(Boolean, default=True, index=True)
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True), default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class Product(Base):
    __tablename__ = "products"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    sku: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    name: Mapped[str] = mapped_column(String(255), index=True)
    unit_cost_cents: Mapped[int | None] = mapped_column(Integer)
    is_active: Mapped[bool] = mapped_column(Boolean, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class SimOrder(Base):
    __tablename__ = "sim_orders"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    order_number: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    order_date: Mapped[date] = mapped_column(Date, default=date.today)
    supplier: Mapped[str | None] = mapped_column(String(255))
    status: Mapped[str] = mapped_column(String(40), default="INVIATO")
    notes: Mapped[str | None] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )
    lines: Mapped[list["SimOrderLine"]] = relationship(cascade="all, delete-orphan")


class SimOrderLine(Base):
    __tablename__ = "sim_order_lines"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    order_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("sim_orders.id", ondelete="CASCADE"), index=True)
    product_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("products.id", ondelete="RESTRICT"), index=True)
    quantity: Mapped[int] = mapped_column(Integer)
    description: Mapped[str | None] = mapped_column(Text)
    product: Mapped[Product] = relationship()


class SimInventory(Base):
    __tablename__ = "sim_inventory"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    iccid: Mapped[str] = mapped_column(String(20), unique=True, index=True)
    product_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("products.id", ondelete="RESTRICT"), index=True)
    order_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("sim_orders.id", ondelete="SET NULL"), index=True)
    status: Mapped[str] = mapped_column(String(40), default="IN_MAGAZZINO", index=True)
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"), index=True)
    msisdn: Mapped[str | None] = mapped_column(String(30), index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )
    product: Mapped[Product] = relationship()
    order: Mapped[SimOrder | None] = relationship()
    customer: Mapped[Customer | None] = relationship()


class SimInventoryImport(Base):
    __tablename__ = "sim_inventory_imports"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    file_name: Mapped[str] = mapped_column(String(255))
    added_count: Mapped[int] = mapped_column(Integer, default=0)
    updated_count: Mapped[int] = mapped_column(Integer, default=0)
    rejected_count: Mapped[int] = mapped_column(Integer, default=0)
    errors: Mapped[list[Any]] = mapped_column(JSONB, default=list)
    imported_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class TariffPlan(Base):
    __tablename__ = "tariff_plans"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    plan_code: Mapped[str] = mapped_column(String(50), unique=True, index=True)
    name: Mapped[str] = mapped_column(String(255), unique=True, index=True)
    plan_type: Mapped[str] = mapped_column(String(30), index=True)
    ga_list_code: Mapped[str | None] = mapped_column(String(80), index=True)
    cb_list_code: Mapped[str | None] = mapped_column(String(80), index=True)
    valid_from: Mapped[date | None] = mapped_column(Date)
    valid_to: Mapped[date | None] = mapped_column(Date)
    subscribable: Mapped[bool] = mapped_column(Boolean, default=True, index=True)
    monthly_fee_cents: Mapped[int] = mapped_column(Integer, default=0)
    secure_web_cents: Mapped[int] = mapped_column(Integer, default=0)
    activation_cost_cents: Mapped[int] = mapped_column(Integer, default=0)
    sim_cost_cents: Mapped[int] = mapped_column(Integer, default=0)
    national_gb: Mapped[str | None] = mapped_column(String(100))
    national_minutes: Mapped[str | None] = mapped_column(String(100))
    national_sms: Mapped[str | None] = mapped_column(String(100))
    eu_limited: Mapped[bool] = mapped_column(Boolean, default=False)
    eu_gb: Mapped[str | None] = mapped_column(String(100))
    eu_minutes: Mapped[str | None] = mapped_column(String(100))
    eu_sms: Mapped[str | None] = mapped_column(String(100))
    eu_international_calls: Mapped[str | None] = mapped_column(Text)
    roaming_countries: Mapped[str | None] = mapped_column(Text)
    international_included: Mapped[bool] = mapped_column(Boolean, default=False)
    international_countries: Mapped[str | None] = mapped_column(Text)
    custom_discounts: Mapped[list[dict[str, Any]]] = mapped_column(JSONB, default=list)
    technical_pdf_path: Mapped[str | None] = mapped_column(String(500))
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))


class LetterheadTemplate(Base):
    __tablename__ = "letterhead_templates"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    company_name: Mapped[str] = mapped_column(String(255), default="Cornet Solutions")
    company_address: Mapped[str | None] = mapped_column(String(500))
    tax_id: Mapped[str | None] = mapped_column(String(80))
    phone: Mapped[str | None] = mapped_column(String(100))
    email: Mapped[str | None] = mapped_column(String(255))
    pec: Mapped[str | None] = mapped_column(String(255))
    website: Mapped[str | None] = mapped_column(String(255))
    logo_url: Mapped[str | None] = mapped_column(String(500))
    logo_size: Mapped[int] = mapped_column(Integer, default=80)
    logo_horizontal: Mapped[str] = mapped_column(String(20), default="left")
    logo_vertical: Mapped[str] = mapped_column(String(20), default="top")
    recipient_offset_mm: Mapped[int] = mapped_column(Integer, default=0)
    primary_color: Mapped[str] = mapped_column(String(20), default="#4f46e5")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))


class TerminalCatalogItem(Base):
    __tablename__ = "terminal_catalog_items"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    channel: Mapped[str] = mapped_column(String(10), index=True)
    brand: Mapped[str | None] = mapped_column(String(100), index=True)
    model: Mapped[str] = mapped_column(String(255), index=True)
    memory: Mapped[str | None] = mapped_column(String(80))
    gsi_code: Mapped[str] = mapped_column(String(100), index=True)
    product_type: Mapped[str] = mapped_column(String(40), default="SMARTPHONE", index=True)
    offer_name: Mapped[str | None] = mapped_column(String(255), index=True)
    customer_band: Mapped[str | None] = mapped_column(String(40), index=True)
    list_price_cents: Mapped[int] = mapped_column(Integer, default=0)
    standard_installment_cents: Mapped[int] = mapped_column(Integer, default=0)
    kasko_cents: Mapped[int] = mapped_column(Integer, default=0)
    kasko_premium_cents: Mapped[int] = mapped_column(Integer, default=0)
    upfront_cents: Mapped[int] = mapped_column(Integer, default=0)
    monthly_installment_cents: Mapped[int] = mapped_column(Integer, default=0)
    final_installment_cents: Mapped[int] = mapped_column(Integer, default=0)
    discount_percent: Mapped[int] = mapped_column(Integer, default=0)
    promotion_name: Mapped[str | None] = mapped_column(String(255))
    promotion_id: Mapped[str | None] = mapped_column(String(100))
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))


class TerminalCatalogMetadata(Base):
    __tablename__ = "terminal_catalog_metadata"
    channel: Mapped[str] = mapped_column(String(10), primary_key=True)
    file_name: Mapped[str | None] = mapped_column(String(255))
    sheet_name: Mapped[str | None] = mapped_column(String(255))
    row_count: Mapped[int] = mapped_column(Integer, default=0)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class TerminalInventoryItem(Base):
    __tablename__ = "terminal_inventory_items"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    channel: Mapped[str] = mapped_column(String(10), index=True)
    model: Mapped[str] = mapped_column(String(255), index=True)
    gsi_code: Mapped[str] = mapped_column(String(100), index=True)
    pieces: Mapped[int] = mapped_column(Integer, default=0)
    skip_availability: Mapped[bool] = mapped_column(Boolean, default=False)
    notes: Mapped[str | None] = mapped_column(Text)
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )


class TerminalInventoryMetadata(Base):
    __tablename__ = "terminal_inventory_metadata"
    channel: Mapped[str] = mapped_column(String(10), primary_key=True)
    file_name: Mapped[str | None] = mapped_column(String(255))
    sheet_name: Mapped[str | None] = mapped_column(String(255))
    row_count: Mapped[int] = mapped_column(Integer, default=0)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class Quote(Base):
    __tablename__ = "quotes"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    customer_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("customers.id", ondelete="CASCADE"), index=True)
    quote_type: Mapped[str] = mapped_column(String(40), default="CONFIGURATORE", index=True)
    file_name: Mapped[str] = mapped_column(String(255))
    status: Mapped[str] = mapped_column(String(30), default="GENERATO", index=True)
    line_count: Mapped[int] = mapped_column(Integer, default=0)
    current_mrr_cents: Mapped[int] = mapped_column(Integer, default=0)
    proposed_mrr_cents: Mapped[int] = mapped_column(Integer, default=0)
    payload: Mapped[dict[str, Any]] = mapped_column(JSONB, default=dict)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class DdtShipmentSim(Base):
    __tablename__ = "ddt_shipment_sims"
    ddt_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("ddt_shipments.id", ondelete="CASCADE"), primary_key=True)
    sim_inventory_id: Mapped[uuid.UUID] = mapped_column(ForeignKey("sim_inventory.id", ondelete="RESTRICT"), primary_key=True)


class DdtShipment(Base):
    __tablename__ = "ddt_shipments"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    ddt_number: Mapped[str] = mapped_column(String(40), unique=True, index=True)
    document_date: Mapped[date] = mapped_column(Date, default=date.today, index=True)
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"), index=True)
    recipient_name: Mapped[str] = mapped_column(String(255), index=True)
    recipient_address: Mapped[str] = mapped_column(String(500))
    goods_description: Mapped[str] = mapped_column(Text)
    carrier: Mapped[str | None] = mapped_column(String(255), index=True)
    tracking_number: Mapped[str | None] = mapped_column(String(255), index=True)
    status: Mapped[str] = mapped_column(String(40), default="IN_PREPARAZIONE", index=True)
    sender_snapshot: Mapped[dict[str, Any]] = mapped_column(JSONB)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc),
    )
    customer: Mapped[Customer | None] = relationship()
    sims: Mapped[list[SimInventory]] = relationship(secondary="ddt_shipment_sims")


class WindTrePanelRequest(Base):
    __tablename__ = "windtre_panel_requests"
    id: Mapped[uuid.UUID] = mapped_column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    customer_id: Mapped[uuid.UUID | None] = mapped_column(ForeignKey("customers.id", ondelete="SET NULL"), index=True)
    customer_name: Mapped[str] = mapped_column(String(255), index=True)
    template_key: Mapped[str] = mapped_column(String(80), index=True)
    template_title: Mapped[str] = mapped_column(String(255))
    recipient: Mapped[str | None] = mapped_column(String(255))
    subject: Mapped[str] = mapped_column(String(500))
    body: Mapped[str] = mapped_column(Text)
    form_data: Mapped[dict[str, Any]] = mapped_column(JSONB, default=dict)
    sent_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)
    response_date: Mapped[date | None] = mapped_column(Date)
    status: Mapped[str] = mapped_column(String(20), default="ATTESA", index=True)
    operator_uid: Mapped[str] = mapped_column(String(255), default="local-user")
    customer: Mapped[Customer | None] = relationship()


engine = create_engine(DATABASE_URL, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine)
pwd = CryptContext(schemes=["bcrypt"], deprecated="auto")


class LoginRequest(BaseModel):
    email: str
    password: str


class StoreSettingsRequest(BaseModel):
    store_name: str
    legal_name: str | None = None
    tax_id: str | None = None
    fiscal_code: str | None = None
    dealer_code: str | None = None
    address: str | None = None
    city: str | None = None
    postal_code: str | None = None
    province: str | None = None
    phone: str | None = None
    whatsapp: str | None = None
    email: str | None = None
    website: str | None = None


class DdtShipmentRequest(BaseModel):
    document_date: date = date.today()
    customer_id: uuid.UUID | None = None
    recipient_name: str
    recipient_address: str
    goods_description: str
    carrier: str | None = None
    tracking_number: str | None = None
    status: str = "IN_PREPARAZIONE"
    sim_ids: list[uuid.UUID] = []


class DdtStatusRequest(BaseModel):
    status: str
    tracking_number: str | None = None


class TariffPlanRequest(BaseModel):
    name: str
    plan_type: str
    ga_list_code: str | None = None
    cb_list_code: str | None = None
    valid_from: date | None = None
    valid_to: date | None = None
    subscribable: bool = True
    monthly_fee: float = 0
    secure_web: float = 0
    activation_cost: float = 0
    sim_cost: float = 0
    national_gb: str | None = None
    national_minutes: str | None = None
    national_sms: str | None = None
    eu_limited: bool = False
    eu_gb: str | None = None
    eu_minutes: str | None = None
    eu_sms: str | None = None
    eu_international_calls: str | None = None
    roaming_countries: str | None = None
    international_included: bool = False
    international_countries: str | None = None
    custom_discounts: list[dict[str, Any]] = []


class LetterheadTemplateRequest(BaseModel):
    company_name: str
    company_address: str | None = None
    tax_id: str | None = None
    phone: str | None = None
    email: str | None = None
    pec: str | None = None
    website: str | None = None
    logo_url: str | None = None
    logo_size: int = 80
    logo_horizontal: str = "left"
    logo_vertical: str = "top"
    recipient_offset_mm: int = 0
    primary_color: str = "#4f46e5"


class LetterheadGenerateRequest(BaseModel):
    mode: str = "A4"
    customer_id: uuid.UUID | None = None
    recipient_name: str
    recipient_address: str
    body_text: str | None = None


class TerminalCBRequest(BaseModel):
    brand: str
    model: str
    memory: str | None = None
    gsi_code: str
    product_type: str = "SMARTPHONE"
    customer_band: str = "START"
    list_price: float = 0
    upfront: float = 0
    monthly_installment: float = 0
    final_installment: float = 0


class ConfiguratorRequest(BaseModel):
    asset_keys: list[str]


class WindTrePanelRequestCreate(BaseModel):
    customer_id: uuid.UUID | None = None
    customer_name: str
    template_key: str
    recipient: str | None = None
    subject: str
    body: str
    form_data: dict[str, Any] = {}
    operator_uid: str = "local-user"


class WindTrePanelResponseUpdate(BaseModel):
    response_date: date | None = None


class ProductRequest(BaseModel):
    sku: str
    name: str
    unit_cost: float | None = None


class OrderLineRequest(BaseModel):
    product_id: uuid.UUID
    quantity: int
    description: str | None = None


class SimOrderRequest(BaseModel):
    order_number: str | None = None
    order_date: date = date.today()
    supplier: str | None = None
    notes: str | None = None
    status: str = "INVIATO"
    lines: list[OrderLineRequest]


class OrderStatusRequest(BaseModel):
    status: str


class SimInventoryRequest(BaseModel):
    iccid: str
    product_id: uuid.UUID
    order_id: uuid.UUID | None = None
    status: str = "IN_MAGAZZINO"
    customer_id: uuid.UUID | None = None
    msisdn: str | None = None


class SimInventoryUpdate(BaseModel):
    status: str | None = None
    customer_id: uuid.UUID | None = None
    msisdn: str | None = None


class QuickCustomerRequest(BaseModel):
    business_name: str
    segment: str = "MICROBUSINESS"
    first_name: str | None = None
    last_name: str | None = None
    tax_id: str | None = None
    fiscal_code: str | None = None
    address: str | None = None
    postal_code: str | None = None
    city: str | None = None
    province: str | None = None
    email: str | None = None
    phone: str | None = None
    birth_date: date | None = None
    birth_place: str | None = None
    birth_province: str | None = None
    gender: str | None = None
    document_type: str | None = None
    document_number: str | None = None
    document_issue_date: date | None = None
    document_expiry_date: date | None = None
    document_issuer: str | None = None


class CustomerAccountCodeRequest(BaseModel):
    operator: str = "WINDTRE"
    market: str
    customer_code: str
    is_primary: bool = True


class IncentiveCompetitionRequest(BaseModel):
    name: str
    operator: str = "WINDTRE"
    market: str = "MIXED"
    dealer_code: str | None = None
    start_date: date
    end_date: date
    status: str = "DRAFT"
    source_document: str | None = None
    configuration: dict[str, Any]


class IncentiveActivationRequest(BaseModel):
    activation_date: date
    customer_id: uuid.UUID | None = None
    seller_name: str | None = None
    track: str
    offer: str | None = None
    asset_number: str | None = None
    customer_code: str | None = None
    contract_code: str | None = None
    monthly_fee: float = 0
    direct_bonus: float = 0
    attributes: dict[str, Any] = Field(default_factory=dict)
    status: str = "VALID"
    notes: str | None = None


class PostActivationTaskUpdate(BaseModel):
    action_required: bool | None = None
    mark_completed: bool = False
    completed_date: date | None = None
    notes: str | None = None


class PostActivationRuleRequest(BaseModel):
    operator: str = "WINDTRE"
    item_type: str
    item_name: str
    can_deactivate: bool = False
    default_action_required: bool = False
    is_active: bool = True
    notes: str | None = None


class PostActivationRuleUpdate(BaseModel):
    can_deactivate: bool | None = None
    default_action_required: bool | None = None
    is_active: bool | None = None
    notes: str | None = None


UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "/app/uploads"))
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def seed():
    with SessionLocal() as db:
        if not db.scalar(select(User).where(User.email == "admin@cornet.local")):
            db.add(
                User(
                    email="admin@cornet.local",
                    full_name="Luca Corniello",
                    hashed_password=pwd.hash("Cornet123!"),
                    role="Amministratore",
                )
            )
            db.commit()


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def serialize_account_code(item: CustomerAccountCode) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "customer_id": str(item.customer_id),
        "operator": item.operator,
        "market": item.market,
        "customer_code": item.customer_code,
        "is_primary": item.is_primary,
    }


def normalize_header(value: Any) -> str:
    text = clean(value).upper()
    text = re.sub(r"[^A-Z0-9]+", "_", text)
    return text.strip("_")


ALIASES = {
    "business_name": ["RAGIONE_SOCIALE", "DENOMINAZIONE", "NOME_CLIENTE", "CLIENTE"],
    "customer_code": ["CODICE_CLIENTE", "COD_CLIENTE", "COD_CLIE", "CUSTOMER_CODE"],
    "tax_id": ["PARTITA_IVA", "P_IVA", "PIVA"],
    "fiscal_code": ["CODICE_FISCALE", "COD_FISCALE", "CF"],
    "email": ["EMAIL", "E_MAIL", "MAIL"],
    "phone": ["TELEFONO", "RECAPITO", "CELLULARE"],
    "asset_number": ["MSISDN", "NUMERO_LINEA", "NUMERO_TELEFONO", "NUMERO", "ASSET"],
    "asset_type": ["TIPO_ASSET", "TIPO_SERVIZIO", "CATEGORIA_SERVIZIO", "CATEGORIA"],
    "plan": ["PIANO_TARIFFARIO_ATTUALE", "PIANO_TARIFFARIO", "OFFERTA", "PIANO"],
    "status": ["STATO_MSISDN_MESE_0", "STATO_LINEA", "STATO_ASSET", "STATO"],
    "monthly_fee": ["CANONE_ATTUALIZZATO", "CANONE_SIM", "CANONE_LINEA", "CANONE_ACCESSO", "CANONE"],
    "iccid": ["ICCID", "SERIALE_SIM"],
    "activation_date": [
        "DATA_ATTIVAZIONE",
        "DATA_ATTIVAZIONE_LINEA",
        "DATA_ATTIVAZIONE_MSISDN",
        "DATA_ATTIVAZIONE_ASSET",
        "DATA_INIZIO",
        "DATA_INIZIO_VALIDITA",
        "DATA_DECORRENZA",
    ],
}
SIGNIFICANT_FIELDS = {
    "DATA_ATTIVAZIONE",
    "DATA_ATTIVAZIONE_LINEA",
    "PIANO_TARIFFARIO_ATTUALE",
    "OPZIONE_ATTIVA",
    "TIPO_PACCHETTO_DATI",
    "DESCRIZIONE_TERMINALE",
    "NUM_LICENZE_MKP",
    "DES_LICENZA_MKP",
    "DES_PRODOTTO_MKP",
    "CANONE_SIM",
    "CANONE_LINEA",
    "CANONE_ACCESSO",
    "CANONE_ATTUALIZZATO",
    "STATO_MSISDN_MESE_0",
    "FLAG_OPZIONE_CONVERGENZA",
    "OPZIONE_ITZ",
    "FLAG_OFFICE_SMART",
    "COPERTURA_ACCESSO",
    "CLUSTER_VALORE_CLIENTE",
    "I_CREDIT_SCORE",
}
ASSET_DETAIL_FIELDS = {
    "DATA_ATTIVAZIONE": "Data attivazione",
    "DATA_ATTIVAZIONE_LINEA": "Data attivazione linea",
    "ICCID": "ICCID",
    "OPZIONE_ATTIVA": "Opzione attiva",
    "TIPO_PACCHETTO_DATI": "Pacchetto dati",
    "DESCRIZIONE_TERMINALE": "Terminale",
    "NUM_LICENZE_MKP": "Numero licenze",
    "DES_LICENZA_MKP": "Licenza Marketplace",
    "DES_PRODOTTO_MKP": "Prodotto Marketplace",
    "FLAG_OPZIONE_CONVERGENZA": "Convergenza",
    "OPZIONE_ITZ": "Opzione internazionale",
    "FLAG_OFFICE_SMART": "Office Smart",
    "COPERTURA_ACCESSO": "Copertura accesso",
    "CLUSTER_VALORE_CLIENTE": "Cluster cliente",
    "I_CREDIT_SCORE": "Credit score",
}


def pick(row: dict[str, str], field: str) -> str:
    for name in ALIASES[field]:
        if row.get(name):
            return row[name]
    return ""


def parse_monthly_fee(value: Any) -> float:
    normalized = clean(value).replace("€", "").replace(" ", "")
    if not normalized:
        return 0.0
    if "," in normalized and "." in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    elif "," in normalized:
        normalized = normalized.replace(",", ".")
    normalized = re.sub(r"[^0-9.\-]", "", normalized)
    try:
        return round(float(normalized), 2)
    except (TypeError, ValueError):
        return 0.0


def readable_field_name(key: str) -> str:
    return key.replace("_", " ").strip().title()


def find_activation_date(raw_data: dict[str, Any] | None) -> str:
    raw_data = raw_data or {}
    direct = pick(raw_data, "activation_date")
    if direct:
        return direct
    for key, value in raw_data.items():
        tokens = set(key.upper().split("_"))
        is_date = "DATA" in tokens or key.upper().startswith("DT_")
        is_activation = bool(tokens & {"ATTIVAZIONE", "ATTIV", "DECORRENZA"}) or (
            "INIZIO" in tokens and bool(tokens & {"VALIDITA", "SERVIZIO", "LINEA", "CONTRATTO"})
        )
        if is_date and is_activation and clean(value):
            return clean(value)
    return ""


def asset_details(raw_data: dict[str, Any] | None) -> list[dict[str, str]]:
    raw_data = raw_data or {}
    details = [
        {"key": key, "label": label, "value": clean(raw_data.get(key))}
        for key, label in ASSET_DETAIL_FIELDS.items()
        if clean(raw_data.get(key))
    ]
    known_keys = {detail["key"] for detail in details}
    for key, value in raw_data.items():
        upper = key.upper()
        if key in known_keys or not clean(value):
            continue
        if ("DATA" in upper or upper.startswith("DT_")) and any(
            token in upper for token in ("ATTIV", "DECORRENZA", "INIZIO_VALIDITA")
        ):
            details.append({"key": key, "label": readable_field_name(key), "value": clean(value)})
    return details


def classify_asset(row: WindTreImportRow) -> str:
    raw = row.raw_data or {}
    sim_type = normalize_header(raw.get("TIPOLOGIA_SIM"))
    sim_class = normalize_header(raw.get("CLASSE_SIM"))

    if sim_type == "FONIA_MOBILE":
        if sim_class == "FONIA":
            return "MOBILE_VOICE"
        if sim_class == "DATI":
            return "MOBILE_DATA"
        if sim_class == "M2M":
            return "MOBILE_M2M"
        return "MOBILE_OTHER"
    if sim_type == "MARKETPLACE":
        return "ICT"
    if sim_type in {"VOIP", "ADSL", "FONIA_FISSA", "DATI"}:
        return "FIXED_DATA"

    asset_text = normalize_header(
        " ".join(
            clean(value)
            for value in (
                row.asset_type,
                row.current_plan,
                raw.get("TIPO_ASSET"),
                raw.get("TIPO_SERVIZIO"),
                raw.get("CATEGORIA_SERVIZIO"),
                raw.get("CATEGORIA"),
                raw.get("TIPO_ACCESSO"),
                raw.get("DES_TIPO_ACCESSO"),
                raw.get("DESCRIZIONE_ACCESSO"),
                raw.get("DES_PRODOTTO"),
                raw.get("PRODOTTO"),
            )
            if clean(value)
        )
    )
    fixed_tokens = (
        "FISSO",
        "DATI",
        "FIBRA",
        "FTTH",
        "FTTC",
        "FWA",
        "ADSL",
        "ACCESSO",
        "OFFICE_PLUS",
        "SUPER_OFFICE",
    )
    # Alcune estrazioni WINDTRE valorizzano MSISDN anche sulle linee fisse:
    # gli indicatori specifici Fisso/Dati devono quindi avere la precedenza.
    if (
        clean(raw.get("CANONE_LINEA"))
        or clean(raw.get("CANONE_ACCESSO"))
        or any(token in asset_text for token in fixed_tokens)
    ):
        return "FIXED_DATA"
    if (
        clean(raw.get("CANONE_SIM"))
        or any(token in asset_text for token in ("MOBILE", "SIM", "MSISDN"))
        or clean(raw.get("MSISDN"))
    ):
        if sim_class == "FONIA":
            return "MOBILE_VOICE"
        if sim_class == "DATI":
            return "MOBILE_DATA"
        if sim_class == "M2M":
            return "MOBILE_M2M"
        return "MOBILE_OTHER"
    return "OTHER"


def is_active_service(row: WindTreImportRow) -> bool:
    status = normalize_header(row.current_status or "")
    return not status or status in {"ATT", "ACTIVE", "ATTIVO", "ATTIVA"}


def customer_service_pivot(rows: list[WindTreImportRow]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, float | int]] = defaultdict(lambda: {"count": 0, "mrr": 0.0})
    for row in rows:
        if not is_active_service(row):
            continue
        plan = clean(row.current_plan) or "Piano non indicato"
        grouped[plan]["count"] = int(grouped[plan]["count"]) + 1
        grouped[plan]["mrr"] = round(float(grouped[plan]["mrr"]) + parse_monthly_fee(row.monthly_fee), 2)
    return [
        {"plan": plan, "count": values["count"], "mrr": values["mrr"]}
        for plan, values in sorted(grouped.items(), key=lambda item: item[0].casefold())
    ]


def safe_export_name(value: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return normalized[:80] or "cliente"


def italian_currency(value: float) -> str:
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " EUR"


def latest_customer_snapshot(db, customer_id: uuid.UUID):
    latest_import = db.scalar(
        select(WindTreImport)
        .order_by(WindTreImport.competence_month.desc(), WindTreImport.uploaded_at.desc())
        .limit(1)
    )
    rows = (
        db.scalars(
            select(WindTreImportRow)
            .where(
                WindTreImportRow.customer_id == customer_id,
                WindTreImportRow.import_id == latest_import.id,
            )
            .order_by(WindTreImportRow.row_number)
        ).all()
        if latest_import
        else []
    )
    return latest_import, rows


def find_header(sheet) -> tuple[int, list[str]]:
    best: tuple[int, int, list[str]] | None = None
    known = {item for values in ALIASES.values() for item in values}
    for number, values in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True), 1):
        headers = [normalize_header(v) for v in values]
        score = sum(bool(h) for h in headers) + 5 * sum(h in known or h.startswith(("V_", "C_", "I_")) for h in headers)
        if best is None or score > best[0]:
            best = (score, number, headers)
    if not best or best[0] < 5:
        raise ValueError("Intestazioni Excel non riconosciute")
    return best[1], best[2]


def parse_workbook(contents: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    sheet = workbook.active
    header_row, headers = find_header(sheet)
    parsed: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        raw = {headers[index]: clean(value) for index, value in enumerate(values) if index < len(headers) and headers[index] and clean(value)}
        if not raw:
            continue
        business_name = pick(raw, "business_name")
        customer_code = pick(raw, "customer_code")
        tax_id = pick(raw, "tax_id")
        fiscal_code = pick(raw, "fiscal_code")
        asset_number = pick(raw, "asset_number")
        iccid = pick(raw, "iccid")
        if not any((business_name, customer_code, tax_id, fiscal_code, asset_number)):
            continue
        customer_key = customer_code or tax_id or fiscal_code or normalize_header(business_name)
        asset_key = asset_number or iccid or f"{customer_key}:ROW:{row_number}"
        campaigns = {key: value for key, value in raw.items() if key.startswith(("V_", "C_", "I_")) and value}
        parsed.append(
            {
                "row_number": row_number,
                "customer_key": customer_key,
                "asset_key": asset_key,
                "business_name": business_name or f"Cliente {customer_key}",
                "customer_code": customer_code,
                "tax_id": tax_id,
                "fiscal_code": fiscal_code,
                "email": pick(raw, "email"),
                "phone": pick(raw, "phone"),
                "asset_type": pick(raw, "asset_type"),
                "asset_number": asset_number,
                "plan": pick(raw, "plan"),
                "status": pick(raw, "status"),
                "monthly_fee": pick(raw, "monthly_fee"),
                "activation_date": find_activation_date(raw),
                "raw": raw,
                "campaigns": campaigns,
            }
        )
    return parsed


def validate_excel_upload(filename: str | None, contents: bytes) -> None:
    if not filename or not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Carica un file Excel .xlsx o .xlsm")
    if len(contents) > 40 * 1024 * 1024:
        raise HTTPException(413, "Il file supera il limite di 40 MB")


def build_preview(filename: str, contents: bytes, parsed: list[dict[str, Any]]) -> dict[str, Any]:
    customer_keys = {row["customer_key"] for row in parsed}
    asset_keys = [row["asset_key"] for row in parsed]
    campaign_names = sorted({name for row in parsed for name in row["campaigns"]})
    duplicate_assets = len(asset_keys) - len(set(asset_keys))
    rows_without_tax_id = sum(not (row["tax_id"] or row["fiscal_code"]) for row in parsed)
    rows_without_customer_code = sum(not row["customer_code"] for row in parsed)
    return {
        "file_name": filename,
        "file_sha256": sha256(contents).hexdigest(),
        "row_count": len(parsed),
        "customer_count": len(customer_keys),
        "asset_count": len(set(asset_keys)),
        "campaign_count": len(campaign_names),
        "campaign_names": campaign_names,
        "quality": {
            "duplicate_asset_rows": duplicate_assets,
            "rows_without_tax_id": rows_without_tax_id,
            "rows_without_customer_code": rows_without_customer_code,
        },
        "sample_rows": [
            {
                "row_number": row["row_number"],
                "business_name": row["business_name"],
                "customer_code": row["customer_code"],
                "tax_id": row["tax_id"],
                "asset_type": row["asset_type"],
                "asset_number": row["asset_number"],
                "plan": row["plan"],
                "status": row["status"],
                "monthly_fee": row["monthly_fee"],
                "campaigns": row["campaigns"],
            }
            for row in parsed[:20]
        ],
    }


def text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def serialize_import(item: WindTreImport) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "competence_month": item.competence_month,
        "file_name": item.file_name,
        "status": item.status,
        "row_count": item.row_count,
        "customer_count": item.customer_count,
        "new_customers": item.new_customers,
        "missing_customers": item.missing_customers,
        "new_assets": item.new_assets,
        "removed_assets": item.removed_assets,
        "field_changes": item.field_changes,
        "campaign_changes": item.campaign_changes,
        "uploaded_at": item.uploaded_at.isoformat(),
    }


PRODUCT_COLUMN_ALIASES = {
    "sku": {"CODICE", "CODICE_ARTICOLO", "SKU", "ARTICOLO"},
    "name": {"NOME", "NOME_PRODOTTO", "DESCRIZIONE", "DESCRIZIONE_COMMERCIALE"},
    "unit_cost": {"COSTO", "COSTO_UNITARIO", "PREZZO_ACQUISTO"},
}
SIM_COLUMN_ALIASES = {
    "iccid": {"ICCID", "SERIALE", "SERIALE_SIM", "SERIALE_ICCID"},
    "sku": {"CODICE", "CODICE_ARTICOLO", "SKU", "ARTICOLO"},
    "status": {"STATO", "STATO_SIM"},
    "order_number": {"ORDINE", "ID_ORDINE", "NUMERO_ORDINE"},
    "msisdn": {"MSISDN", "NUMERO", "NUMERO_TELEFONICO", "NUMERO_ASSEGNATO_MSISDN"},
    "customer": {"CLIENTE", "NOME_CLIENTE", "ASSEGNATA_A_CLIENTE"},
    "order_date": {"DATA_ORDINE", "DATA"},
}
SIM_STATUSES = {"IN_MAGAZZINO", "ASSEGNATA", "ATTIVATA", "DISABILITATA", "SOSPESA"}
ORDER_STATUSES = {"INVIATO", "IN_ATTESA", "IN_LAVORAZIONE", "RICEVUTO", "EVASO"}


def parse_uploaded_table(filename: str, contents: bytes) -> list[dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [normalize_header(value) for value in next(values, [])]
        return [
            {
                headers[index]: clean(value)
                for index, value in enumerate(row)
                if index < len(headers) and headers[index] and clean(value)
            }
            for row in values
            if any(clean(value) for value in row)
        ]
    if suffix == ".csv":
        decoded = None
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                decoded = contents.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise HTTPException(422, "Codifica CSV non riconosciuta")
        sample = decoded[:65536]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(decoded.splitlines(), dialect=dialect)
        return [
            {normalize_header(key): clean(value) for key, value in row.items() if key and clean(value)}
            for row in reader
            if any(clean(value) for value in row.values())
        ]
    raise HTTPException(422, "Carica un file CSV, XLSX o XLSM")


def mapped_value(row: dict[str, str], aliases: set[str]) -> str:
    return next((row[key] for key in aliases if row.get(key)), "")


def parse_compact_date(value: str) -> date:
    normalized = re.sub(r"\D", "", value or "")
    for pattern in ("%Y%m%d", "%d%m%Y"):
        try:
            return datetime.strptime(normalized, pattern).date()
        except ValueError:
            pass
    return date.today()


def cents_from_value(value: Any) -> int | None:
    parsed = parse_monthly_fee(value)
    return round(parsed * 100) if clean(value) else None


def serialize_product(item: Product) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "sku": item.sku,
        "name": item.name,
        "unit_cost": round(item.unit_cost_cents / 100, 2) if item.unit_cost_cents is not None else None,
        "is_active": item.is_active,
    }


def serialize_order(item: SimOrder) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "order_number": item.order_number,
        "order_date": item.order_date.isoformat(),
        "supplier": item.supplier,
        "status": item.status,
        "notes": item.notes,
        "lines": [
            {
                "id": str(line.id),
                "product_id": str(line.product_id),
                "sku": line.product.sku,
                "product_name": line.product.name,
                "quantity": line.quantity,
                "description": line.description,
            }
            for line in item.lines
        ],
        "total_quantity": sum(line.quantity for line in item.lines),
    }


def serialize_inventory(item: SimInventory) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "iccid": item.iccid,
        "product_id": str(item.product_id),
        "sku": item.product.sku,
        "product_name": item.product.name,
        "order_id": str(item.order_id) if item.order_id else None,
        "order_number": item.order.order_number if item.order else None,
        "status": item.status,
        "customer_id": str(item.customer_id) if item.customer_id else None,
        "customer_name": item.customer.business_name if item.customer else None,
        "msisdn": item.msisdn,
        "created_at": item.created_at.isoformat(),
    }


def normalize_iccid(value: Any) -> str:
    return re.sub(r"\D", "", clean(value))


def order_share_text(item: SimOrder) -> str:
    lines = [
        f"ORDINE SIM {item.order_number}",
        f"Data: {item.order_date.strftime('%d/%m/%Y')}",
        f"Fornitore: {item.supplier or 'Non indicato'}",
        "",
    ]
    lines.extend(f"- {line.product.sku} | {line.product.name}: {line.quantity} pz" for line in item.lines)
    if item.notes:
        lines.extend(["", f"Note: {item.notes}"])
    lines.append(f"\nTotale: {sum(line.quantity for line in item.lines)} SIM")
    return "\n".join(lines)


def march_2026_incentive_configuration() -> dict[str, Any]:
    return {
        "version": 1,
        "source": "Lettera WINDTRE Incentivazione Marzo 2026",
        "tracks": {
            "MOBILE": {
                "label": "Attivazioni mobili",
                "access_threshold": 25,
                "thresholds": [
                    {"target": 25, "label": "1a soglia", "mnp_multiplier": 2.25, "no_mnp_multiplier": 0.5, "tied_extra_multiplier": 0.75},
                    {"target": 55, "label": "2a soglia", "mnp_multiplier": 3.25, "no_mnp_multiplier": 1.0, "tied_extra_multiplier": 0.75},
                    {"target": 105, "label": "3a soglia", "mnp_multiplier": 3.75, "no_mnp_multiplier": 1.5, "tied_extra_multiplier": 1.0},
                    {"target": 143, "label": "4a soglia", "mnp_multiplier": 4.25, "no_mnp_multiplier": 2.0, "tied_extra_multiplier": 1.0},
                ],
                "rules": [
                    {"code": "VERY", "label": "Very Mobile (esclusa MNP intrabrand)", "points": 0.5},
                    {"code": "SECURE", "label": "Più Sicuri Mobile/Pro contestuale", "points": 1.25},
                    {"code": "PHONE_INCLUDED", "label": "Tied con Telefono Incluso standard o Rata Smart 1° device", "extra_points": 1.25},
                    {"code": "TIED_PREMIUM", "label": "Tied Start/Unlimited/Unlimited Pro/Pack Reload Exchange", "extra_points": 0.5},
                    {"code": "TIED_PIVA", "label": "Tied Partita IVA", "extra_points": 1.0},
                    {"code": "TRAVEL", "label": "Offerte Travel", "excluded": True},
                ],
            },
            "FIXED": {
                "label": "Fisso e convergenza",
                "access_threshold": 4,
                "thresholds": [
                    {"target": 4, "label": "1a soglia", "convergent_multiplier": 4.75, "standard_multiplier": 3.25},
                    {"target": 8, "label": "2a soglia", "convergent_multiplier": 5.0, "standard_multiplier": 3.5},
                    {"target": 15, "label": "3a soglia", "convergent_multiplier": 5.5, "standard_multiplier": 4.0},
                    {"target": 23, "label": "4a soglia", "convergent_multiplier": 6.25, "standard_multiplier": 4.75},
                    {"target": 36, "label": "Soglia extra", "convergent_multiplier": 7.25, "standard_multiplier": 5.75},
                ],
                "rules": [
                    {"code": "PIVA", "label": "Cliente P.IVA (1a e 2a linea)", "extra_multiplier": 1.0},
                    {"code": "FTTH", "label": "FTTH prima linea", "extra_multiplier": 1.0},
                    {"code": "FWA", "label": "FWA Indoor 2P/Outdoor prima linea", "extra_multiplier": 1.5},
                    {"code": "SECOND_PRO", "label": "Seconda linea Professional", "points": 1.5},
                    {"code": "PRO_BOX", "label": "Professional Box con FRITZ!Box", "points": 4.0},
                ],
            },
            "CUSTOMER_BASE": {
                "label": "Customer Base",
                "access_threshold": 0,
                "thresholds": [],
                "rules": [{"code": "DIRECT", "label": "Gettoni configurabili per evento"}],
            },
            "RELOAD": {
                "label": "Smartphone Reload",
                "access_threshold": 0,
                "thresholds": [
                    {"target": 25, "label": "Attachment 25%", "bonus_multiplier": 2.0},
                    {"target": 40, "label": "Attachment 40%", "bonus_multiplier": 3.0},
                    {"target": 55, "label": "Attachment 55%", "bonus_multiplier": 4.0},
                ],
                "rules": [{"code": "ATTACHMENT", "label": "Incidenza Reload sul totale device"}],
            },
        },
        "validation_rules": [
            "Evento inserito e attivato nel periodo di gara",
            "Linea attiva e non sospesa alla consuntivazione",
            "MNP espletata; escluse POPI e migrazioni non ammesse",
            "Verifica rinnovi, silenza, disconoscimenti e storni successivi",
        ],
    }


def pdf_document_text(contents: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(contents))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise HTTPException(422, f"Impossibile leggere il PDF: {exc}") from exc


def regex_value(text: str, pattern: str, flags: int = re.IGNORECASE) -> str:
    match = re.search(pattern, text, flags)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def parsed_date(value: str) -> date | None:
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value.strip(), pattern).date()
        except (ValueError, AttributeError):
            pass
    return None


def format_date_it(value: str | None) -> str:
    parsed = parsed_date(value or "")
    return parsed.strftime("%d/%m/%Y") if parsed else (value or "")


def pdc_reload_base_bonus(device_price: float) -> float:
    if device_price < 150:
        return 5.0
    if device_price < 300:
        return 8.0
    if device_price < 500:
        return 10.0
    if device_price < 700:
        return 11.0
    return 14.0


def first_business_day_next_month(value: date) -> date:
    candidate = (value.replace(day=28) + timedelta(days=4)).replace(day=1)
    while candidate.weekday() >= 5:
        candidate += timedelta(days=1)
    return candidate


def pdc_post_activation_items(offer: str, options: str) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    seen: set[str] = set()

    def add(item_type: str, name: str):
        normalized_name = re.sub(r"\s+", " ", name).strip(" -:;.")
        key = normalize_header(f"{item_type}_{normalized_name}")[:255]
        if normalized_name and key not in seen:
            seen.add(key)
            items.append({"key": key, "type": item_type, "name": normalized_name})

    add("OFFER", offer)
    known_options = (
        r"GIGA illimitati per le tue SIM",
        r"Giga Illimitati super Fibra",
        r"Pi[uù]'?\s*Sicuri Casa&Ufficio",
        r"Pi[uù]'?\s*Sicuri Mobile(?: Easy Pay)?",
        r"Contributo rateizzato",
        r"Convergenza Special",
        r"Smartphone Reload",
        r"Sconto Rata Telefono Incluso PRE New",
        r"Vendita a rate",
        r"Reload Forever Basic",
        r"Segreteria Telefonica",
    )
    for pattern in known_options:
        match = re.search(pattern, options, re.IGNORECASE)
        if match:
            add("OPTION", match.group(0))
    return items


def parse_windtre_pdc(contents: bytes, file_name: str = "") -> dict[str, Any]:
    text = pdf_document_text(contents)
    upper = text.upper()
    if "PROPOSTA DI CONTRATTO WINDTRE" not in upper and "PROPOSTA DI CONTRATTO" not in upper:
        raise HTTPException(422, "Il documento non è stato riconosciuto come PDC WINDTRE")
    is_fixed_ga = "LINEA TELEFONICA SU CUI ATTIVARE IL SERVIZIO" in upper
    is_mobile_ga = "DATI DELLA SIM WINDTRE" in upper and any(
        marker in upper for marker in ("GO UNLIMITED", "SUPER UNLIMITED", "MIA UNLIMITED", "START 5G")
    )
    first_name = regex_value(text, r"Cognome:\s*[A-ZÀ-Ý' ]+\s+Nome:\s*([A-ZÀ-Ý' ]+?)\s+Sesso:")
    last_name = regex_value(text, r"Cognome:\s*([A-ZÀ-Ý' ]+?)\s+Nome:")
    if is_fixed_ga:
        first_name = regex_value(text, r"(?:^|\n)Nome:\s*([^\n]+)")
        last_name = regex_value(text, r"(?:^|\n)Cognome:\s*([^\n]+)")
    fiscal_code = regex_value(text, r"Codice Fiscale:\s*([A-Z0-9]{16})")
    birth_date_text = regex_value(text, r"Data di nascita:\s*(\d{2}/\d{2}/\d{4})")
    activation_date_text = regex_value(text, r"(?:^|\n)Data:\s*(\d{2}/\d{2}/\d{4})")
    phone = re.sub(r"^39", "", re.sub(r"\D", "", regex_value(
        text, r"(?:N\.\s*Telefono|Numero di telefono):\s*(\+?\d{9,15})"
    )))
    if is_mobile_ga:
        phone = re.sub(r"^39", "", re.sub(r"\D", "", regex_value(
            text, r"Numero di telefono:\s*(\+?\d{9,15})\s+Seriale SIM"
        )))
    if is_fixed_ga:
        phone = re.sub(r"\D", "", regex_value(text, r"Sul Numero di Telefono:\s*(\d{8,15})"))
    contact_phone = re.sub(r"^39", "", re.sub(r"\D", "", regex_value(
        text, r"(?:Intestatario Contratto[\s\S]{0,100}?Numero di telefono|Recapito):\s*(\+?\d{9,15})"
    )))
    iccid = regex_value(text, r"(?:Seriale SIM \(ICCID\)|Numero della SIM):\s*(\d{19,20})")
    customer_code = regex_value(
        text, r"Codice Cliente\s*:?\s*(?:PROPOSTA DI CONTRATTO\s*)?(\d{6,})"
    )
    contract_code = regex_value(text, r"(?:Codice Contratto|Numero Contratto)\s*:?\s*([A-Z0-9.]+)")
    dealer_code = regex_value(text, r"(?:Codice Rivenditore|Codice Dealer)\s*:?\s*([A-Z0-9.]+)")
    plan = regex_value(text, r"Piano Telefonico:\s*(.+?)\s*(?:\n|Opzioni/servizi:)")
    options = regex_value(text, r"Opzioni/servizi:\s*([\s\S]+?)(?:Costo attivazione|Data:)")
    fixed_offer = regex_value(text, r"Offerta\s*\(\*\)\s*\n\s*([^\n]+)")
    fixed_options = regex_value(text, r"Opzioni aggiuntive\s*\n([\s\S]+?)\(\*\)")
    mobile_offer = regex_value(options, r"^\s*([^:\n]+?)\s*:")
    mobile_monthly_fee = parse_monthly_fee(regex_value(
        options, r"^[^:\n]+?:\s*Costo mensile:\s*([\d.,]+)\s*euro"
    ))
    imei = regex_value(text, r"Numero IMEI:\s*(\d{15})")
    device_model = regex_value(text, r"Modello:\s*(.+?)\s*(?:\n|Data di acquisto:)")
    device_price = parse_monthly_fee(regex_value(text, r"Prezzo device:\s*([\d.,]+)\s*euro"))
    upfront = parse_monthly_fee(regex_value(text, r"Anticipo:\s*([\d.,]+)\s*euro"))
    installments = int(regex_value(text, r"Durata Rateizzazione:\s*(\d+)") or 0)
    installment_amount = parse_monthly_fee(regex_value(text, r"\d+\s*rate da\s*([\d.,]+)\s*euro"))
    reload_active = any(marker in upper for marker in (
        "DATI DEL SERVIZIO  SMARTPHONE RELOAD", "DATI DEL SERVIZIO SMARTPHONE RELOAD",
        "DATI DEL SERVIZIO  RELOAD FOREVER", "DATI DEL SERVIZIO RELOAD FOREVER",
    ))
    reload_cost = parse_monthly_fee(regex_value(text, r"Costo servizio:\s*([\d.,]+)\s*€"))
    payment_method = "SDD" if "ADDEBITO DIRETTO SU C/C" in upper else "ALTRO"
    address = regex_value(text, r"Residenza:\s*(.+?)\s+Provincia:")
    if is_fixed_ga:
        street_type = regex_value(text, r"Via, Piazza, etc:\s*([^\n]+?)\s+(?:Recapito:|N\.:|Prov\.:)")
        street_name = regex_value(text, r"Indirizzo:\s*([^\n]+?)\s+(?:Città|Cap:)")
        street_number = regex_value(text, r"N\.:\s*(\d+)")
        address = " ".join(value for value in (street_type, street_name, street_number) if value)
    postal_code = regex_value(text, r"CAP:\s*(\d{5})")
    city = regex_value(text, r"Comune:\s*([A-ZÀ-Ý' ]+?)\s+Nazione:")
    if is_fixed_ga:
        city = regex_value(text, r"Città, Località:\s*([A-ZÀ-Ý' ]+?)\s+Cap:")
    email = regex_value(text, r"Email:\s*([^\s]+@[^\s]+)")
    identity_document_type = regex_value(text, r"Documento d'Identit(?:à|a'):\s*([^;\n]+)")
    document_number = regex_value(text, r"Numero:\s*([A-Z0-9]+)\s+Data Rilascio:")
    issue_date_text = regex_value(text, r"Data Rilascio:\s*(\d{2}/\d{2}/\d{4})")
    activation_date = parsed_date(activation_date_text)
    company_tax_id = regex_value(
        text,
        r"Partita Iva/Cod\.Fiscale Azienda:[ \t]*(\d{11})(?:\s|$)",
    )
    is_piva = bool(company_tax_id)
    proposed_entries = []
    if is_fixed_ga:
        proposed_entries.append({
            "track": "FIXED", "label": "Nuova attivazione Fisso GA",
            "offer": fixed_offer or "Linea fissa WINDTRE", "direct_bonus": 0,
            "monthly_fee": 0, "points": 1,
            "attributes": {
                "new_line": "SU UNA NUOVA LINEA TELEFONICA" in upper,
                "first_line": True, "piva": is_piva,
                "convergent": "GIGA ILLIMITATI" in fixed_options.upper(),
                "ftth": False, "fwa": False, "technology_to_verify": True,
            },
        })
    elif is_mobile_ga:
        proposed_entries.append({
            "track": "MOBILE", "label": "Nuova attivazione Mobile GA",
            "offer": mobile_offer or plan, "direct_bonus": 0,
            "monthly_fee": mobile_monthly_fee, "points": 1,
            "attributes": {
                "mnp": bool(regex_value(text, r"Operatore di provenienza:\s*([^\n]+)")),
                "tied": "EASY PAY" in options.upper(), "piva": is_piva,
                "secure_option": "PIU' SICURI" in options.upper() or "PIÙ SICURI" in options.upper(),
                "convergent": "CONVERGENZA" in options.upper() or "GIGA ILLIMITATI" in options.upper(),
                "phone_included": False, "device_sale": False,
            },
        })
    elif imei:
        proposed_entries.append({
            "track": "CUSTOMER_BASE", "label": "Telefono Incluso - vendita a rate",
            "offer": f"Telefono Incluso · {device_model}", "direct_bonus": 8.0,
            "monthly_fee": 0, "points": 1,
            "attributes": {"device_sale": True, "phone_included": True, "piva": is_piva, "payment_type": "VAR"},
        })
    if reload_active:
        proposed_entries.append({
            "track": "RELOAD", "label": "Reload Forever" if "RELOAD FOREVER" in upper else "Smartphone Reload",
            "offer": f"{'Reload Forever' if 'RELOAD FOREVER' in upper else 'Smartphone Reload'} · {device_model}".rstrip(" ·"),
            "direct_bonus": pdc_reload_base_bonus(device_price) if device_price else 0,
            "monthly_fee": 0, "points": 1,
            "attributes": {"reload_service": True, "device_sale": False, "device_price": device_price, "reload_cost": reload_cost},
            "status": "VALID" if device_price else "TO_VERIFY",
        })
    warnings = []
    if not activation_date:
        warnings.append("Data di attivazione non rilevata")
    if not fiscal_code:
        warnings.append("Codice fiscale non rilevato")
    if not customer_code:
        warnings.append("Codice cliente non rilevato")
    if reload_active and not device_price:
        warnings.append("Reload rilevato, ma il prezzo di listino del terminale non è presente: quota Reload da verificare")
    if is_fixed_ga and not fixed_offer:
        warnings.append("Offerta fissa non rilevata")
    if is_fixed_ga:
        warnings.append("Tecnologia FTTH/FTTC e canone non sono esposti nella PDC: completare dopo l'esito tecnico")
    if not is_fixed_ga and not is_mobile_ga and (not plan or plan.upper() in {"WIND BASIC", "NEW BASIC"}):
        warnings.append("La PDC non espone un'offerta mobile ricorrente remunerabile: non viene creata automaticamente una nuova attivazione Mobile")
    if is_fixed_ga:
        document_type = "WINDTRE_PDC_GA_FIXED"
        classification = {"new_mobile_activation": False, "reason": "Nuova attivazione Fisso GA"}
    elif is_mobile_ga:
        document_type = "WINDTRE_PDC_GA_MOBILE"
        classification = {"new_mobile_activation": True, "reason": "Nuova attivazione Mobile GA"}
    else:
        document_type = "WINDTRE_PDC_DEVICE_RELOAD" if reload_active else "WINDTRE_PDC"
        classification = {
            "new_mobile_activation": False,
            "reason": "PDC con vendita a rate/Reload su linea esistente; classificata Customer Base salvo verifica operatore",
        }
    effective_offer = fixed_offer if is_fixed_ga else (mobile_offer or plan)
    effective_options = fixed_options if is_fixed_ga else options
    return {
        "document_type": document_type,
        "file_name": file_name,
        "customer": {
            "first_name": first_name.title(), "last_name": last_name.title(),
            "business_name": f"{first_name.title()} {last_name.title()}".strip(),
            "fiscal_code": fiscal_code, "segment": "CONSUMER", "birth_date": birth_date_text,
            "address": address.title(), "postal_code": postal_code, "city": city.title(),
            "email": email, "contact_phone": contact_phone,
            "document_type": identity_document_type, "document_number": document_number,
            "document_issue_date": issue_date_text,
        },
        "contract": {
            "operator": "WINDTRE", "market": "CONSUMER", "customer_code": customer_code,
            "contract_code": contract_code, "activation_date": activation_date.isoformat() if activation_date else None,
            "dealer_code": dealer_code, "phone": phone, "iccid": iccid,
            "plan": effective_offer,
            "options": effective_options, "payment_method": payment_method,
        },
        "device": {
            "imei": imei, "model": device_model, "price": device_price, "upfront": upfront,
            "installments": installments, "installment_amount": installment_amount,
            "reload_active": reload_active, "reload_cost": reload_cost,
        },
        "classification": classification,
        "proposed_entries": proposed_entries,
        "post_activation_items": pdc_post_activation_items(effective_offer, effective_options),
        "warnings": warnings,
    }


def post_activation_rule_key(item_type: str, item_name: str) -> str:
    return normalize_header(f"{item_type}_{item_name}")[:255]


def ensure_post_activation_rule(
    db, item_type: str, item_name: str, operator: str = "WINDTRE"
) -> PostActivationRule:
    item_key = post_activation_rule_key(item_type, item_name)
    rule = db.scalar(select(PostActivationRule).where(
        PostActivationRule.operator == operator,
        PostActivationRule.item_type == item_type,
        PostActivationRule.item_key == item_key,
    ))
    if rule:
        if rule.item_name != item_name:
            rule.item_name = item_name
        return rule
    rule = PostActivationRule(
        operator=operator, item_type=item_type, item_key=item_key,
        item_name=item_name, can_deactivate=False,
        default_action_required=False, is_active=True,
    )
    db.add(rule)
    db.flush()
    return rule


def ensure_post_activation_tasks_for_record(db, record: IncentivePdcImport) -> int:
    extracted = record.extracted_data or {}
    contract = extracted.get("contract", {})
    activation_date = parsed_date(contract.get("activation_date", ""))
    if not activation_date:
        return 0
    items = extracted.get("post_activation_items") or pdc_post_activation_items(
        contract.get("plan", ""), contract.get("options", "")
    )
    existing = set(db.scalars(
        select(PostActivationTask.item_key).where(PostActivationTask.pdc_import_id == record.id)
    ).all())
    added = 0
    for item in items:
        rule = ensure_post_activation_rule(db, item["type"], item["name"])
        if item["key"] in existing:
            continue
        action_required = bool(rule.is_active and rule.can_deactivate and rule.default_action_required)
        db.add(PostActivationTask(
            pdc_import_id=record.id, customer_id=record.customer_id,
            item_key=item["key"], item_type=item["type"], item_name=item["name"],
            customer_code=contract.get("customer_code") or None,
            contract_code=contract.get("contract_code") or None,
            asset_number=contract.get("phone") or ("NUOVA LINEA" if extracted.get("document_type") == "WINDTRE_PDC_GA_FIXED" else None),
            activation_date=activation_date, action_required=action_required,
            status="PENDING" if action_required else "REVIEW",
            due_date=first_business_day_next_month(activation_date) if action_required else None,
        ))
        existing.add(item["key"])
        added += 1
    return added


def serialize_post_activation_task(
    item: PostActivationTask, customer_name: str = "", rule: PostActivationRule | None = None
) -> dict[str, Any]:
    today = date.today()
    if item.status == "DONE":
        alert_state = "DONE"
    elif not item.action_required:
        alert_state = "REVIEW"
    elif item.due_date and item.due_date < today:
        alert_state = "OVERDUE"
    elif item.due_date == today:
        alert_state = "DUE_TODAY"
    else:
        alert_state = "SCHEDULED"
    return {
        "id": str(item.id), "pdc_import_id": str(item.pdc_import_id),
        "customer_id": str(item.customer_id) if item.customer_id else None,
        "customer_name": customer_name, "item_key": item.item_key,
        "item_type": item.item_type, "item_name": item.item_name,
        "customer_code": item.customer_code or "", "contract_code": item.contract_code or "",
        "asset_number": item.asset_number or "", "activation_date": item.activation_date.isoformat(),
        "action_required": item.action_required, "status": item.status,
        "due_date": item.due_date.isoformat() if item.due_date else None,
        "completed_date": item.completed_date.isoformat() if item.completed_date else None,
        "alert_state": alert_state, "notes": item.notes or "",
        "operator": rule.operator if rule else "WINDTRE",
        "can_deactivate": bool(rule and rule.is_active and rule.can_deactivate),
        "configuration_status": "CONFIGURED" if rule and rule.can_deactivate else "NOT_ENABLED",
    }


def serialize_post_activation_rule(item: PostActivationRule) -> dict[str, Any]:
    return {
        "id": str(item.id), "operator": item.operator, "item_type": item.item_type,
        "item_key": item.item_key, "item_name": item.item_name,
        "can_deactivate": item.can_deactivate,
        "default_action_required": item.default_action_required,
        "is_active": item.is_active, "notes": item.notes or "",
        "created_at": item.created_at.isoformat(), "updated_at": item.updated_at.isoformat(),
    }


def serialize_incentive_competition(item: IncentiveCompetition) -> dict[str, Any]:
    return {
        "id": str(item.id), "name": item.name, "operator": item.operator, "market": item.market,
        "dealer_code": item.dealer_code or "", "start_date": item.start_date.isoformat(),
        "end_date": item.end_date.isoformat(), "status": item.status,
        "source_document": item.source_document or "", "configuration": item.configuration or {},
        "created_at": item.created_at.isoformat(),
    }


def serialize_incentive_activation(item: IncentiveActivation, customer_name: str | None = None) -> dict[str, Any]:
    return {
        "id": str(item.id), "competition_id": str(item.competition_id),
        "customer_id": str(item.customer_id) if item.customer_id else None,
        "customer_name": customer_name or "", "activation_date": item.activation_date.isoformat(),
        "source_type": item.source_type, "source_key": item.source_key, "seller_name": item.seller_name or "",
        "track": item.track, "offer": item.offer or "", "asset_number": item.asset_number or "",
        "customer_code": item.customer_code or (item.attributes or {}).get("customer_code", ""),
        "contract_code": item.contract_code or (item.attributes or {}).get("contract_code", ""),
        "monthly_fee": item.monthly_fee_cents / 100, "direct_bonus": item.direct_bonus_cents / 100,
        "attributes": item.attributes or {}, "status": item.status, "notes": item.notes or "",
    }


def incentive_threshold(track_config: dict[str, Any], points: float) -> dict[str, Any] | None:
    reached = [item for item in track_config.get("thresholds", []) if points >= float(item.get("target", 0))]
    return max(reached, key=lambda item: float(item.get("target", 0))) if reached else None


def incentive_activation_points(item: IncentiveActivation) -> float:
    if item.status != "VALID":
        return 0.0
    attrs = item.attributes or {}
    offer = (item.offer or "").upper()
    if attrs.get("excluded") or (item.track == "MOBILE" and "TRAVEL" in offer):
        return 0.0
    if item.track == "MOBILE":
        if attrs.get("very_mobile"):
            return 0.0 if attrs.get("intrabrand_mnp") else 0.5
        points = 1.25 if attrs.get("secure_option") else 1.0
        if attrs.get("phone_included"):
            points += 1.25
        if attrs.get("premium_tied_offer"):
            points += 0.5
        if attrs.get("tied") and attrs.get("piva"):
            points += 1.0
        return points
    if item.track == "FIXED":
        if attrs.get("professional_box"):
            return 4.0
        if attrs.get("second_professional"):
            return 1.5
        if attrs.get("piva") and (attrs.get("first_line", True) or attrs.get("fwa")):
            return 1.5
        return 1.0
    return 1.0


def incentive_report(db, competition: IncentiveCompetition) -> dict[str, Any]:
    activations = db.scalars(
        select(IncentiveActivation)
        .where(IncentiveActivation.competition_id == competition.id)
        .order_by(IncentiveActivation.activation_date, IncentiveActivation.created_at)
    ).all()
    config = competition.configuration or {}
    track_configs = config.get("tracks", {})
    points_by_track: dict[str, float] = defaultdict(float)
    for item in activations:
        points_by_track[item.track] += incentive_activation_points(item)
    reload_events = [item for item in activations if item.track == "RELOAD" and item.status == "VALID"]
    reload_device_total = sum(1 for item in activations if item.status == "VALID" and (item.attributes or {}).get("device_sale"))
    if not reload_device_total:
        reload_device_total = sum(max(1, int((item.attributes or {}).get("device_total", 1))) for item in reload_events)
    reload_rate = (len(reload_events) / reload_device_total * 100) if reload_device_total else 0
    rows = []
    commissioning_total = 0.0
    for item in activations:
        attrs = item.attributes or {}
        points = incentive_activation_points(item)
        track_config = track_configs.get(item.track, {})
        threshold_basis = reload_rate if item.track == "RELOAD" else points_by_track[item.track]
        threshold = incentive_threshold(track_config, threshold_basis)
        access = float(track_config.get("access_threshold", 0))
        eligible = item.status == "VALID" and points > 0 and points_by_track[item.track] >= access
        multiplier = 0.0
        if eligible and threshold:
            if item.track == "MOBILE":
                multiplier = float(threshold.get("mnp_multiplier" if attrs.get("mnp") else "no_mnp_multiplier", 0))
                if attrs.get("tied"):
                    multiplier += float(threshold.get("tied_extra_multiplier", 0))
            elif item.track == "FIXED":
                multiplier = float(threshold.get("convergent_multiplier" if attrs.get("convergent") else "standard_multiplier", 0))
                if attrs.get("piva"):
                    multiplier += 1.0
                if attrs.get("ftth") and attrs.get("first_line", True):
                    multiplier += 1.0
                if attrs.get("fwa") and attrs.get("first_line", True):
                    multiplier += 1.5
            elif item.track == "RELOAD":
                multiplier = float(threshold.get("bonus_multiplier", 1))
        direct_bonus = item.direct_bonus_cents / 100 if eligible else 0.0
        commission = round((item.monthly_fee_cents / 100) * multiplier + direct_bonus * (multiplier if item.track == "RELOAD" and threshold else 1), 2)
        commissioning_total += commission
        rows.append({
            **serialize_incentive_activation(item, db.get(Customer, item.customer_id).business_name if item.customer_id and db.get(Customer, item.customer_id) else ""),
            "points": round(points, 2), "eligible": eligible,
            "threshold": threshold.get("label") if threshold else "Non raggiunta",
            "multiplier": multiplier, "commission": commission,
        })
    track_summaries = []
    for code, track_config in track_configs.items():
        points = round(points_by_track.get(code, 0), 2)
        basis = reload_rate if code == "RELOAD" else points
        reached = incentive_threshold(track_config, basis)
        thresholds = sorted(track_config.get("thresholds", []), key=lambda item: float(item.get("target", 0)))
        next_threshold = next((item for item in thresholds if basis < float(item.get("target", 0))), None)
        track_summaries.append({
            "track": code, "label": track_config.get("label", code), "points": points,
            "events": sum(item.track == code for item in activations),
            "valid_events": sum(item.track == code and item.status == "VALID" for item in activations),
            "access_threshold": track_config.get("access_threshold", 0),
            "reached": reached.get("label") if reached else "Non raggiunta",
            "next_target": next_threshold.get("target") if next_threshold else None,
            "remaining": round(max(0, float(next_threshold.get("target", 0)) - basis), 2) if next_threshold else 0,
            "attachment_rate": round(reload_rate, 2) if code == "RELOAD" else None,
            "commission": round(sum(row["commission"] for row in rows if row["track"] == code), 2),
        })
    return {
        "competition": serialize_incentive_competition(competition),
        "tracks": track_summaries, "activations": rows,
        "total_events": len(activations), "valid_events": sum(item.status == "VALID" for item in activations),
        "commissioning_total": round(commissioning_total, 2),
    }


@asynccontextmanager
async def lifespan(app: FastAPI):
    seed()
    yield


app = FastAPI(title="Cornet ERP API", version="0.2.0", lifespan=lifespan)
app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[os.getenv("FRONTEND_ORIGIN", "http://localhost:5173")],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health")
def health():
    return {"status": "ok", "version": "0.2.0"}


def serialize_store_settings(item: StoreSettings) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "store_name": item.store_name,
        "legal_name": item.legal_name or "",
        "tax_id": item.tax_id or "",
        "fiscal_code": item.fiscal_code or "",
        "dealer_code": item.dealer_code or "",
        "address": item.address or "",
        "city": item.city or "",
        "postal_code": item.postal_code or "",
        "province": item.province or "",
        "phone": item.phone or "",
        "whatsapp": item.whatsapp or "",
        "email": item.email or "",
        "website": item.website or "",
        "logo_url": f"/uploads/{item.logo_path}" if item.logo_path else None,
        "partner_logo_url": f"/uploads/{item.partner_logo_path}" if item.partner_logo_path else None,
        "updated_at": item.updated_at.isoformat(),
    }


def get_or_create_store_settings(db) -> StoreSettings:
    item = db.scalar(select(StoreSettings).limit(1))
    if not item:
        item = StoreSettings()
        db.add(item)
        db.flush()
    return item


DEFAULT_OPERATOR_BRANDS = (
    ("WINDTRE", "WINDTRE"), ("VERY MOBILE", "Very Mobile"),
    ("VODAFONE", "Vodafone"), ("TIM", "TIM"), ("FASTWEB", "Fastweb"),
    ("ILIAD", "Iliad"), ("EOLO", "EOLO"), ("SKY WIFI", "Sky Wifi"),
)


def ensure_operator_brands(db):
    existing = set(db.scalars(select(OperatorBrand.operator)).all())
    for operator, display_name in DEFAULT_OPERATOR_BRANDS:
        if operator not in existing:
            db.add(OperatorBrand(operator=operator, display_name=display_name, is_active=True))
    db.flush()


def serialize_operator_brand(item: OperatorBrand) -> dict[str, Any]:
    return {
        "id": str(item.id), "operator": item.operator, "display_name": item.display_name,
        "logo_url": f"/uploads/{item.logo_path}" if item.logo_path else None,
        "is_active": item.is_active, "updated_at": item.updated_at.isoformat(),
    }


@app.post("/api/v1/auth/login")
def login(data: LoginRequest):
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == data.email))
        if not user or not pwd.verify(data.password, user.hashed_password):
            raise HTTPException(401, "Credenziali non valide")
        token = jwt.encode(
            {"sub": str(user.id), "exp": datetime.now(timezone.utc) + timedelta(hours=8)},
            SECRET_KEY,
            algorithm="HS256",
        )
        return {
            "access_token": token,
            "user": {"id": str(user.id), "email": user.email, "full_name": user.full_name, "role": user.role},
        }


@app.get("/api/v1/settings/store")
def store_settings():
    with SessionLocal() as db:
        item = get_or_create_store_settings(db)
        db.commit()
        db.refresh(item)
        return serialize_store_settings(item)


@app.put("/api/v1/settings/store")
def update_store_settings(data: StoreSettingsRequest):
    if not data.store_name.strip():
        raise HTTPException(422, "Il nome del punto vendita è obbligatorio")
    with SessionLocal() as db:
        item = get_or_create_store_settings(db)
        for field, value in data.model_dump().items():
            setattr(item, field, value.strip() if isinstance(value, str) else value)
        db.commit()
        db.refresh(item)
        return serialize_store_settings(item)


@app.post("/api/v1/settings/store/logo")
async def upload_store_logo(file: UploadFile = File(...)):
    allowed_types = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}
    extension = allowed_types.get(file.content_type or "")
    if not extension:
        raise HTTPException(422, "Carica un logo PNG, JPG o WEBP")
    contents = await file.read()
    if len(contents) > 5 * 1024 * 1024:
        raise HTTPException(413, "Il logo supera il limite di 5 MB")
    filename = f"store-logo{extension}"
    target = UPLOAD_DIR / filename
    target.write_bytes(contents)
    with SessionLocal() as db:
        item = get_or_create_store_settings(db)
        previous = item.logo_path
        item.logo_path = filename
        db.commit()
        db.refresh(item)
        if previous and previous != filename:
            previous_path = UPLOAD_DIR / previous
            if previous_path.is_file():
                previous_path.unlink()
        return serialize_store_settings(item)


@app.post("/api/v1/settings/store/partner-logo")
async def upload_partner_logo(file: UploadFile = File(...)):
    allowed_types = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp"}
    extension = allowed_types.get(file.content_type or "")
    if not extension:
        raise HTTPException(422, "Carica un logo PNG, JPG o WEBP")
    contents = await file.read()
    if len(contents) > 5 * 1024 * 1024:
        raise HTTPException(413, "Il logo supera il limite di 5 MB")
    filename = f"partner-logo{extension}"
    (UPLOAD_DIR / filename).write_bytes(contents)
    with SessionLocal() as db:
        item = get_or_create_store_settings(db)
        previous = item.partner_logo_path
        item.partner_logo_path = filename
        db.commit()
        db.refresh(item)
        if previous and previous != filename:
            previous_path = UPLOAD_DIR / previous
            if previous_path.is_file():
                previous_path.unlink()
        return serialize_store_settings(item)


@app.get("/api/v1/settings/operator-brands")
def operator_brands():
    with SessionLocal() as db:
        ensure_operator_brands(db)
        db.commit()
        items = db.scalars(select(OperatorBrand).order_by(OperatorBrand.display_name)).all()
        return [serialize_operator_brand(item) for item in items]


@app.post("/api/v1/settings/operator-brands/{operator}/logo")
async def upload_operator_brand_logo(operator: str, file: UploadFile = File(...)):
    allowed_types = {"image/png": ".png", "image/jpeg": ".jpg", "image/webp": ".webp", "image/svg+xml": ".svg"}
    extension = allowed_types.get(file.content_type or "")
    if not extension:
        raise HTTPException(422, "Carica un logo PNG, JPG, WEBP o SVG")
    contents = await file.read()
    if len(contents) > 5 * 1024 * 1024:
        raise HTTPException(413, "Il logo supera il limite di 5 MB")
    operator_code = clean(operator).upper()
    with SessionLocal() as db:
        ensure_operator_brands(db)
        item = db.scalar(select(OperatorBrand).where(OperatorBrand.operator == operator_code))
        if not item:
            item = OperatorBrand(operator=operator_code, display_name=operator_code, is_active=True)
            db.add(item)
            db.flush()
        safe_code = normalize_header(operator_code).lower() or "operatore"
        filename = f"operator-{safe_code}-{item.id}{extension}"
        (UPLOAD_DIR / filename).write_bytes(contents)
        previous = item.logo_path
        item.logo_path = filename
        db.commit()
        db.refresh(item)
        if previous and previous != filename:
            previous_path = UPLOAD_DIR / previous
            if previous_path.is_file():
                previous_path.unlink()
        return serialize_operator_brand(item)


def get_or_create_letterhead(db) -> LetterheadTemplate:
    item = db.scalar(select(LetterheadTemplate).limit(1))
    if not item:
        store = get_or_create_store_settings(db)
        address = ", ".join(part for part in [store.address, " ".join(part for part in [store.postal_code, store.city] if part), store.province] if part)
        item = LetterheadTemplate(
            company_name=store.legal_name or store.store_name,
            company_address=address, tax_id=store.tax_id or store.fiscal_code,
            phone=store.phone, email=store.email, website=store.website,
            logo_url=f"/uploads/{store.logo_path}" if store.logo_path else None,
        )
        db.add(item)
        db.flush()
    return item


def serialize_letterhead(item: LetterheadTemplate) -> dict[str, Any]:
    return {
        "id": str(item.id), "company_name": item.company_name, "company_address": item.company_address or "",
        "tax_id": item.tax_id or "", "phone": item.phone or "", "email": item.email or "",
        "pec": item.pec or "", "website": item.website or "", "logo_url": item.logo_url or "",
        "logo_size": item.logo_size, "logo_horizontal": item.logo_horizontal, "logo_vertical": item.logo_vertical,
        "recipient_offset_mm": item.recipient_offset_mm, "primary_color": item.primary_color,
        "updated_at": item.updated_at.isoformat(),
    }


def validate_letterhead(data: LetterheadTemplateRequest):
    if not data.company_name.strip():
        raise HTTPException(422, "Il nome dell'azienda è obbligatorio")
    if not 32 <= data.logo_size <= 160:
        raise HTTPException(422, "La dimensione del logo deve essere compresa tra 32 e 160 px")
    if data.logo_horizontal not in {"left", "center", "right"} or data.logo_vertical not in {"top", "center", "bottom"}:
        raise HTTPException(422, "Allineamento logo non valido")
    if not 0 <= data.recipient_offset_mm <= 100:
        raise HTTPException(422, "L'offset destinatario deve essere compreso tra 0 e 100 mm")


@app.get("/api/v1/letterhead-template")
def letterhead_template():
    with SessionLocal() as db:
        item = get_or_create_letterhead(db)
        db.commit()
        db.refresh(item)
        return serialize_letterhead(item)


@app.put("/api/v1/letterhead-template")
def update_letterhead_template(data: LetterheadTemplateRequest):
    validate_letterhead(data)
    with SessionLocal() as db:
        item = get_or_create_letterhead(db)
        for field, value in data.model_dump().items():
            setattr(item, field, value.strip() if isinstance(value, str) else value)
        db.commit()
        db.refresh(item)
        return serialize_letterhead(item)


@app.post("/api/v1/letterhead-template/sync-store")
def sync_letterhead_store():
    with SessionLocal() as db:
        store = get_or_create_store_settings(db)
        item = get_or_create_letterhead(db)
        item.company_name = store.legal_name or store.store_name
        item.company_address = ", ".join(part for part in [store.address, " ".join(part for part in [store.postal_code, store.city] if part), store.province] if part)
        item.tax_id, item.phone, item.email, item.website = store.tax_id or store.fiscal_code, store.phone, store.email, store.website
        item.logo_url = f"/uploads/{store.logo_path}" if store.logo_path else item.logo_url
        db.commit()
        db.refresh(item)
        return serialize_letterhead(item)


def draw_wrapped_text(pdf, text: str, x: float, y: float, max_width: float, font: str = "Helvetica", size: int = 10, leading: int = 14):
    pdf.setFont(font, size)
    words, line, lines = text.split(), "", []
    for word in words:
        candidate = f"{line} {word}".strip()
        if pdf.stringWidth(candidate, font, size) <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    for value in lines:
        pdf.drawString(x, y, value)
        y -= leading
    return y


@app.post("/api/v1/letterhead-template/pdf")
def letterhead_pdf(data: LetterheadGenerateRequest):
    if data.mode not in {"A4", "DL"}:
        raise HTTPException(422, "Formato non valido")
    if not data.recipient_name.strip() or not data.recipient_address.strip():
        raise HTTPException(422, "Destinatario e indirizzo sono obbligatori")
    with SessionLocal() as db:
        item = get_or_create_letterhead(db)
        page_size = A4 if data.mode == "A4" else (215 * mm, 110 * mm)
        width, height = page_size
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=page_size)
        primary = colors.HexColor(item.primary_color if re.fullmatch(r"#[0-9A-Fa-f]{6}", item.primary_color or "") else "#4f46e5")
        margin = 20 * mm if data.mode == "A4" else 10 * mm
        logo_path = None
        if item.logo_url and item.logo_url.startswith("/uploads/"):
            candidate = UPLOAD_DIR / Path(item.logo_url).name
            logo_path = candidate if candidate.is_file() else None
        logo_height = min(item.logo_size * .264583 * mm, 35 * mm if data.mode == "A4" else 18 * mm)
        logo_width = logo_height * 2.4
        logo_x = margin if item.logo_horizontal == "left" else (width - logo_width) / 2 if item.logo_horizontal == "center" else width - margin - logo_width
        logo_y = height - margin - logo_height if item.logo_vertical == "top" else (height - logo_height) / 2 if item.logo_vertical == "center" else margin + 15 * mm
        if logo_path:
            pdf.drawImage(ImageReader(str(logo_path)), logo_x, logo_y, width=logo_width, height=logo_height, preserveAspectRatio=True, anchor="c", mask="auto")
        if data.mode == "A4":
            company_y = height - margin - (logo_height + 5 * mm if logo_path and item.logo_horizontal == "left" else 3 * mm)
            pdf.setFillColor(primary); pdf.setFont("Helvetica-Bold", 14); pdf.drawString(margin, company_y, item.company_name)
            pdf.setStrokeColor(primary); pdf.setLineWidth(1.2); pdf.line(margin, company_y - 5 * mm, width - margin, company_y - 5 * mm)
            recipient_x = max(margin, 118 * mm - item.recipient_offset_mm * mm)
            recipient_y = height - 72 * mm
            pdf.setFillColor(colors.black); pdf.setFont("Helvetica-Bold", 11); pdf.drawString(recipient_x, recipient_y, data.recipient_name)
            draw_wrapped_text(pdf, data.recipient_address, recipient_x, recipient_y - 6 * mm, width - margin - recipient_x, size=10)
            body = data.body_text or "Spazio riservato al contenuto della comunicazione."
            body_y = height - 115 * mm
            for paragraph in body.splitlines():
                body_y = draw_wrapped_text(pdf, paragraph or " ", margin, body_y, width - 2 * margin, size=10, leading=15) - 5
            footer_y = 18 * mm
            pdf.setStrokeColor(primary); pdf.line(margin, footer_y + 10 * mm, width - margin, footer_y + 10 * mm)
            pdf.setFillColor(colors.black); pdf.setFont("Helvetica", 7.5)
            footer = " | ".join(part for part in [item.company_name, item.company_address, f"P.IVA/CF {item.tax_id}" if item.tax_id else None, item.phone, item.email, item.pec, item.website] if part)
            draw_wrapped_text(pdf, footer, margin, footer_y + 5 * mm, width - 2 * margin, size=7.5, leading=9)
        else:
            sender_y = logo_y - 3 * mm if logo_path else height - margin - 4 * mm
            pdf.setFillColor(primary); pdf.setFont("Helvetica-Bold", 10); pdf.drawString(margin, sender_y, item.company_name)
            draw_wrapped_text(pdf, item.company_address or "", margin, sender_y - 5 * mm, 90 * mm, size=7.5, leading=9)
            recipient_x = max(margin, 115 * mm - item.recipient_offset_mm * mm)
            recipient_y = 42 * mm
            pdf.setFillColor(colors.black); pdf.setFont("Helvetica-Bold", 11); pdf.drawString(recipient_x, recipient_y, data.recipient_name)
            draw_wrapped_text(pdf, data.recipient_address, recipient_x, recipient_y - 6 * mm, width - margin - recipient_x, size=9.5, leading=12)
        pdf.showPage(); pdf.save(); buffer.seek(0)
        filename = "carta-intestata-a4.pdf" if data.mode == "A4" else "busta-intestata-dl.pdf"
        return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{filename}"'})


TERMINAL_TYPES = {"SMARTPHONE", "TABLET", "ROUTER", "ACCESSORIO"}
CUSTOMER_BANDS = {"START", "SMERALDO", "RUBINO", "ZAFFIRO"}
TERMINAL_INVENTORY_ALIASES = {
    "model": {"MODELLO", "COD_CL_TM", "CODCLTM"},
    "gsi": {"CODICE_GSI", "CODICEGSI"},
    "pieces": {"PEZZI", "QUANTITA", "QTY"},
    "skip": {"SKIP_DISPONIBILITA", "SKIPDISPONIBILITA", "SKIP"},
    "notes": {"NOTE", "COMMENTI"},
}
GA_ALIASES = {
    "model": {"MODELLO", "MODELLO_TERMINALE", "TERMINALE", "DESCRIZIONE"},
    "gsi": {"CODICE_GSI", "GSI", "COD_GSI", "CODICE"},
    "offer": {"OFFERTA", "OFFERTA_ASSOCIATA", "PROFILO"},
    "list_price": {"PREZZO_LISTINO", "PREZZO_DI_LISTINO", "LISTINO"},
    "standard": {"RATA_MENSILE_LISTINO_STANDARD", "RATA_STANDARD"},
    "kasko": {"KASKO", "CANONE_KASKO"},
    "kasko_premium": {"KASKO_PREMIUM", "CANONE_KASKO_PREMIUM"},
    "upfront": {"ANTICIPO"},
    "monthly": {"RATA_MENSILE", "RATA"},
    "final": {"RATA_FINALE", "MAXIRATA", "MAXI_RATA"},
    "discount": {"SCONTO", "SCONTO_PERCENTUALE", "SCONTO"},
    "promotion": {"PROMOZIONE", "NOME_PROMOZIONE", "PROMO"},
    "promotion_id": {"ID_PROMOZIONE", "ID_PROMO"},
}


def excel_sheet_rows(contents: bytes, filename: str, sheet_name: str) -> list[list[Any]]:
    try:
        if filename.lower().endswith(".xls"):
            workbook = xlrd.open_workbook(file_contents=contents)
            if sheet_name not in workbook.sheet_names():
                raise HTTPException(422, "Foglio Excel non trovato")
            sheet = workbook.sheet_by_name(sheet_name)
            return [sheet.row_values(index) for index in range(sheet.nrows)]
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise HTTPException(422, "Foglio Excel non trovato")
        return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"File Excel non leggibile: {exc}")


def find_terminal_inventory_header(rows: list[list[Any]]) -> tuple[int, list[str]]:
    known = {alias for aliases in TERMINAL_INVENTORY_ALIASES.values() for alias in aliases}
    candidates = []
    for index, values in enumerate(rows[:25]):
        headers = [normalize_header(value) for value in values]
        candidates.append((sum(header in known for header in headers), index, headers))
    if not candidates:
        raise HTTPException(422, {"message": "Il foglio è vuoto", "missing_columns": ["Modello", "Codice GSI"], "detected_columns": []})
    _, index, headers = max(candidates, key=lambda candidate: candidate[0])
    missing = []
    if not any(header in TERMINAL_INVENTORY_ALIASES["model"] for header in headers):
        missing.append("Modello")
    if not any(header in TERMINAL_INVENTORY_ALIASES["gsi"] for header in headers):
        missing.append("Codice GSI")
    if missing:
        raise HTTPException(422, {
            "message": "Tracciato Excel non valido",
            "missing_columns": missing,
            "detected_columns": [header for header in headers if header],
        })
    return index, headers


def inventory_integer(value: Any) -> int:
    text = clean(value).replace(".", "").replace(",", ".")
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def inventory_boolean(value: Any) -> bool:
    return normalize_header(value) in {"1", "TRUE", "VERO", "SI", "S", "YES", "Y", "X"}


def serialize_terminal_inventory(item: TerminalInventoryItem) -> dict[str, Any]:
    available = item.pieces > 0 or item.skip_availability
    return {
        "id": str(item.id),
        "channel": item.channel,
        "model": item.model,
        "gsi_code": item.gsi_code,
        "pieces": item.pieces,
        "skip_availability": item.skip_availability,
        "notes": item.notes or "",
        "available": available,
        "availability_label": "Disponibile" if item.pieces > 0 else ("Su ordinazione" if item.skip_availability else "Non disponibile"),
        "updated_at": item.updated_at.isoformat(),
    }


@app.get("/api/v1/terminal-inventory")
def terminal_inventory(channel: str, search: str = ""):
    if channel not in {"GA", "CB"}:
        raise HTTPException(422, "Canale non valido")
    with SessionLocal() as db:
        query = select(TerminalInventoryItem).where(TerminalInventoryItem.channel == channel)
        if search.strip():
            pattern = f"%{search.strip()}%"
            query = query.where(or_(
                TerminalInventoryItem.model.ilike(pattern),
                TerminalInventoryItem.gsi_code.ilike(pattern),
                TerminalInventoryItem.notes.ilike(pattern),
            ))
        items = db.scalars(query.order_by(TerminalInventoryItem.model, TerminalInventoryItem.gsi_code)).all()
        meta = db.get(TerminalInventoryMetadata, channel)
        totals = db.execute(
            select(func.count(TerminalInventoryItem.id), func.coalesce(func.sum(TerminalInventoryItem.pieces), 0))
            .where(TerminalInventoryItem.channel == channel)
        ).one()
        return {
            "items": [serialize_terminal_inventory(item) for item in items],
            "kpi": {"models": totals[0], "pieces": totals[1]},
            "metadata": {
                "file_name": meta.file_name,
                "sheet_name": meta.sheet_name,
                "row_count": meta.row_count,
                "updated_at": meta.updated_at.isoformat(),
            } if meta else None,
        }


@app.post("/api/v1/terminal-inventory/sheets")
async def terminal_inventory_sheets(file: UploadFile = File(...)):
    contents = await file.read()
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(422, "Carica un file Excel .xlsx o .xls")
    try:
        sheets = xlrd.open_workbook(file_contents=contents).sheet_names() if filename.lower().endswith(".xls") else load_workbook(BytesIO(contents), read_only=True, data_only=True).sheetnames
    except Exception as exc:
        raise HTTPException(422, f"File Excel non leggibile: {exc}")
    return {"sheets": sheets}


@app.post("/api/v1/terminal-inventory/import")
async def import_terminal_inventory(
    file: UploadFile = File(...),
    sheet_name: str = Form(...),
    channel: str = Form(...),
):
    if channel not in {"GA", "CB"}:
        raise HTTPException(422, "Canale non valido")
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(422, "Carica un file Excel .xlsx o .xls")
    contents = await file.read()
    sheet_rows = excel_sheet_rows(contents, filename, sheet_name)
    header_index, headers = find_terminal_inventory_header(sheet_rows)
    parsed: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    for row_number, values in enumerate(sheet_rows[header_index + 1:], header_index + 2):
        raw = {headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index] and clean(value)}
        if not raw:
            continue
        model = mapped_value(raw, TERMINAL_INVENTORY_ALIASES["model"])
        gsi = mapped_value(raw, TERMINAL_INVENTORY_ALIASES["gsi"])
        if not model or not gsi:
            skipped.append({"row": row_number, "error": "Modello o Codice GSI mancante"})
            continue
        parsed.append({
            "model": model,
            "gsi_code": gsi,
            "pieces": inventory_integer(mapped_value(raw, TERMINAL_INVENTORY_ALIASES["pieces"])),
            "skip_availability": inventory_boolean(mapped_value(raw, TERMINAL_INVENTORY_ALIASES["skip"])),
            "notes": mapped_value(raw, TERMINAL_INVENTORY_ALIASES["notes"]) or None,
        })
    if not parsed:
        raise HTTPException(422, {
            "message": "Nessun dato valido: la giacenza precedente non è stata modificata",
            "missing_columns": [],
            "detected_columns": [header for header in headers if header],
        })
    now = datetime.now(timezone.utc)
    with SessionLocal() as db:
        try:
            db.execute(delete(TerminalInventoryItem).where(TerminalInventoryItem.channel == channel))
            db.add_all(TerminalInventoryItem(channel=channel, updated_at=now, **row) for row in parsed)
            metadata = db.get(TerminalInventoryMetadata, channel) or TerminalInventoryMetadata(channel=channel)
            metadata.file_name = filename
            metadata.sheet_name = sheet_name
            metadata.row_count = len(parsed)
            metadata.updated_at = now
            db.add(metadata)
            db.commit()
        except Exception:
            db.rollback()
            raise HTTPException(500, "Importazione non salvata: la giacenza precedente è rimasta invariata")
    return {"channel": channel, "imported": len(parsed), "skipped": skipped, "updated_at": now.isoformat()}


def serialize_terminal(item: TerminalCatalogItem) -> dict[str, Any]:
    return {
        "id": str(item.id), "channel": item.channel, "brand": item.brand or "", "model": item.model,
        "memory": item.memory or "", "gsi_code": item.gsi_code, "product_type": item.product_type,
        "offer_name": item.offer_name or "", "customer_band": item.customer_band or "",
        "list_price": item.list_price_cents / 100, "standard_installment": item.standard_installment_cents / 100,
        "kasko": item.kasko_cents / 100, "kasko_premium": item.kasko_premium_cents / 100,
        "upfront": item.upfront_cents / 100, "monthly_installment": item.monthly_installment_cents / 100,
        "final_installment": item.final_installment_cents / 100, "discount_percent": item.discount_percent,
        "promotion_name": item.promotion_name or "", "promotion_id": item.promotion_id or "",
        "updated_at": item.updated_at.isoformat(),
    }


def terminal_query(db, channel: str, search: str = "", product_type: str | None = None, customer_band: str | None = None):
    query = select(TerminalCatalogItem).where(TerminalCatalogItem.channel == channel).order_by(TerminalCatalogItem.brand, TerminalCatalogItem.model)
    if product_type:
        query = query.where(TerminalCatalogItem.product_type == product_type)
    if customer_band:
        query = query.where(TerminalCatalogItem.customer_band == customer_band)
    if search.strip():
        pattern = f"%{search.strip()}%"
        query = query.where(or_(TerminalCatalogItem.model.ilike(pattern), TerminalCatalogItem.brand.ilike(pattern), TerminalCatalogItem.gsi_code.ilike(pattern), TerminalCatalogItem.offer_name.ilike(pattern)))
    return query


@app.get("/api/v1/terminals")
def terminals(channel: str, search: str = "", product_type: str | None = None, customer_band: str | None = None):
    if channel not in {"GA", "CB"}:
        raise HTTPException(422, "Canale non valido")
    with SessionLocal() as db:
        items = db.scalars(terminal_query(db, channel, search, product_type, customer_band)).all()
        meta = db.get(TerminalCatalogMetadata, channel)
        return {
            "items": [serialize_terminal(item) for item in items],
            "metadata": {"file_name": meta.file_name, "sheet_name": meta.sheet_name, "row_count": meta.row_count, "updated_at": meta.updated_at.isoformat()} if meta else None,
        }


@app.post("/api/v1/terminals-ga/sheets")
async def terminal_ga_sheets(file: UploadFile = File(...)):
    contents = await file.read()
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm", ".xls")):
        raise HTTPException(422, "Carica un file Excel .xlsx, .xlsm o .xls")
    try:
        sheets = xlrd.open_workbook(file_contents=contents).sheet_names() if filename.endswith(".xls") else load_workbook(BytesIO(contents), read_only=True, data_only=True).sheetnames
    except Exception as exc:
        raise HTTPException(422, f"File Excel non leggibile: {exc}")
    return {"sheets": sheets}


def find_ga_header(rows):
    best = None
    known = {value for values in GA_ALIASES.values() for value in values}
    for row_number, values in enumerate(rows[:25], 1):
        headers = [normalize_header(value) for value in values]
        score = sum(header in known for header in headers)
        if not best or score > best[0]:
            best = (score, row_number, headers)
    if not best or best[0] < 2 or not any(value in GA_ALIASES["model"] for value in best[2]) or not any(value in GA_ALIASES["gsi"] for value in best[2]):
        raise HTTPException(422, "Colonne obbligatorie Modello e Codice GSI non riconosciute")
    return best[1], best[2]


@app.post("/api/v1/terminals-ga/import")
async def import_terminals_ga(file: UploadFile = File(...), sheet_name: str = Form(...)):
    contents = await file.read()
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm", ".xls")):
        raise HTTPException(422, "Carica un file Excel .xlsx, .xlsm o .xls")
    try:
        if filename.endswith(".xls"):
            workbook = xlrd.open_workbook(file_contents=contents)
            if sheet_name not in workbook.sheet_names():
                raise HTTPException(422, "Foglio Excel non trovato")
            source = workbook.sheet_by_name(sheet_name)
            sheet_rows = [source.row_values(index) for index in range(source.nrows)]
        else:
            workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
            if sheet_name not in workbook.sheetnames:
                raise HTTPException(422, "Foglio Excel non trovato")
            sheet_rows = list(workbook[sheet_name].iter_rows(values_only=True))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"File Excel non leggibile: {exc}")
    header_row, headers = find_ga_header(sheet_rows)
    rows, skipped = [], []
    for row_number, values in enumerate(sheet_rows[header_row:], header_row + 1):
        raw = {headers[i]: clean(value) for i, value in enumerate(values) if i < len(headers) and headers[i] and clean(value)}
        if not raw:
            continue
        model, gsi = mapped_value(raw, GA_ALIASES["model"]), mapped_value(raw, GA_ALIASES["gsi"])
        if not model or not gsi:
            skipped.append({"row": row_number, "error": "Modello o Codice GSI mancante"})
            continue
        rows.append(TerminalCatalogItem(
            channel="GA", model=model, gsi_code=gsi, product_type="SMARTPHONE",
            offer_name=mapped_value(raw, GA_ALIASES["offer"]) or None,
            list_price_cents=tariff_cents(mapped_value(raw, GA_ALIASES["list_price"])),
            standard_installment_cents=tariff_cents(mapped_value(raw, GA_ALIASES["standard"])),
            kasko_cents=tariff_cents(mapped_value(raw, GA_ALIASES["kasko"])),
            kasko_premium_cents=tariff_cents(mapped_value(raw, GA_ALIASES["kasko_premium"])),
            upfront_cents=tariff_cents(mapped_value(raw, GA_ALIASES["upfront"])),
            monthly_installment_cents=tariff_cents(mapped_value(raw, GA_ALIASES["monthly"])),
            final_installment_cents=tariff_cents(mapped_value(raw, GA_ALIASES["final"])),
            discount_percent=round(parse_monthly_fee(mapped_value(raw, GA_ALIASES["discount"]))),
            promotion_name=mapped_value(raw, GA_ALIASES["promotion"]) or None,
            promotion_id=mapped_value(raw, GA_ALIASES["promotion_id"]) or None,
        ))
    if not rows:
        raise HTTPException(422, "Il foglio non contiene terminali validi")
    with SessionLocal() as db:
        db.execute(delete(TerminalCatalogItem).where(TerminalCatalogItem.channel == "GA"))
        db.add_all(rows)
        metadata = db.get(TerminalCatalogMetadata, "GA") or TerminalCatalogMetadata(channel="GA")
        metadata.file_name, metadata.sheet_name, metadata.row_count, metadata.updated_at = file.filename, sheet_name, len(rows), datetime.now(timezone.utc)
        db.add(metadata)
        db.commit()
    return {"imported": len(rows), "skipped": skipped[:100], "sheet_name": sheet_name}


@app.get("/api/v1/terminals-ga.pdf")
def terminals_ga_pdf(search: str = ""):
    with SessionLocal() as db:
        items = db.scalars(terminal_query(db, "GA", search)).all()
        store = get_or_create_store_settings(db)
        meta = db.get(TerminalCatalogMetadata, "GA")
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
        styles = getSampleStyleSheet()
        story = [Paragraph(f"<b>{store.legal_name or store.store_name} - Listino Terminali GA</b>", styles["Title"]), Paragraph(f"Aggiornato: {meta.updated_at.strftime('%d/%m/%Y %H:%M') if meta else '—'}", styles["BodyText"]), Spacer(1, 4*mm)]
        data = [["Modello","GSI","Offerta","Listino","Anticipo","Rata","Finale","Kasko","Sconto"]]
        data.extend([[item.model,item.gsi_code,item.offer_name or "—",f"€ {item.list_price_cents/100:.2f}",f"€ {item.upfront_cents/100:.2f}",f"€ {item.monthly_installment_cents/100:.2f}",f"€ {item.final_installment_cents/100:.2f}",f"€ {item.kasko_cents/100:.2f}",f"{item.discount_percent}%"] for item in items])
        table = Table(data, repeatRows=1, colWidths=[45*mm,25*mm,42*mm,20*mm,20*mm,20*mm,20*mm,20*mm,16*mm])
        table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1e3a8a")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#cbd5e1")),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f8fafc")])]))
        story.append(table); doc.build(story); buffer.seek(0)
        return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": 'inline; filename="listino-terminali-ga.pdf"'})


def apply_cb(item: TerminalCatalogItem, data: TerminalCBRequest):
    if data.product_type not in TERMINAL_TYPES or data.customer_band not in CUSTOMER_BANDS:
        raise HTTPException(422, "Tipologia o fascia non valida")
    if not data.brand.strip() or not data.model.strip() or not data.gsi_code.strip():
        raise HTTPException(422, "Marca, modello e Codice GSI sono obbligatori")
    item.brand, item.model, item.gsi_code = data.brand.strip(), data.model.strip(), data.gsi_code.strip()
    item.memory, item.product_type, item.customer_band = (data.memory or "").strip() or None, data.product_type, data.customer_band
    item.list_price_cents, item.upfront_cents = tariff_cents(data.list_price), tariff_cents(data.upfront)
    item.monthly_installment_cents, item.final_installment_cents = tariff_cents(data.monthly_installment), tariff_cents(data.final_installment)


@app.post("/api/v1/terminals-cb")
def create_terminal_cb(data: TerminalCBRequest):
    with SessionLocal() as db:
        item = TerminalCatalogItem(channel="CB", model=data.model, gsi_code=data.gsi_code)
        apply_cb(item, data); db.add(item); db.commit(); db.refresh(item)
        return serialize_terminal(item)


@app.put("/api/v1/terminals-cb/{item_id}")
def update_terminal_cb(item_id: uuid.UUID, data: TerminalCBRequest):
    with SessionLocal() as db:
        item = db.get(TerminalCatalogItem, item_id)
        if not item or item.channel != "CB":
            raise HTTPException(404, "Terminale CB non trovato")
        apply_cb(item, data); db.commit(); db.refresh(item)
        return serialize_terminal(item)


@app.delete("/api/v1/terminals-cb/{item_id}")
def delete_terminal_cb(item_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(TerminalCatalogItem, item_id)
        if not item or item.channel != "CB":
            raise HTTPException(404, "Terminale CB non trovato")
        db.delete(item); db.commit(); return {"ok": True}


@app.delete("/api/v1/terminals-cb")
def clear_terminals_cb(customer_band: str | None = None):
    with SessionLocal() as db:
        query = delete(TerminalCatalogItem).where(TerminalCatalogItem.channel == "CB")
        if customer_band:
            query = query.where(TerminalCatalogItem.customer_band == customer_band)
        result = db.execute(query); db.commit(); return {"deleted": result.rowcount}


DEFAULT_CB = [
    ("Apple","iPhone 16","128GB","GSI-IP16-128","SMARTPHONE"),
    ("Apple","iPhone 16 Pro","256GB","GSI-IP16P-256","SMARTPHONE"),
    ("Apple","iPhone 17","256GB","GSI-IP17-256","SMARTPHONE"),
    ("Apple","iPad Air","128GB","GSI-IPADAIR-128","TABLET"),
    ("Samsung","Galaxy S25","256GB","GSI-S25-256","SMARTPHONE"),
    ("Samsung","Galaxy S26","256GB","GSI-S26-256","SMARTPHONE"),
    ("Samsung","Galaxy Z Fold","512GB","GSI-ZFOLD-512","SMARTPHONE"),
    ("Xiaomi","Xiaomi 15","256GB","GSI-XIAOMI15-256","SMARTPHONE"),
    ("Honor","Honor Magic","256GB","GSI-HONOR-MAGIC","SMARTPHONE"),
    ("Motorola","Edge","256GB","GSI-MOTO-EDGE","SMARTPHONE"),
    ("TCL","TCL 60","128GB","GSI-TCL60-128","SMARTPHONE"),
    ("ZTE","Router 5G","—","GSI-ZTE-5G","ROUTER"),
    ("Samsung","Galaxy Buds","—","GSI-BUDS","ACCESSORIO"),
]


@app.post("/api/v1/terminals-cb/reset")
def reset_terminals_cb():
    with SessionLocal() as db:
        db.execute(delete(TerminalCatalogItem).where(TerminalCatalogItem.channel == "CB"))
        for band in ("START","SMERALDO","RUBINO","ZAFFIRO"):
            for brand, model, memory, gsi, product_type in DEFAULT_CB:
                db.add(TerminalCatalogItem(channel="CB",brand=brand,model=model,memory=memory,gsi_code=f"{gsi}-{band}",product_type=product_type,customer_band=band))
        metadata = db.get(TerminalCatalogMetadata, "CB") or TerminalCatalogMetadata(channel="CB")
        metadata.file_name, metadata.sheet_name, metadata.row_count, metadata.updated_at = "Matrice predefinita", None, len(DEFAULT_CB)*4, datetime.now(timezone.utc)
        db.add(metadata); db.commit()
        return {"created": len(DEFAULT_CB)*4}


WINDTRE_PANEL_TEMPLATES = [
    {"key": "DOCUMENTI", "title": "Richiesta Documenti", "recipient": "", "category": "CLIENTE"},
    {"key": "DISDETTA_1928", "title": "Disdetta Business SME (1928)", "recipient": "CustomerCareWindTreBusiness@pec.windtre.it", "category": "PEC"},
    {"key": "CESSAZIONE_159", "title": "Cessazione Consumer/Micro (159)", "recipient": "servizioclienti159@pec.windtre.it", "category": "PEC"},
    {"key": "SOSTITUZIONE_SIM", "title": "Sostituzione SIM (Furto / Smarrimento)", "recipient": "CustomerCareWindTreBusiness@pec.windtre.it", "category": "PEC"},
    {"key": "CAMBIO_AGENZIA", "title": "Richiesta Cambio Agenzia", "recipient": "lucacorniello@partner.windtre.it", "category": "PARTNER"},
    {"key": "BUSINESS_CONSUMER", "title": "Passaggio Business -> Consumer", "recipient": "lucacorniello@partner.windtre.it", "category": "PARTNER", "printable": True},
    {"key": "SPEDIZIONE_TERMINALI", "title": "Autorizzazione Spedizione Terminali", "recipient": "lucacorniello@partner.windtre.it", "category": "PARTNER"},
    {"key": "SUBENTRO", "title": "Richiesta Subentro", "recipient": "", "category": "CLIENTE"},
    {"key": "BENVENUTO", "title": "Benvenuto & Conferma", "recipient": "", "category": "CLIENTE"},
    {"key": "SOLLECITO", "title": "Sollecito Pratica", "recipient": "", "category": "CLIENTE"},
]
WINDTRE_PANEL_TEMPLATE_MAP = {item["key"]: item for item in WINDTRE_PANEL_TEMPLATES}
WINDTRE_PANEL_TAX_REQUIRED = {"DOCUMENTI", "DISDETTA_1928", "CESSAZIONE_159", "SOSTITUZIONE_SIM", "CAMBIO_AGENZIA", "BUSINESS_CONSUMER", "SPEDIZIONE_TERMINALI"}


def serialize_windtre_panel_request(item: WindTrePanelRequest) -> dict[str, Any]:
    return {
        "id": str(item.id), "customer_id": str(item.customer_id) if item.customer_id else None,
        "customer_name": item.customer_name, "template_key": item.template_key,
        "template_title": item.template_title, "recipient": item.recipient or "",
        "subject": item.subject, "body": item.body, "form_data": item.form_data or {},
        "sent_at": item.sent_at.isoformat(), "response_date": item.response_date.isoformat() if item.response_date else None,
        "status": item.status, "operator_uid": item.operator_uid,
    }


@app.get("/api/v1/windtre-panel/templates")
def windtre_panel_templates():
    return WINDTRE_PANEL_TEMPLATES


@app.get("/api/v1/windtre-panel/requests")
def windtre_panel_requests(search: str = "", status: str | None = None, limit: int = Query(50, ge=1, le=200)):
    with SessionLocal() as db:
        query = select(WindTrePanelRequest).order_by(WindTrePanelRequest.sent_at.desc()).limit(limit)
        if search.strip():
            query = query.where(WindTrePanelRequest.customer_name.ilike(f"%{search.strip()}%"))
        if status:
            query = query.where(WindTrePanelRequest.status == status)
        return [serialize_windtre_panel_request(item) for item in db.scalars(query).all()]


@app.post("/api/v1/windtre-panel/requests")
def create_windtre_panel_request(data: WindTrePanelRequestCreate):
    template = WINDTRE_PANEL_TEMPLATE_MAP.get(data.template_key)
    if not template:
        raise HTTPException(422, "Template WindTre non valido")
    if not data.customer_name.strip() or not data.subject.strip() or not data.body.strip():
        raise HTTPException(422, "Cliente, oggetto e corpo della comunicazione sono obbligatori")
    tax_id = re.sub(r"\D", "", str(data.form_data.get("tax_id") or ""))
    if data.template_key in WINDTRE_PANEL_TAX_REQUIRED and len(tax_id) != 11:
        raise HTTPException(422, "La Partita IVA deve contenere esattamente 11 cifre")
    if data.template_key == "DOCUMENTI":
        phone = re.sub(r"\D", "", str(data.form_data.get("otp_phone") or ""))
        if len(phone) not in {9, 10}:
            raise HTTPException(422, "Il telefono per OTP deve contenere 9 o 10 cifre")
    with SessionLocal() as db:
        if data.customer_id and not db.get(Customer, data.customer_id):
            raise HTTPException(404, "Cliente non trovato")
        item = WindTrePanelRequest(
            customer_id=data.customer_id, customer_name=data.customer_name.strip(),
            template_key=data.template_key, template_title=template["title"],
            recipient=(data.recipient or "").strip() or None, subject=data.subject.strip(),
            body=data.body, form_data=data.form_data, operator_uid=data.operator_uid or "local-user",
        )
        db.add(item); db.commit(); db.refresh(item)
        return serialize_windtre_panel_request(item)


@app.patch("/api/v1/windtre-panel/requests/{request_id}")
def update_windtre_panel_request(request_id: uuid.UUID, data: WindTrePanelResponseUpdate):
    with SessionLocal() as db:
        item = db.get(WindTrePanelRequest, request_id)
        if not item:
            raise HTTPException(404, "Pratica non trovata")
        item.response_date = data.response_date
        item.status = "GESTITA" if data.response_date else "ATTESA"
        db.commit(); db.refresh(item)
        return serialize_windtre_panel_request(item)


@app.delete("/api/v1/windtre-panel/requests/{request_id}")
def delete_windtre_panel_request(request_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(WindTrePanelRequest, request_id)
        if not item:
            raise HTTPException(404, "Pratica non trovata")
        db.delete(item); db.commit()
        return {"ok": True}


TARIFF_PLAN_TYPES = {"VOCE", "DATI", "FISSO", "DATI_M2M"}


def tariff_cents(value: Any) -> int:
    return max(0, round(parse_monthly_fee(value) * 100))


def serialize_tariff_plan(item: TariffPlan) -> dict[str, Any]:
    return {
        "id": str(item.id), "plan_code": item.plan_code, "name": item.name, "plan_type": item.plan_type,
        "ga_list_code": item.ga_list_code or "", "cb_list_code": item.cb_list_code or "",
        "valid_from": item.valid_from.isoformat() if item.valid_from else None,
        "valid_to": item.valid_to.isoformat() if item.valid_to else None, "subscribable": item.subscribable,
        "monthly_fee": item.monthly_fee_cents / 100, "secure_web": item.secure_web_cents / 100,
        "activation_cost": item.activation_cost_cents / 100, "sim_cost": item.sim_cost_cents / 100,
        "national_gb": item.national_gb or "", "national_minutes": item.national_minutes or "",
        "national_sms": item.national_sms or "", "eu_limited": item.eu_limited, "eu_gb": item.eu_gb or "",
        "eu_minutes": item.eu_minutes or "", "eu_sms": item.eu_sms or "",
        "eu_international_calls": item.eu_international_calls or "", "roaming_countries": item.roaming_countries or "",
        "international_included": item.international_included, "international_countries": item.international_countries or "",
        "custom_discounts": item.custom_discounts or [],
        "technical_pdf_url": f"/uploads/{item.technical_pdf_path}" if item.technical_pdf_path else None,
        "total_monthly": (item.monthly_fee_cents + item.secure_web_cents) / 100,
        "updated_at": item.updated_at.isoformat(),
    }


def apply_tariff_data(item: TariffPlan, data: TariffPlanRequest):
    if data.plan_type not in TARIFF_PLAN_TYPES:
        raise HTTPException(422, "Tipologia piano non valida")
    if not data.name.strip():
        raise HTTPException(422, "Il nome del piano è obbligatorio")
    if data.valid_from and data.valid_to and data.valid_to < data.valid_from:
        raise HTTPException(422, "La data di fine validità precede la data iniziale")
    text_fields = (
        "ga_list_code", "cb_list_code", "national_gb", "national_minutes", "national_sms", "eu_gb",
        "eu_minutes", "eu_sms", "eu_international_calls", "roaming_countries", "international_countries",
    )
    item.name, item.plan_type = data.name.strip(), data.plan_type
    item.valid_from, item.valid_to, item.subscribable = data.valid_from, data.valid_to, data.subscribable
    item.monthly_fee_cents, item.secure_web_cents = tariff_cents(data.monthly_fee), tariff_cents(data.secure_web)
    item.activation_cost_cents, item.sim_cost_cents = tariff_cents(data.activation_cost), tariff_cents(data.sim_cost)
    item.eu_limited, item.international_included = data.eu_limited, data.international_included
    for field in text_fields:
        value = getattr(data, field)
        setattr(item, field, value.strip() if value else None)
    item.custom_discounts = [
        {
            "cod": clean(row.get("cod")), "desc": clean(row.get("desc")),
            "prz": parse_monthly_fee(row.get("prz")), "sw": parse_monthly_fee(row.get("sw")),
            "ga": bool(row.get("ga")), "cb": bool(row.get("cb")),
        }
        for row in data.custom_discounts if clean(row.get("cod")) or clean(row.get("desc"))
    ]


def next_plan_code(db) -> str:
    base = int(datetime.now(timezone.utc).timestamp() * 1000)
    code = f"PT-{base}"
    while db.scalar(select(TariffPlan.id).where(TariffPlan.plan_code == code)):
        base += 1
        code = f"PT-{base}"
    return code


@app.get("/api/v1/tariff-plans")
def tariff_plans(search: str = "", plan_type: str | None = None, subscribable: bool | None = None):
    with SessionLocal() as db:
        query = select(TariffPlan).order_by(TariffPlan.subscribable.desc(), TariffPlan.name)
        if plan_type:
            query = query.where(TariffPlan.plan_type == plan_type)
        if subscribable is not None:
            query = query.where(TariffPlan.subscribable == subscribable)
        if search.strip():
            pattern = f"%{search.strip()}%"
            query = query.where(or_(TariffPlan.name.ilike(pattern), TariffPlan.ga_list_code.ilike(pattern), TariffPlan.cb_list_code.ilike(pattern), TariffPlan.plan_code.ilike(pattern)))
        return [serialize_tariff_plan(item) for item in db.scalars(query).all()]


@app.post("/api/v1/tariff-plans")
def create_tariff_plan(data: TariffPlanRequest):
    with SessionLocal() as db:
        if db.scalar(select(TariffPlan).where(func.lower(TariffPlan.name) == data.name.strip().lower())):
            raise HTTPException(409, "Esiste già un piano con questo nome")
        item = TariffPlan(plan_code=next_plan_code(db), name=data.name.strip(), plan_type=data.plan_type)
        apply_tariff_data(item, data)
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_tariff_plan(item)


@app.put("/api/v1/tariff-plans/{plan_id}")
def update_tariff_plan(plan_id: uuid.UUID, data: TariffPlanRequest):
    with SessionLocal() as db:
        item = db.get(TariffPlan, plan_id)
        if not item:
            raise HTTPException(404, "Piano non trovato")
        duplicate = db.scalar(select(TariffPlan).where(func.lower(TariffPlan.name) == data.name.strip().lower(), TariffPlan.id != plan_id))
        if duplicate:
            raise HTTPException(409, "Esiste già un piano con questo nome")
        apply_tariff_data(item, data)
        db.commit()
        db.refresh(item)
        return serialize_tariff_plan(item)


@app.delete("/api/v1/tariff-plans/{plan_id}")
def delete_tariff_plan(plan_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(TariffPlan, plan_id)
        if not item:
            raise HTTPException(404, "Piano non trovato")
        attachment = UPLOAD_DIR / item.technical_pdf_path if item.technical_pdf_path else None
        db.delete(item)
        db.commit()
        if attachment and attachment.is_file():
            attachment.unlink()
        return {"ok": True}


@app.post("/api/v1/tariff-plans/{plan_id}/pdf")
async def upload_tariff_pdf(plan_id: uuid.UUID, file: UploadFile = File(...)):
    if file.content_type != "application/pdf" and not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(422, "Carica una scheda tecnica PDF")
    contents = await file.read()
    if len(contents) > 15 * 1024 * 1024:
        raise HTTPException(413, "Il PDF supera il limite di 15 MB")
    filename = f"tariff-plan-{plan_id}.pdf"
    (UPLOAD_DIR / filename).write_bytes(contents)
    with SessionLocal() as db:
        item = db.get(TariffPlan, plan_id)
        if not item:
            raise HTTPException(404, "Piano non trovato")
        item.technical_pdf_path = filename
        db.commit()
        db.refresh(item)
        return serialize_tariff_plan(item)


@app.get("/api/v1/tariff-plans/{plan_id}/share")
def share_tariff_plan(plan_id: uuid.UUID, discount_code: str | None = None):
    with SessionLocal() as db:
        item = db.get(TariffPlan, plan_id)
        if not item:
            raise HTTPException(404, "Piano non trovato")
        scenario = next((row for row in item.custom_discounts or [] if row.get("cod") == discount_code), None)
        monthly = float(scenario.get("prz", 0)) if scenario else item.monthly_fee_cents / 100
        secure = float(scenario.get("sw", 0)) if scenario else item.secure_web_cents / 100
        text = "\n".join(filter(None, [
            f"📱 *{item.name}*", f"Tipologia: {item.plan_type.replace('_', ' ')}",
            f"💶 Canone: € {monthly:.2f}/mese", f"🛡️ Secure Web: € {secure:.2f}/mese" if secure else None,
            f"*Totale mensile: € {monthly + secure:.2f}*",
            f"🌐 Giga: {item.national_gb}" if item.national_gb else None,
            f"📞 Minuti: {item.national_minutes}" if item.national_minutes else None,
            f"💬 SMS: {item.national_sms}" if item.national_sms else None,
            f"🎁 Promozione: {scenario.get('desc')}" if scenario else None,
        ]))
        return {"subject": f"Proposta {item.name}", "text": text, "total_monthly": monthly + secure}


@app.get("/api/v1/tariff-plans-export.csv")
def export_tariff_plans():
    with SessionLocal() as db:
        items = db.scalars(select(TariffPlan).order_by(TariffPlan.name)).all()
        output = StringIO()
        fields = ["nome_piano","tipologia","listino_ga","listino_cb","data_inizio","data_fine","sottoscrivibile","canone_mensile","secure_web_base","costo_attivazione","costo_sim","gb_naz","min_naz","sms_naz","limitazioni_ue","gb_ue","min_ue","sms_ue","int_da_ue","paesi_roaming","flag_int","paesi_int","sconti_custom"]
        writer = csv.DictWriter(output, fieldnames=fields, delimiter=";")
        writer.writeheader()
        for item in items:
            value = serialize_tariff_plan(item)
            writer.writerow({
                "nome_piano": value["name"], "tipologia": value["plan_type"], "listino_ga": value["ga_list_code"], "listino_cb": value["cb_list_code"],
                "data_inizio": value["valid_from"] or "", "data_fine": value["valid_to"] or "", "sottoscrivibile": str(value["subscribable"]).lower(),
                "canone_mensile": value["monthly_fee"], "secure_web_base": value["secure_web"], "costo_attivazione": value["activation_cost"], "costo_sim": value["sim_cost"],
                "gb_naz": value["national_gb"], "min_naz": value["national_minutes"], "sms_naz": value["national_sms"], "limitazioni_ue": str(value["eu_limited"]).lower(),
                "gb_ue": value["eu_gb"], "min_ue": value["eu_minutes"], "sms_ue": value["eu_sms"], "int_da_ue": value["eu_international_calls"],
                "paesi_roaming": value["roaming_countries"], "flag_int": str(value["international_included"]).lower(), "paesi_int": value["international_countries"],
                "sconti_custom": json.dumps(value["custom_discounts"], ensure_ascii=False),
            })
        payload = ("\ufeff" + output.getvalue()).encode("utf-8")
        return StreamingResponse(BytesIO(payload), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="piani_tariffari.csv"'})


@app.post("/api/v1/tariff-plans-import")
async def import_tariff_plans(file: UploadFile = File(...)):
    rows = parse_uploaded_table(file.filename or "", await file.read())
    imported, skipped, errors = 0, 0, []
    with SessionLocal() as db:
        existing = {item.name.casefold() for item in db.scalars(select(TariffPlan)).all()}
        for index, row in enumerate(rows, 2):
            try:
                name = row.get("NOME_PIANO", "").strip()
                if not name or name.casefold() in existing:
                    skipped += 1
                    continue
                discounts = json.loads(row.get("SCONTI_CUSTOM", "[]") or "[]")
                data = TariffPlanRequest(
                    name=name, plan_type=row.get("TIPOLOGIA", "VOCE"), ga_list_code=row.get("LISTINO_GA"),
                    cb_list_code=row.get("LISTINO_CB"), valid_from=row.get("DATA_INIZIO") or None, valid_to=row.get("DATA_FINE") or None,
                    subscribable=normalize_header(row.get("SOTTOSCRIVIBILE", "TRUE")) in {"TRUE","SI","1","Y"},
                    monthly_fee=parse_monthly_fee(row.get("CANONE_MENSILE")), secure_web=parse_monthly_fee(row.get("SECURE_WEB_BASE")),
                    activation_cost=parse_monthly_fee(row.get("COSTO_ATTIVAZIONE")), sim_cost=parse_monthly_fee(row.get("COSTO_SIM")),
                    national_gb=row.get("GB_NAZ"), national_minutes=row.get("MIN_NAZ"), national_sms=row.get("SMS_NAZ"),
                    eu_limited=normalize_header(row.get("LIMITAZIONI_UE")) in {"TRUE","SI","1","Y"}, eu_gb=row.get("GB_UE"),
                    eu_minutes=row.get("MIN_UE"), eu_sms=row.get("SMS_UE"), eu_international_calls=row.get("INT_DA_UE"),
                    roaming_countries=row.get("PAESI_ROAMING"), international_included=normalize_header(row.get("FLAG_INT")) in {"TRUE","SI","1","Y"},
                    international_countries=row.get("PAESI_INT"), custom_discounts=discounts,
                )
                item = TariffPlan(plan_code=next_plan_code(db), name=name, plan_type=data.plan_type)
                apply_tariff_data(item, data)
                db.add(item)
                db.flush()
                existing.add(name.casefold())
                imported += 1
            except Exception as exc:
                errors.append({"row": index, "error": str(exc)})
        db.commit()
    return {"imported": imported, "skipped": skipped, "errors": errors[:100]}


DDT_STATUSES = {"IN_PREPARAZIONE", "SPEDITO", "CONSEGNATO", "ANNULLATO"}


def sender_snapshot(item: StoreSettings) -> dict[str, Any]:
    full_address = ", ".join(
        part for part in [
            item.address,
            " ".join(part for part in [item.postal_code, item.city] if part),
            f"({item.province})" if item.province else None,
        ] if part
    )
    return {
        "name": item.legal_name or item.store_name,
        "address": full_address,
        "tax_id": item.tax_id or item.fiscal_code or "",
        "logo_path": item.logo_path,
        "partner_logo_path": item.partner_logo_path,
    }


def generate_ddt_number(db, document_date: date) -> str:
    prefix = f"DDT-{document_date.strftime('%Y%m%d')}-"
    for _ in range(100):
        number = prefix + str(uuid.uuid4().int % 1000).zfill(3)
        if not db.scalar(select(DdtShipment.id).where(DdtShipment.ddt_number == number)):
            return number
    raise HTTPException(503, "Impossibile generare il numero DDT")


def serialize_ddt(item: DdtShipment) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "ddt_number": item.ddt_number,
        "document_date": item.document_date.isoformat(),
        "customer_id": str(item.customer_id) if item.customer_id else None,
        "recipient_name": item.recipient_name,
        "recipient_address": item.recipient_address,
        "goods_description": item.goods_description,
        "carrier": item.carrier or "",
        "tracking_number": item.tracking_number or "",
        "status": item.status,
        "sims": [serialize_inventory(sim) for sim in item.sims],
        "sim_ids": [str(sim.id) for sim in item.sims],
        "sender_snapshot": item.sender_snapshot,
        "created_at": item.created_at.isoformat(),
        "updated_at": item.updated_at.isoformat(),
    }


def validate_ddt(data: DdtShipmentRequest):
    if data.status not in DDT_STATUSES:
        raise HTTPException(422, "Stato DDT non valido")
    if not data.recipient_name.strip():
        raise HTTPException(422, "Il destinatario è obbligatorio")
    if not data.recipient_address.strip():
        raise HTTPException(422, "L'indirizzo di consegna è obbligatorio")
    if not data.goods_description.strip() and not data.sim_ids:
        raise HTTPException(422, "La descrizione dei beni è obbligatoria")
    if data.status == "SPEDITO" and not (data.tracking_number or "").strip():
        raise HTTPException(422, "Inserisci il tracking prima di segnare il DDT come spedito")


def resolve_ddt_sims(db, sim_ids: list[uuid.UUID], customer_id: uuid.UUID | None) -> list[SimInventory]:
    if not sim_ids:
        return []
    items = db.scalars(select(SimInventory).where(SimInventory.id.in_(sim_ids))).all()
    if len(items) != len(set(sim_ids)):
        raise HTTPException(422, "Una o più SIM selezionate non sono presenti in magazzino")
    if customer_id and any(item.customer_id != customer_id for item in items):
        raise HTTPException(422, "Puoi inserire soltanto SIM associate al cliente selezionato")
    return items


@app.get("/api/v1/ddt")
def ddt_shipments(search: str = "", status: str | None = None):
    with SessionLocal() as db:
        query = select(DdtShipment).order_by(DdtShipment.document_date.desc(), DdtShipment.created_at.desc())
        if status:
            query = query.where(DdtShipment.status == status)
        if search.strip():
            pattern = f"%{search.strip()}%"
            query = query.where(or_(
                DdtShipment.ddt_number.ilike(pattern),
                DdtShipment.recipient_name.ilike(pattern),
                DdtShipment.carrier.ilike(pattern),
                DdtShipment.tracking_number.ilike(pattern),
                DdtShipment.goods_description.ilike(pattern),
            ))
        items = db.scalars(query).all()
        return {
            "items": [serialize_ddt(item) for item in items],
            "counts": {state: sum(item.status == state for item in items) for state in DDT_STATUSES},
        }


@app.post("/api/v1/ddt")
def create_ddt(data: DdtShipmentRequest):
    validate_ddt(data)
    with SessionLocal() as db:
        store = get_or_create_store_settings(db)
        item = DdtShipment(
            ddt_number=generate_ddt_number(db, data.document_date),
            document_date=data.document_date,
            customer_id=data.customer_id,
            recipient_name=data.recipient_name.strip(),
            recipient_address=data.recipient_address.strip(),
            goods_description=data.goods_description.strip(),
            carrier=(data.carrier or "").strip() or None,
            tracking_number=(data.tracking_number or "").strip() or None,
            status=data.status,
            sender_snapshot=sender_snapshot(store),
            sims=resolve_ddt_sims(db, data.sim_ids, data.customer_id),
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_ddt(item)


@app.put("/api/v1/ddt/{ddt_id}")
def update_ddt(ddt_id: uuid.UUID, data: DdtShipmentRequest):
    validate_ddt(data)
    with SessionLocal() as db:
        item = db.get(DdtShipment, ddt_id)
        if not item:
            raise HTTPException(404, "DDT non trovato")
        if item.status == "ANNULLATO":
            raise HTTPException(409, "Un DDT annullato non può essere modificato")
        sim_ids = data.sim_ids
        for field, value in data.model_dump(exclude={"sim_ids"}).items():
            setattr(item, field, value.strip() if isinstance(value, str) else value)
        item.sims = resolve_ddt_sims(db, sim_ids, data.customer_id)
        db.commit()
        db.refresh(item)
        return serialize_ddt(item)


@app.patch("/api/v1/ddt/{ddt_id}/status")
def update_ddt_status(ddt_id: uuid.UUID, data: DdtStatusRequest):
    if data.status not in DDT_STATUSES:
        raise HTTPException(422, "Stato DDT non valido")
    with SessionLocal() as db:
        item = db.get(DdtShipment, ddt_id)
        if not item:
            raise HTTPException(404, "DDT non trovato")
        if item.status == "ANNULLATO":
            raise HTTPException(409, "Un DDT annullato non può cambiare stato")
        if data.status == "SPEDITO" and not (data.tracking_number or item.tracking_number):
            raise HTTPException(422, "Inserisci il tracking prima della spedizione")
        item.status = data.status
        if data.tracking_number is not None:
            item.tracking_number = data.tracking_number.strip() or None
        db.commit()
        db.refresh(item)
        return serialize_ddt(item)


@app.delete("/api/v1/ddt/{ddt_id}")
def delete_ddt(ddt_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(DdtShipment, ddt_id)
        if not item:
            raise HTTPException(404, "DDT non trovato")
        if item.status != "IN_PREPARAZIONE":
            raise HTTPException(409, "Puoi eliminare solo DDT in preparazione; usa Annulla per conservarne lo storico")
        db.delete(item)
        db.commit()
        return {"ok": True}


@app.get("/api/v1/ddt/{ddt_id}/pdf")
def ddt_pdf(ddt_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(DdtShipment, ddt_id)
        if not item:
            raise HTTPException(404, "DDT non trovato")
        snapshot = item.sender_snapshot or {}
        lines = [line.strip() for line in item.goods_description.splitlines() if line.strip()]
        if item.sims:
            lines.append("SIM SELEZIONATE:")
            lines.extend(
                f"{sim.product.name} | ICCID: {sim.iccid} | Numero: {sim.msisdn or '—'} | Stato: {sim.status.replace('_', ' ')}"
                for sim in item.sims
            )
        lines = lines or ["—"]
        chunks = [lines[index:index + 15] for index in range(0, len(lines), 15)]
        styles = getSampleStyleSheet()
        body = ParagraphStyle("ddt-body", parent=styles["BodyText"], fontSize=9, leading=12)
        small = ParagraphStyle("ddt-small", parent=body, fontSize=8, leading=10)
        heading = ParagraphStyle("ddt-heading", parent=styles["Heading1"], fontSize=13, leading=16, alignment=1)
        buffer = BytesIO()
        document = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
        story = []
        for page_index, chunk in enumerate(chunks):
            logo_path = UPLOAD_DIR / snapshot.get("logo_path", "") if snapshot.get("logo_path") else None
            partner_path = UPLOAD_DIR / snapshot.get("partner_logo_path", "") if snapshot.get("partner_logo_path") else None
            left_logo = Image(str(logo_path), width=35 * mm, height=15 * mm, kind="proportional") if logo_path and logo_path.is_file() else Paragraph("", body)
            right_logo = Image(str(partner_path), width=35 * mm, height=15 * mm, kind="proportional") if partner_path and partner_path.is_file() else Paragraph("", body)
            header = Table([[left_logo, Paragraph("DOCUMENTO DI TRASPORTO (D.D.T.)<br/><font size='8'>D.P.R. 472/96</font>", heading), right_logo]], colWidths=[45 * mm, 87 * mm, 45 * mm])
            header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (2, 0), (2, 0), "RIGHT")]))
            story.extend([header, Spacer(1, 4 * mm)])
            refs = Table([[
                Paragraph(f"<b>Numero:</b> {item.ddt_number}", body),
                Paragraph(f"<b>Data:</b> {item.document_date.strftime('%d/%m/%Y')}", body),
                Paragraph(f"<b>Pagina:</b> {page_index + 1} di {len(chunks)}", body),
            ]], colWidths=[75 * mm, 55 * mm, 47 * mm])
            refs.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .6, colors.HexColor("#9ca3af")), ("INNERGRID", (0, 0), (-1, -1), .4, colors.HexColor("#d1d5db")), ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")), ("PADDING", (0, 0), (-1, -1), 6)]))
            story.extend([refs, Spacer(1, 4 * mm)])
            parties = Table([[
                Paragraph(f"<b>MITTENTE</b><br/>{snapshot.get('name', '')}<br/>{snapshot.get('address', '')}<br/>P.IVA/C.F.: {snapshot.get('tax_id', '')}", body),
                Paragraph(f"<b>DESTINATARIO</b><br/>{item.recipient_name}<br/>{item.recipient_address}", body),
            ]], colWidths=[88.5 * mm, 88.5 * mm])
            parties.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .7, colors.HexColor("#475569")), ("INNERGRID", (0, 0), (-1, -1), .4, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 8)]))
            story.extend([parties, Spacer(1, 5 * mm)])
            goods = [[Paragraph("<b>DESCRIZIONE DEI BENI / SERIALI / MATRICOLЕ</b>", body)]]
            goods.extend([[Paragraph(line, body)] for line in chunk])
            if page_index < len(chunks) - 1:
                goods.append([Paragraph("<i>Segue a pagina successiva…</i>", small)])
            goods_table = Table(goods, colWidths=[177 * mm], repeatRows=1)
            goods_table.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .7, colors.HexColor("#475569")), ("INNERGRID", (0, 0), (-1, -1), .3, colors.HexColor("#cbd5e1")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")), ("PADDING", (0, 0), (-1, -1), 6)]))
            story.append(goods_table)
            if page_index == len(chunks) - 1:
                story.extend([Spacer(1, 7 * mm), Table([[
                    Paragraph(f"<b>VETTORE</b><br/>{item.carrier or '—'}", body),
                    Paragraph(f"<b>TRACKING</b><br/>{item.tracking_number or '—'}", body),
                ]], colWidths=[88.5 * mm, 88.5 * mm], style=[("BOX", (0, 0), (-1, -1), .7, colors.HexColor("#475569")), ("INNERGRID", (0, 0), (-1, -1), .4, colors.HexColor("#cbd5e1")), ("PADDING", (0, 0), (-1, -1), 8)]), Spacer(1, 14 * mm)])
                signatures = Table([["Firma Mittente", "Firma Vettore", "Firma Destinatario"], ["\n\n____________________", "\n\n____________________", "\n\n____________________"]], colWidths=[59 * mm] * 3)
                signatures.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
                story.append(signatures)
            if page_index < len(chunks) - 1:
                story.append(PageBreak())
        document.build(story)
        buffer.seek(0)
        headers = {"Content-Disposition": f'inline; filename="{item.ddt_number}.pdf"'}
        return StreamingResponse(buffer, media_type="application/pdf", headers=headers)


@app.get("/api/v1/products")
def products(search: str = "", include_inactive: bool = False):
    with SessionLocal() as db:
        query = select(Product).order_by(Product.name)
        if not include_inactive:
            query = query.where(Product.is_active.is_(True))
        if search.strip():
            pattern = f"%{search.strip()}%"
            query = query.where(or_(Product.sku.ilike(pattern), Product.name.ilike(pattern)))
        items = db.scalars(query).all()
        costs = [item.unit_cost_cents for item in items if item.unit_cost_cents is not None]
        return {
            "items": [serialize_product(item) for item in items],
            "total": len(items),
            "average_cost": round(sum(costs) / len(costs) / 100, 2) if costs else 0,
        }


@app.post("/api/v1/products")
def create_product(data: ProductRequest):
    sku, name = data.sku.strip().upper(), data.name.strip()
    if not sku or not name:
        raise HTTPException(422, "Codice articolo e nome sono obbligatori")
    with SessionLocal() as db:
        if db.scalar(select(Product).where(Product.sku == sku)):
            raise HTTPException(409, "Codice articolo già presente")
        item = Product(sku=sku, name=name, unit_cost_cents=cents_from_value(data.unit_cost))
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_product(item)


@app.put("/api/v1/products/{product_id}")
def update_product(product_id: uuid.UUID, data: ProductRequest):
    with SessionLocal() as db:
        item = db.get(Product, product_id)
        if not item:
            raise HTTPException(404, "Prodotto non trovato")
        sku, name = data.sku.strip().upper(), data.name.strip()
        duplicate = db.scalar(select(Product).where(Product.sku == sku, Product.id != product_id))
        if duplicate:
            raise HTTPException(409, "Codice articolo già presente")
        item.sku, item.name = sku, name
        item.unit_cost_cents = cents_from_value(data.unit_cost)
        item.is_active = True
        db.commit()
        return serialize_product(item)


@app.delete("/api/v1/products/{product_id}")
def delete_product(product_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(Product, product_id)
        if not item:
            raise HTTPException(404, "Prodotto non trovato")
        item.is_active = False
        db.commit()
        return {"ok": True}


@app.post("/api/v1/products/import")
async def import_products(file: UploadFile = File(...)):
    contents = await file.read()
    rows = parse_uploaded_table(file.filename or "", contents)
    added = updated = rejected = 0
    errors = []
    with SessionLocal() as db:
        for index, row in enumerate(rows, 2):
            sku = mapped_value(row, PRODUCT_COLUMN_ALIASES["sku"]).strip().upper()
            name = mapped_value(row, PRODUCT_COLUMN_ALIASES["name"]).strip()
            cost = mapped_value(row, PRODUCT_COLUMN_ALIASES["unit_cost"])
            if not sku or not name:
                rejected += 1
                errors.append({"row": index, "error": "Codice o nome mancante"})
                continue
            item = db.scalar(select(Product).where(Product.sku == sku))
            if item:
                item.name = name
                item.unit_cost_cents = cents_from_value(cost)
                item.is_active = True
                updated += 1
            else:
                db.add(Product(sku=sku, name=name, unit_cost_cents=cents_from_value(cost)))
                added += 1
        db.commit()
    return {"added": added, "updated": updated, "rejected": rejected, "errors": errors[:50]}


@app.get("/api/v1/sim-orders")
def sim_orders():
    with SessionLocal() as db:
        return [serialize_order(item) for item in db.scalars(select(SimOrder).order_by(SimOrder.order_date.desc())).unique().all()]


@app.post("/api/v1/sim-orders")
def create_sim_order(data: SimOrderRequest):
    if not data.lines:
        raise HTTPException(422, "Inserisci almeno una riga d'ordine")
    if data.status not in ORDER_STATUSES:
        raise HTTPException(422, "Stato ordine non valido")
    with SessionLocal() as db:
        number = (data.order_number or "").strip().upper()
        if not number:
            year = data.order_date.year
            sequence = (db.scalar(select(func.count()).select_from(SimOrder).where(
                func.extract("year", SimOrder.order_date) == year
            )) or 0) + 1
            number = f"ORD-{year}-{sequence:03d}"
        if db.scalar(select(SimOrder).where(SimOrder.order_number == number)):
            raise HTTPException(409, "Numero ordine già presente")
        products_by_id = {
            item.id: item
            for item in db.scalars(select(Product).where(Product.id.in_([line.product_id for line in data.lines]))).all()
        }
        if len(products_by_id) != len(set(line.product_id for line in data.lines)):
            raise HTTPException(422, "Uno o più prodotti non esistono")
        if any(line.quantity <= 0 for line in data.lines):
            raise HTTPException(422, "Le quantità devono essere maggiori di zero")
        item = SimOrder(
            order_number=number,
            order_date=data.order_date,
            supplier=(data.supplier or "").strip() or None,
            notes=(data.notes or "").strip() or None,
            status=data.status,
        )
        db.add(item)
        db.flush()
        for line in data.lines:
            db.add(SimOrderLine(
                order_id=item.id,
                product_id=line.product_id,
                quantity=line.quantity,
                description=(line.description or "").strip() or None,
            ))
        db.commit()
        db.refresh(item)
        return serialize_order(item)


@app.patch("/api/v1/sim-orders/{order_id}/status")
def update_sim_order_status(order_id: uuid.UUID, data: OrderStatusRequest):
    if data.status not in ORDER_STATUSES:
        raise HTTPException(422, "Stato ordine non valido")
    with SessionLocal() as db:
        item = db.get(SimOrder, order_id)
        if not item:
            raise HTTPException(404, "Ordine non trovato")
        item.status = data.status
        db.commit()
        return serialize_order(item)


@app.get("/api/v1/sim-orders/{order_id}/share")
def share_sim_order(order_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(SimOrder, order_id)
        if not item:
            raise HTTPException(404, "Ordine non trovato")
        return {"subject": f"Ordine SIM {item.order_number}", "text": order_share_text(item)}


@app.get("/api/v1/sim-inventory")
def sim_inventory(
    search: str = "",
    status: str | None = None,
    product_id: uuid.UUID | None = None,
    order_id: uuid.UUID | None = None,
    customer_id: uuid.UUID | None = None,
    limit: int = Query(500, ge=1, le=2000),
):
    with SessionLocal() as db:
        query = select(SimInventory).order_by(SimInventory.created_at.desc()).limit(limit)
        if status:
            query = query.where(SimInventory.status == status)
        if product_id:
            query = query.where(SimInventory.product_id == product_id)
        if order_id:
            query = query.where(SimInventory.order_id == order_id)
        if customer_id:
            query = query.where(SimInventory.customer_id == customer_id)
        if search.strip():
            pattern = f"%{search.strip()}%"
            query = query.outerjoin(Customer).join(Product).outerjoin(SimOrder).where(or_(
                SimInventory.iccid.ilike(pattern),
                SimInventory.msisdn.ilike(pattern),
                Customer.business_name.ilike(pattern),
                Product.sku.ilike(pattern),
                Product.name.ilike(pattern),
                SimOrder.order_number.ilike(pattern),
            ))
        items = db.scalars(query).unique().all()
        return [serialize_inventory(item) for item in items]


@app.post("/api/v1/sim-inventory")
def create_sim_inventory(data: SimInventoryRequest):
    iccid = normalize_iccid(data.iccid)
    if len(iccid) not in {19, 20}:
        raise HTTPException(422, "L'ICCID deve contenere 19 o 20 cifre")
    if data.status not in SIM_STATUSES:
        raise HTTPException(422, "Stato SIM non valido")
    with SessionLocal() as db:
        if db.scalar(select(SimInventory).where(SimInventory.iccid == iccid)):
            raise HTTPException(409, "ICCID già presente")
        if not db.get(Product, data.product_id):
            raise HTTPException(422, "Prodotto non trovato")
        item = SimInventory(
            iccid=iccid,
            product_id=data.product_id,
            order_id=data.order_id,
            status=data.status,
            customer_id=data.customer_id,
            msisdn=normalize_iccid(data.msisdn) or None,
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_inventory(item)


@app.patch("/api/v1/sim-inventory/{inventory_id}")
def update_sim_inventory(inventory_id: uuid.UUID, data: SimInventoryUpdate):
    with SessionLocal() as db:
        item = db.get(SimInventory, inventory_id)
        if not item:
            raise HTTPException(404, "SIM non trovata")
        if data.status is not None:
            if data.status not in SIM_STATUSES:
                raise HTTPException(422, "Stato SIM non valido")
            item.status = data.status
        item.customer_id = data.customer_id
        item.msisdn = normalize_iccid(data.msisdn) or None
        db.commit()
        return serialize_inventory(item)


@app.post("/api/v1/sim-inventory/import")
async def import_sim_inventory(
    file: UploadFile = File(...),
    product_sku: str | None = Form(None),
    order_number: str | None = Form(None),
    default_status: str = Form("IN_MAGAZZINO"),
):
    if default_status not in SIM_STATUSES:
        raise HTTPException(422, "Stato SIM non valido")
    contents = await file.read()
    rows = parse_uploaded_table(file.filename or "", contents)
    added = updated = rejected = 0
    errors = []
    with SessionLocal() as db:
        products = {item.sku.upper(): item for item in db.scalars(select(Product)).all()}
        customers = {item.business_name.strip().casefold(): item for item in db.scalars(select(Customer)).all()}
        orders = {item.order_number.upper(): item for item in db.scalars(select(SimOrder)).unique().all()}
        created_products = created_customers = created_orders = 0

        # Il formato esportato da Cornet contiene già articolo, ordine e cliente:
        # creiamo prima le anagrafiche mancanti, poi importiamo le singole SIM.
        for row in rows:
            sku = (mapped_value(row, SIM_COLUMN_ALIASES["sku"]) or product_sku or "").strip().upper()
            if sku and sku not in products:
                product = Product(sku=sku, name=sku, is_active=True)
                db.add(product)
                db.flush()
                products[sku] = product
                created_products += 1
            customer_name = mapped_value(row, SIM_COLUMN_ALIASES["customer"]).strip()
            customer_key = customer_name.casefold()
            if customer_name and customer_key not in customers:
                customer = Customer(
                    business_name=customer_name,
                    segment="MICROBUSINESS",
                    portfolio_status="ACTIVE",
                )
                db.add(customer)
                db.flush()
                customers[customer_key] = customer
                created_customers += 1

        order_lines: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
        order_dates: dict[str, date] = {}
        for row in rows:
            source_order = (mapped_value(row, SIM_COLUMN_ALIASES["order_number"]) or order_number or "").strip().upper()
            sku = (mapped_value(row, SIM_COLUMN_ALIASES["sku"]) or product_sku or "").strip().upper()
            if source_order and sku:
                order_lines[source_order][sku] += 1
                order_dates.setdefault(source_order, parse_compact_date(mapped_value(row, SIM_COLUMN_ALIASES["order_date"])))
        for source_order, lines in order_lines.items():
            if source_order in orders:
                continue
            order = SimOrder(
                order_number=source_order,
                order_date=order_dates[source_order],
                supplier="WINDTRE",
                status="RICEVUTO",
            )
            db.add(order)
            db.flush()
            for sku, quantity in lines.items():
                db.add(SimOrderLine(order_id=order.id, product_id=products[sku].id, quantity=quantity))
            orders[source_order] = order
            created_orders += 1

        for index, row in enumerate(rows, 2):
            iccid = normalize_iccid(mapped_value(row, SIM_COLUMN_ALIASES["iccid"]))
            sku = (mapped_value(row, SIM_COLUMN_ALIASES["sku"]) or product_sku or "").strip().upper()
            status = normalize_header(mapped_value(row, SIM_COLUMN_ALIASES["status"]) or default_status)
            source_order = (mapped_value(row, SIM_COLUMN_ALIASES["order_number"]) or order_number or "").strip().upper()
            msisdn = normalize_iccid(mapped_value(row, SIM_COLUMN_ALIASES["msisdn"])) or None
            customer_name = mapped_value(row, SIM_COLUMN_ALIASES["customer"]).strip()
            customer = customers.get(customer_name.casefold()) if customer_name else None
            if customer or msisdn:
                status = "ASSEGNATA"
            product = products.get(sku)
            order = orders.get(source_order) if source_order else None
            if len(iccid) not in {19, 20} or not product or status not in SIM_STATUSES:
                rejected += 1
                errors.append({"row": index, "iccid": iccid, "error": "ICCID, prodotto o stato non valido"})
                continue
            item = db.scalar(select(SimInventory).where(SimInventory.iccid == iccid))
            if item:
                item.product_id = product.id
                item.order_id = order.id if order else None
                item.status = status
                item.msisdn = msisdn
                item.customer_id = customer.id if customer else None
                updated += 1
            else:
                db.add(SimInventory(
                    iccid=iccid,
                    product_id=product.id,
                    order_id=order.id if order else None,
                    status=status,
                    msisdn=msisdn,
                    customer_id=customer.id if customer else None,
                ))
                added += 1
        record = SimInventoryImport(
            file_name=file.filename or "import",
            added_count=added,
            updated_count=updated,
            rejected_count=rejected,
            errors=errors[:100],
        )
        db.add(record)
        db.commit()
        return {
            "id": str(record.id),
            "added": added,
            "updated": updated,
            "rejected": rejected,
            "created_products": created_products,
            "created_orders": created_orders,
            "created_customers": created_customers,
            "errors": errors[:50],
        }


@app.get("/api/v1/sim-inventory/imports")
def sim_inventory_import_history():
    with SessionLocal() as db:
        items = db.scalars(select(SimInventoryImport).order_by(SimInventoryImport.imported_at.desc()).limit(50)).all()
        return [{
            "id": str(item.id),
            "file_name": item.file_name,
            "added": item.added_count,
            "updated": item.updated_count,
            "rejected": item.rejected_count,
            "errors": item.errors,
            "imported_at": item.imported_at.isoformat(),
        } for item in items]


@app.get("/api/v1/sim-report")
def sim_report():
    with SessionLocal() as db:
        inventory = db.scalars(select(SimInventory).order_by(SimInventory.created_at)).all()
        orders = db.scalars(select(SimOrder).order_by(SimOrder.order_date.desc())).unique().all()
        total = len(inventory)
        available = sum(item.status == "IN_MAGAZZINO" for item in inventory)
        grouped: dict[str, dict[str, Any]] = {}
        for order in orders:
            group = grouped[str(order.id)] = {
                "order_id": str(order.id),
                "order_number": order.order_number,
                "order_date": order.order_date.isoformat(),
                "status": order.status,
                "products": {},
            }
            for line in order.lines:
                group["products"][str(line.product_id)] = {
                    "product_id": str(line.product_id),
                    "sku": line.product.sku,
                    "product_name": line.product.name,
                    "ordered": line.quantity,
                    "total": 0,
                    "available": 0,
                }
        for item in inventory:
            key = str(item.order_id) if item.order_id else "NO_ORDER"
            group = grouped.setdefault(key, {
                "order_id": None,
                "order_number": "SENZA ORDINE",
                "order_date": None,
                "status": None,
                "products": {},
            })
            product = group["products"].setdefault(str(item.product_id), {
                "product_id": str(item.product_id),
                "sku": item.product.sku,
                "product_name": item.product.name,
                "ordered": 0,
                "total": 0,
                "available": 0,
            })
            product["total"] += 1
            product["available"] += item.status == "IN_MAGAZZINO"
        return {
            "total": total,
            "available": available,
            "assigned": total - available,
            "assignment_rate": round((total - available) / total * 100, 1) if total else 0,
            "orders": [
                {**group, "products": list(group["products"].values())}
                for group in grouped.values()
            ],
        }


@app.post("/api/v1/customers")
def quick_create_customer(data: QuickCustomerRequest):
    name = data.business_name.strip()
    if not name:
        raise HTTPException(422, "Nome cliente obbligatorio")
    if data.segment not in {"CONSUMER", "MICROBUSINESS", "BUSINESS_SME"}:
        raise HTTPException(422, "Segmento cliente non valido")
    with SessionLocal() as db:
        tax_id = (data.tax_id or "").strip() or None
        fiscal_code = (data.fiscal_code or "").strip() or None
        if tax_id and db.scalar(select(Customer).where(Customer.tax_id == tax_id)):
            raise HTTPException(409, "Cliente già presente")
        if fiscal_code and db.scalar(select(Customer).where(Customer.fiscal_code == fiscal_code)):
            raise HTTPException(409, "Cliente già presente")
        item = Customer(
            business_name=name,
            tax_id=tax_id,
            fiscal_code=fiscal_code,
            address=(data.address or "").strip() or None,
            segment=data.segment,
            portfolio_status="ACTIVE",
        )
        for field in (
            "first_name", "last_name", "postal_code", "city", "province", "email", "phone", "birth_date",
            "birth_place", "birth_province", "gender", "document_type", "document_number",
            "document_issue_date", "document_expiry_date", "document_issuer",
        ):
            value = getattr(data, field)
            setattr(item, field, value.strip() or None if isinstance(value, str) else value)
        db.add(item)
        db.commit()
        return {"id": str(item.id), "business_name": item.business_name}


@app.put("/api/v1/customers/{customer_id}")
def update_customer(customer_id: uuid.UUID, data: QuickCustomerRequest):
    with SessionLocal() as db:
        item = db.get(Customer, customer_id)
        if not item:
            raise HTTPException(404, "Cliente non trovato")
        if data.segment not in {"CONSUMER", "MICROBUSINESS", "BUSINESS_SME"}:
            raise HTTPException(422, "Segmento cliente non valido")
        item.business_name = data.business_name.strip()
        item.segment = data.segment
        for field in (
            "first_name", "last_name", "tax_id", "fiscal_code", "address", "postal_code", "city", "province",
            "email", "phone", "birth_date", "birth_place", "birth_province", "gender", "document_type",
            "document_number", "document_issue_date", "document_expiry_date", "document_issuer",
        ):
            value = getattr(data, field)
            setattr(item, field, value.strip() or None if isinstance(value, str) else value)
        db.commit()
        return {"id": str(item.id), "business_name": item.business_name}


@app.get("/api/v1/customers/{customer_id}/account-codes")
def customer_account_codes(customer_id: uuid.UUID):
    with SessionLocal() as db:
        if not db.get(Customer, customer_id):
            raise HTTPException(404, "Cliente non trovato")
        items = db.scalars(
            select(CustomerAccountCode)
            .where(CustomerAccountCode.customer_id == customer_id)
            .order_by(CustomerAccountCode.operator, CustomerAccountCode.market, CustomerAccountCode.customer_code)
        ).all()
        return [serialize_account_code(item) for item in items]


@app.post("/api/v1/customers/{customer_id}/account-codes")
def create_customer_account_code(customer_id: uuid.UUID, data: CustomerAccountCodeRequest):
    operator = normalize_header(data.operator)
    market = normalize_header(data.market)
    customer_code = clean(data.customer_code).upper()
    if not operator or not customer_code:
        raise HTTPException(422, "Operatore e codice cliente sono obbligatori")
    if market not in {"CONSUMER", "MICROBUSINESS", "BUSINESS_SME"}:
        raise HTTPException(422, "Mercato non valido")
    with SessionLocal() as db:
        customer = db.get(Customer, customer_id)
        if not customer:
            raise HTTPException(404, "Cliente non trovato")
        duplicate = db.scalar(
            select(CustomerAccountCode).where(
                CustomerAccountCode.operator == operator,
                CustomerAccountCode.market == market,
                CustomerAccountCode.customer_code == customer_code,
            )
        )
        if duplicate:
            if duplicate.customer_id != customer_id:
                raise HTTPException(409, "Codice già associato a un altro cliente nello stesso mercato")
            return serialize_account_code(duplicate)
        if data.is_primary:
            for item in db.scalars(
                select(CustomerAccountCode).where(
                    CustomerAccountCode.customer_id == customer_id,
                    CustomerAccountCode.operator == operator,
                    CustomerAccountCode.market == market,
                )
            ).all():
                item.is_primary = False
        item = CustomerAccountCode(
            customer_id=customer_id,
            operator=operator,
            market=market,
            customer_code=customer_code,
            is_primary=data.is_primary,
        )
        db.add(item)
        if operator == "WINDTRE" and market == "BUSINESS_SME":
            customer.windtre_customer_code = customer_code
        db.commit()
        db.refresh(item)
        return serialize_account_code(item)


@app.delete("/api/v1/customers/{customer_id}/account-codes/{account_code_id}")
def delete_customer_account_code(customer_id: uuid.UUID, account_code_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(CustomerAccountCode, account_code_id)
        if not item or item.customer_id != customer_id:
            raise HTTPException(404, "Codice cliente non trovato")
        customer = db.get(Customer, customer_id)
        clears_legacy = (
            item.operator == "WINDTRE"
            and item.market == "BUSINESS_SME"
            and customer
            and customer.windtre_customer_code == item.customer_code
        )
        db.delete(item)
        if clears_legacy:
            replacement = db.scalar(
                select(CustomerAccountCode).where(
                    CustomerAccountCode.customer_id == customer_id,
                    CustomerAccountCode.operator == "WINDTRE",
                    CustomerAccountCode.market == "BUSINESS_SME",
                    CustomerAccountCode.id != account_code_id,
                ).order_by(CustomerAccountCode.is_primary.desc(), CustomerAccountCode.created_at)
            )
            customer.windtre_customer_code = replacement.customer_code if replacement else None
        db.commit()
        return {"deleted": True}


@app.get("/api/v1/incentives")
def incentive_competitions():
    with SessionLocal() as db:
        items = db.scalars(
            select(IncentiveCompetition).order_by(IncentiveCompetition.start_date.desc(), IncentiveCompetition.created_at.desc())
        ).all()
        return [serialize_incentive_competition(item) for item in items]


@app.post("/api/v1/incentives/templates/windtre-march-2026")
def create_march_2026_incentive_template():
    with SessionLocal() as db:
        existing = db.scalar(
            select(IncentiveCompetition).where(
                IncentiveCompetition.operator == "WINDTRE",
                IncentiveCompetition.start_date == date(2026, 3, 1),
                IncentiveCompetition.end_date == date(2026, 3, 31),
            )
        )
        if existing:
            return serialize_incentive_competition(existing)
        store = db.scalar(select(StoreSettings).limit(1))
        item = IncentiveCompetition(
            name="Gara WINDTRE Marzo 2026",
            operator="WINDTRE",
            market="CONSUMER_MICROBUSINESS",
            dealer_code=store.dealer_code if store else None,
            start_date=date(2026, 3, 1),
            end_date=date(2026, 3, 31),
            status="ACTIVE",
            source_document="Gara_Marzo_2026.pdf",
            configuration=march_2026_incentive_configuration(),
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_incentive_competition(item)


@app.post("/api/v1/incentives")
def create_incentive_competition(data: IncentiveCompetitionRequest):
    if data.end_date < data.start_date:
        raise HTTPException(422, "La data finale deve essere successiva a quella iniziale")
    with SessionLocal() as db:
        item = IncentiveCompetition(
            name=data.name.strip(), operator=normalize_header(data.operator), market=normalize_header(data.market),
            dealer_code=clean(data.dealer_code) or None, start_date=data.start_date, end_date=data.end_date,
            status=normalize_header(data.status), source_document=clean(data.source_document) or None,
            configuration=data.configuration,
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_incentive_competition(item)


@app.put("/api/v1/incentives/{competition_id}")
def update_incentive_competition(competition_id: uuid.UUID, data: IncentiveCompetitionRequest):
    if data.end_date < data.start_date:
        raise HTTPException(422, "La data finale deve essere successiva a quella iniziale")
    with SessionLocal() as db:
        item = db.get(IncentiveCompetition, competition_id)
        if not item:
            raise HTTPException(404, "Gara non trovata")
        item.name = data.name.strip()
        item.operator = normalize_header(data.operator)
        item.market = normalize_header(data.market)
        item.dealer_code = clean(data.dealer_code) or None
        item.start_date = data.start_date
        item.end_date = data.end_date
        item.status = normalize_header(data.status)
        item.source_document = clean(data.source_document) or None
        item.configuration = data.configuration
        db.commit()
        return serialize_incentive_competition(item)


@app.post("/api/v1/incentives/{competition_id}/activations")
def create_incentive_activation(competition_id: uuid.UUID, data: IncentiveActivationRequest):
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        if not competition.start_date <= data.activation_date <= competition.end_date:
            raise HTTPException(422, "La data di attivazione non rientra nel periodo di gara")
        track = normalize_header(data.track)
        if track not in (competition.configuration or {}).get("tracks", {}):
            raise HTTPException(422, "Pista non configurata")
        item = IncentiveActivation(
            competition_id=competition_id, customer_id=data.customer_id, activation_date=data.activation_date,
            source_type="MANUAL", source_key=f"MANUAL:{uuid.uuid4()}", seller_name=clean(data.seller_name) or None,
            track=track, offer=clean(data.offer) or None, asset_number=clean(data.asset_number) or None,
            customer_code=clean(data.customer_code) or None, contract_code=clean(data.contract_code) or None,
            monthly_fee_cents=round(data.monthly_fee * 100), direct_bonus_cents=round(data.direct_bonus * 100),
            attributes=data.attributes, status=normalize_header(data.status), notes=clean(data.notes) or None,
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        return serialize_incentive_activation(item)


@app.post("/api/v1/incentives/{competition_id}/pdc/preview")
async def preview_incentive_pdc(competition_id: uuid.UUID, file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(422, "Carica una PDC in formato PDF")
    contents = await file.read()
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        parsed = parse_windtre_pdc(contents, file.filename or "PDC.pdf")
        fiscal_code = parsed["customer"]["fiscal_code"]
        customer = db.scalar(select(Customer).where(Customer.fiscal_code == fiscal_code)) if fiscal_code else None
        parsed["customer_match"] = {
            "found": bool(customer), "customer_id": str(customer.id) if customer else None,
            "customer_name": customer.business_name if customer else None,
        }
        parsed["file_sha256"] = sha256(contents).hexdigest()
        parsed["duplicate"] = bool(db.scalar(
            select(IncentivePdcImport).where(
                IncentivePdcImport.competition_id == competition_id,
                IncentivePdcImport.file_sha256 == parsed["file_sha256"],
            )
        ))
        if parsed["contract"]["dealer_code"] and competition.dealer_code and parsed["contract"]["dealer_code"] != competition.dealer_code:
            parsed["warnings"].append(
                f"Codice dealer PDC {parsed['contract']['dealer_code']} diverso dalla gara {competition.dealer_code}"
            )
        return parsed


@app.post("/api/v1/incentives/{competition_id}/pdc/import")
async def import_incentive_pdc(
    competition_id: uuid.UUID,
    file: UploadFile = File(...),
    seller_name: str | None = Form(None),
    include_mobile: bool = Form(False),
):
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(422, "Carica una PDC in formato PDF")
    contents = await file.read()
    digest = sha256(contents).hexdigest()
    parsed = parse_windtre_pdc(contents, file.filename or "PDC.pdf")
    activation_date = date.fromisoformat(parsed["contract"]["activation_date"]) if parsed["contract"]["activation_date"] else None
    if not activation_date:
        raise HTTPException(422, "Data di attivazione non rilevata: completa la verifica prima di importare")
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        if not competition.start_date <= activation_date <= competition.end_date:
            raise HTTPException(422, "La data della PDC non rientra nel periodo della gara selezionata")
        duplicate = db.scalar(
            select(IncentivePdcImport).where(
                IncentivePdcImport.competition_id == competition_id,
                IncentivePdcImport.file_sha256 == digest,
            )
        )
        if duplicate:
            raise HTTPException(409, "Questa PDC è già stata importata nella gara")
        customer_data = parsed["customer"]
        customer = db.scalar(select(Customer).where(Customer.fiscal_code == customer_data["fiscal_code"]))
        if not customer:
            customer = Customer(
                segment="CONSUMER", business_name=customer_data["business_name"],
                first_name=customer_data["first_name"] or None, last_name=customer_data["last_name"] or None,
                fiscal_code=customer_data["fiscal_code"] or None,
                phone=customer_data.get("contact_phone") or parsed["contract"]["phone"] or None,
                email=customer_data["email"] or None, address=customer_data["address"] or None,
                postal_code=customer_data["postal_code"] or None, city=customer_data["city"] or None,
                birth_date=parsed_date(customer_data["birth_date"]), document_type=customer_data["document_type"] or None,
                document_number=customer_data["document_number"] or None,
                document_issue_date=parsed_date(customer_data["document_issue_date"]),
                portfolio_status="ACTIVE",
            )
            db.add(customer)
            db.flush()
        else:
            for field, value in (
                ("first_name", customer_data["first_name"]), ("last_name", customer_data["last_name"]),
                ("phone", customer_data.get("contact_phone") or parsed["contract"]["phone"]), ("email", customer_data["email"]),
                ("address", customer_data["address"]), ("postal_code", customer_data["postal_code"]),
                ("city", customer_data["city"]), ("document_type", customer_data["document_type"]),
                ("document_number", customer_data["document_number"]),
            ):
                if value and not getattr(customer, field):
                    setattr(customer, field, value)
        code = parsed["contract"]["customer_code"]
        if code and not db.scalar(
            select(CustomerAccountCode).where(
                CustomerAccountCode.operator == "WINDTRE",
                CustomerAccountCode.market == "CONSUMER",
                CustomerAccountCode.customer_code == code,
            )
        ):
            db.add(CustomerAccountCode(
                customer_id=customer.id, operator="WINDTRE", market="CONSUMER",
                customer_code=code, is_primary=True,
            ))
        import_id = uuid.uuid4()
        pdc_dir = UPLOAD_DIR / "pdc"
        pdc_dir.mkdir(parents=True, exist_ok=True)
        stored_path = pdc_dir / f"{import_id}.pdf"
        stored_path.write_bytes(contents)
        activation_ids = []
        entries = list(parsed["proposed_entries"])
        if include_mobile and not any(entry.get("track") == "MOBILE" for entry in entries):
            entries.insert(0, {
                "track": "MOBILE", "label": "Nuova attivazione Mobile da PDC",
                "offer": parsed["contract"]["plan"], "direct_bonus": 5.0 if parsed["device"]["imei"] else 0,
                "monthly_fee": 0, "attributes": {
                    "mnp": False, "tied": parsed["contract"]["payment_method"] == "SDD",
                    "piva": False, "phone_included": bool(parsed["device"]["imei"]),
                    "device_sale": bool(parsed["device"]["imei"]), "classification_to_verify": True,
                },
                "status": "TO_VERIFY",
            })
        for index, entry in enumerate(entries):
            activation = IncentiveActivation(
                competition_id=competition_id, customer_id=customer.id, activation_date=activation_date,
                source_type="PDC", source_key=f"PDC:{digest}:{entry['track']}:{index}",
                seller_name=clean(seller_name) or None, track=entry["track"], offer=entry["offer"],
                asset_number=parsed["contract"]["phone"] or ("NUOVA LINEA" if entry["track"] == "FIXED" else None),
                monthly_fee_cents=round(entry.get("monthly_fee", 0) * 100),
                customer_code=parsed["contract"]["customer_code"] or None,
                contract_code=parsed["contract"]["contract_code"] or None,
                direct_bonus_cents=round(entry.get("direct_bonus", 0) * 100), attributes={
                    **entry.get("attributes", {}), "contract_code": parsed["contract"]["contract_code"],
                    "iccid": parsed["contract"]["iccid"], "imei": parsed["device"]["imei"],
                    "customer_code": parsed["contract"]["customer_code"], "source_pdc_id": str(import_id),
                }, status=entry.get("status", "VALID"),
                notes=entry["label"],
            )
            db.add(activation)
            db.flush()
            activation_ids.append(str(activation.id))
        record = IncentivePdcImport(
            id=import_id, competition_id=competition_id, customer_id=customer.id,
            file_name=file.filename or "PDC.pdf", file_sha256=digest,
            stored_path=str(stored_path), document_type=parsed["document_type"],
            status="IMPORTED", extracted_data=parsed, activation_ids=activation_ids,
        )
        db.add(record)
        db.flush()
        ensure_post_activation_tasks_for_record(db, record)
        db.commit()
        return {
            "id": str(record.id), "customer_id": str(customer.id), "customer_name": customer.business_name,
            "activation_ids": activation_ids, "entries_created": len(activation_ids),
            "report_url": f"/api/v1/incentives/{competition_id}/pdc-imports/{record.id}/report.pdf",
        }


@app.get("/api/v1/incentives/{competition_id}/pdc-imports")
def incentive_pdc_import_history(competition_id: uuid.UUID):
    with SessionLocal() as db:
        items = db.scalars(
            select(IncentivePdcImport)
            .where(IncentivePdcImport.competition_id == competition_id)
            .order_by(IncentivePdcImport.imported_at.desc())
        ).all()
        return [{
            "id": str(item.id), "file_name": item.file_name, "document_type": item.document_type,
            "status": item.status, "customer_id": str(item.customer_id) if item.customer_id else None,
            "customer_name": db.get(Customer, item.customer_id).business_name if item.customer_id and db.get(Customer, item.customer_id) else "",
            "activation_ids": item.activation_ids, "extracted_data": item.extracted_data,
            "imported_at": item.imported_at.isoformat(),
            "report_url": f"/api/v1/incentives/{competition_id}/pdc-imports/{item.id}/report.pdf",
        } for item in items]


@app.get("/api/v1/incentives/{competition_id}/pdc-imports/{pdc_import_id}/report.pdf")
def incentive_pdc_activation_report(competition_id: uuid.UUID, pdc_import_id: uuid.UUID):
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        record = db.get(IncentivePdcImport, pdc_import_id)
        if not competition or not record or record.competition_id != competition_id:
            raise HTTPException(404, "Importazione PDC non trovata")
        activation_ids = [uuid.UUID(value) for value in (record.activation_ids or [])]
        activations = db.scalars(select(IncentiveActivation).where(IncentiveActivation.id.in_(activation_ids))).all() if activation_ids else []
        full_report = incentive_report(db, competition)
        report_rows = [row for row in full_report["activations"] if row["id"] in set(record.activation_ids or [])]
        customer = db.get(Customer, record.customer_id) if record.customer_id else None
        extracted = record.extracted_data or {}
        operator_brand = db.scalar(select(OperatorBrand).where(OperatorBrand.operator == competition.operator))
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=A4, leftMargin=16*mm, rightMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm)
    styles = getSampleStyleSheet()
    story = []
    if operator_brand and operator_brand.logo_path:
        logo_path = UPLOAD_DIR / operator_brand.logo_path
        if logo_path.is_file():
            story.extend([Image(str(logo_path), width=42*mm, height=16*mm), Spacer(1, 4*mm)])
    story.extend([
        Paragraph("Report importazione attivazione PDC", styles["Title"]),
        Paragraph(f"{competition.name} · Dealer {competition.dealer_code or 'non indicato'}", styles["Normal"]),
        Spacer(1, 5*mm),
    ])
    contract = extracted.get("contract", {})
    device = extracted.get("device", {})
    customer_data = extracted.get("customer", {})
    number_iccid = " · ".join(value for value in (contract.get("phone"), contract.get("iccid")) if value) or "Nuova linea - numerazione da assegnare"
    terminal_parts = []
    if device.get("model"):
        terminal_parts.append(device["model"])
    if device.get("imei"):
        terminal_parts.append(f"IMEI {device['imei']}")
    terminal_description = " · ".join(terminal_parts) or "Non presente nella PDC"
    price_installments = (
        f"€ {device.get('price',0):.2f} · {device.get('installments',0)} rate"
        if device.get("price") or device.get("installments")
        else "Non applicabile / non presente"
    )
    summary = [
        ["Cliente", customer.business_name if customer else customer_data.get("business_name", "")],
        ["Codice fiscale", customer_data.get("fiscal_code", "")],
        ["Codice cliente", contract.get("customer_code", "")],
        ["Codice contratto", contract.get("contract_code", "")],
        ["Data attivazione", format_date_it(contract.get("activation_date"))],
        ["Numero / ICCID", number_iccid],
        ["Piano", contract.get("plan", "")],
        ["Terminale", terminal_description],
        ["Prezzo / rate", price_installments],
        ["Pagamento", contract.get("payment_method", "")],
    ]
    table = Table(summary, colWidths=[44*mm, 132*mm])
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(0,-1),colors.HexColor("#EEF1F6")),("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#CCD3DE")),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    story.extend([table, Spacer(1, 6*mm), Paragraph("Quote generate", styles["Heading2"])])
    quota_data = [["Pista", "Voce", "Punti", "Soglia", "Gettone base", "Commissione stimata"]]
    row_by_id = {row["id"]: row for row in report_rows}
    for activation in activations:
        row = row_by_id.get(str(activation.id), {})
        quota_data.append([
            activation.track, f"{activation.notes or activation.offer or ''}{' (da verificare)' if activation.status == 'TO_VERIFY' else ''}", row.get("points", 0),
            row.get("threshold", "Non raggiunta"), f"€ {activation.direct_bonus_cents/100:.2f}",
            f"€ {row.get('commission',0):.2f}",
        ])
    quota = Table(quota_data, repeatRows=1, colWidths=[28*mm,60*mm,18*mm,28*mm,25*mm,32*mm])
    quota.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#18233A")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#CCD3DE")),("FONTSIZE",(0,0),(-1,-1),8),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F6F7F9")])]))
    story.extend([quota, Spacer(1, 5*mm), Paragraph("Nota: la commissione è una stima basata sulla soglia complessiva attualmente raggiunta. Rimane soggetta alle verifiche WINDTRE su attivazione, rinnovi, silenza, MNP, storni e qualità.", styles["Normal"])])
    document.build(story)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="Report_Attivazione_{contract.get("phone") or pdc_import_id}.pdf"'})


@app.delete("/api/v1/incentives/{competition_id}/activations/{activation_id}")
def delete_incentive_activation(competition_id: uuid.UUID, activation_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(IncentiveActivation, activation_id)
        if not item or item.competition_id != competition_id:
            raise HTTPException(404, "Attivazione non trovata")
        db.delete(item)
        db.commit()
        return {"deleted": True}


@app.post("/api/v1/incentives/{competition_id}/sync-windtre")
def sync_incentive_activations(competition_id: uuid.UUID):
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        latest_import = db.scalar(
            select(WindTreImport).order_by(WindTreImport.competence_month.desc(), WindTreImport.uploaded_at.desc()).limit(1)
        )
        if not latest_import:
            raise HTTPException(422, "Non sono presenti estrazioni WINDTRE")
        rows = db.scalars(select(WindTreImportRow).where(WindTreImportRow.import_id == latest_import.id)).all()
        existing_keys = set(db.scalars(
            select(IncentiveActivation.source_key).where(IncentiveActivation.competition_id == competition_id)
        ).all())
        added = skipped = 0
        for row in rows:
            activation_text = find_activation_date(row.raw_data)
            if not activation_text:
                skipped += 1
                continue
            try:
                activation_date = date.fromisoformat(activation_text[:10])
            except ValueError:
                skipped += 1
                continue
            if not competition.start_date <= activation_date <= competition.end_date:
                continue
            category = classify_asset(row)
            track = "MOBILE" if category.startswith("MOBILE") else "FIXED" if category == "FIXED_DATA" else None
            if not track:
                continue
            source_key = f"DBTOOL:{row.customer_id}:{row.asset_key}:{activation_date.isoformat()}"
            if source_key in existing_keys:
                continue
            customer = db.get(Customer, row.customer_id) if row.customer_id else None
            raw = {normalize_header(key): clean(value) for key, value in (row.raw_data or {}).items()}
            raw_text = " ".join(f"{key} {value}" for key, value in raw.items()).upper()
            offer = row.current_plan or ""
            status_text = (row.current_status or "").upper()
            attrs = {
                "mnp": "MNP" in raw_text and not any(token in raw_text for token in ("NO MNP", "NOMNP")),
                "tied": any(token in raw_text or token in offer.upper() for token in ("TIED", "EASY PAY", "EASYPAY")),
                "piva": bool(customer and customer.tax_id),
                "convergent": any("CONVERGEN" in key and value.upper() not in {"", "NO", "N", "0"} for key, value in raw.items()),
                "ftth": "FTTH" in raw_text,
                "fwa": "FWA" in raw_text,
                "first_line": not any(token in raw_text for token in ("2 LINEA", "SECONDA LINEA", "2° LINEA")),
                "second_professional": any(token in raw_text for token in ("2 LINEA PROFESSIONAL", "SECONDA LINEA PROFESSIONAL")),
                "professional_box": "PROFESSIONAL BOX" in raw_text and ("FRITZ" in raw_text or "FRITZ" in offer.upper()),
                "secure_option": "PIU SICURI" in raw_text or "PIÙ SICURI" in raw_text,
                "phone_included": "TELEFONO INCLUSO" in raw_text,
                "premium_tied_offer": any(token in offer.upper() for token in ("START UNLIMITED", "UNLIMITED 5G", "UNLIMITED PRO", "RELOAD EXCHANGE")),
            }
            item = IncentiveActivation(
                competition_id=competition_id, customer_id=row.customer_id, activation_date=activation_date,
                source_type="DB_TOOL", source_key=source_key, track=track, offer=row.current_plan,
                asset_number=row.asset_number, monthly_fee_cents=round(parse_monthly_fee(row.monthly_fee) * 100),
                attributes=attrs, status="VALID" if not status_text or status_text in {"ATT", "ACTIVE", "ATTIVO", "ATTIVA"} else "TO_VERIFY",
            )
            db.add(item)
            existing_keys.add(source_key)
            added += 1
        db.commit()
        return {"added": added, "skipped_without_activation_date": skipped, "source_import": latest_import.file_name}


@app.get("/api/v1/consumer-activations/dashboard")
def consumer_activations_dashboard(
    date_from: date | None = Query(None),
    date_to: date | None = Query(None),
):
    with SessionLocal() as db:
        consumer_pdc_records = db.scalars(
            select(IncentivePdcImport)
            .join(Customer, Customer.id == IncentivePdcImport.customer_id)
            .where(Customer.segment == "CONSUMER")
        ).all()
        for record in consumer_pdc_records:
            ensure_post_activation_tasks_for_record(db, record)
        db.commit()

        statement = (
            select(IncentiveActivation, Customer)
            .join(Customer, Customer.id == IncentiveActivation.customer_id)
            .where(Customer.segment == "CONSUMER")
        )
        if date_from:
            statement = statement.where(IncentiveActivation.activation_date >= date_from)
        if date_to:
            statement = statement.where(IncentiveActivation.activation_date <= date_to)
        pairs = db.execute(
            statement.order_by(IncentiveActivation.activation_date.desc(), IncentiveActivation.created_at.desc())
        ).all()

        commissions: dict[str, float] = {}
        competition_operators: dict[uuid.UUID, str] = {}
        competition_ids = {activation.competition_id for activation, _ in pairs}
        for competition_id in competition_ids:
            competition = db.get(IncentiveCompetition, competition_id)
            if competition:
                competition_operators[competition_id] = competition.operator
                for row in incentive_report(db, competition)["activations"]:
                    commissions[row["id"]] = row["commission"]

        rows = []
        daily: dict[str, dict[str, Any]] = {}
        track_counts: dict[str, int] = defaultdict(int)
        status_counts: dict[str, int] = defaultdict(int)
        customer_ids: set[uuid.UUID] = set()
        for activation, customer in pairs:
            serialized = serialize_incentive_activation(activation, customer.business_name)
            serialized["operator"] = competition_operators.get(activation.competition_id, "WINDTRE")
            commission = commissions.get(str(activation.id), 0)
            serialized["commission"] = commission
            rows.append(serialized)
            customer_ids.add(customer.id)
            track_counts[activation.track] += 1
            status_counts[activation.status] += 1
            day_key = activation.activation_date.isoformat()
            day = daily.setdefault(day_key, {"date": day_key, "events": 0, "commission": 0.0})
            day["events"] += 1
            day["commission"] = round(day["commission"] + commission, 2)

        pdc_statement = (
            select(func.count(IncentivePdcImport.id))
            .join(Customer, Customer.id == IncentivePdcImport.customer_id)
            .where(Customer.segment == "CONSUMER")
        )
        if date_from:
            pdc_statement = pdc_statement.where(
                IncentivePdcImport.extracted_data["contract"]["activation_date"].astext >= date_from.isoformat()
            )
        if date_to:
            pdc_statement = pdc_statement.where(
                IncentivePdcImport.extracted_data["contract"]["activation_date"].astext <= date_to.isoformat()
            )
        pdc_count = db.scalar(pdc_statement) or 0
        task_statement = (
            select(PostActivationTask, Customer)
            .join(Customer, Customer.id == PostActivationTask.customer_id)
            .where(Customer.segment == "CONSUMER")
        )
        if date_from:
            task_statement = task_statement.where(PostActivationTask.activation_date >= date_from)
        if date_to:
            task_statement = task_statement.where(PostActivationTask.activation_date <= date_to)
        task_pairs = db.execute(
            task_statement.order_by(PostActivationTask.activation_date.desc(), PostActivationTask.created_at.desc())
        ).all()
        rule_items = db.scalars(select(PostActivationRule)).all()
        rules_by_key = {
            (rule.operator, rule.item_type, rule.item_key): rule for rule in rule_items
        }
        post_activation_tasks = [
            serialize_post_activation_task(
                task, customer.business_name,
                rules_by_key.get(("WINDTRE", task.item_type, task.item_key)),
            )
            for task, customer in task_pairs
        ]
        post_activation_tasks.sort(key=lambda item: (
            {"OVERDUE": 0, "DUE_TODAY": 1, "SCHEDULED": 2, "REVIEW": 3, "DONE": 4}.get(item["alert_state"], 5),
            item["due_date"] or "9999-12-31",
        ))
        pending_tasks = [
            item for item in post_activation_tasks
            if item["action_required"] and item["status"] != "DONE"
        ]
        review_tasks = [
            item for item in post_activation_tasks
            if not item["action_required"] and item["status"] == "REVIEW"
        ]

        return {
            "filters": {
                "date_from": date_from.isoformat() if date_from else None,
                "date_to": date_to.isoformat() if date_to else None,
            },
            "summary": {
                "pdc_imported": pdc_count,
                "customers": len(customer_ids),
                "events": len(rows),
                "valid_events": status_counts.get("VALID", 0),
                "to_verify": status_counts.get("TO_VERIFY", 0),
                "mobile": track_counts.get("MOBILE", 0),
                "fixed": track_counts.get("FIXED", 0),
                "customer_base": track_counts.get("CUSTOMER_BASE", 0),
                "reload": track_counts.get("RELOAD", 0),
                "monthly_revenue": round(sum(item.monthly_fee_cents for item, _ in pairs) / 100, 2),
                "commissioning": round(sum(commissions.get(str(item.id), 0) for item, _ in pairs), 2),
                "post_activation_pending": len(pending_tasks),
                "post_activation_overdue": sum(item["alert_state"] == "OVERDUE" for item in pending_tasks),
                "post_activation_review": len(review_tasks),
            },
            "tracks": [{"track": key, "events": value} for key, value in sorted(track_counts.items())],
            "statuses": [{"status": key, "events": value} for key, value in sorted(status_counts.items())],
            "daily": [daily[key] for key in sorted(daily)],
            "activations": rows[:100],
            "post_activation_tasks": post_activation_tasks,
        }


@app.patch("/api/v1/post-activation-tasks/{task_id}")
def update_post_activation_task(task_id: uuid.UUID, data: PostActivationTaskUpdate):
    with SessionLocal() as db:
        item = db.get(PostActivationTask, task_id)
        if not item:
            raise HTTPException(404, "Verifica post-attivazione non trovata")
        rule = db.scalar(select(PostActivationRule).where(
            PostActivationRule.operator == "WINDTRE",
            PostActivationRule.item_type == item.item_type,
            PostActivationRule.item_key == item.item_key,
        ))
        if data.mark_completed:
            if not item.action_required:
                raise HTTPException(422, "Seleziona prima l'opzione come da disattivare")
            item.status = "DONE"
            item.completed_date = data.completed_date or date.today()
        elif data.action_required is not None:
            if data.action_required and not (rule and rule.is_active and rule.can_deactivate):
                raise HTTPException(422, "Abilita prima la disattivazione nella configurazione post-vendita")
            item.action_required = data.action_required
            item.status = "PENDING" if data.action_required else "REVIEW"
            item.due_date = first_business_day_next_month(item.activation_date) if data.action_required else None
            item.completed_date = None
        if data.notes is not None:
            item.notes = clean(data.notes) or None
        db.commit()
        db.refresh(item)
        customer = db.get(Customer, item.customer_id) if item.customer_id else None
        return serialize_post_activation_task(item, customer.business_name if customer else "", rule)


@app.get("/api/v1/post-activation-rules")
def post_activation_rules(
    operator: str | None = None,
    item_type: str | None = None,
    include_inactive: bool = True,
):
    with SessionLocal() as db:
        for record in db.scalars(select(IncentivePdcImport)).all():
            ensure_post_activation_tasks_for_record(db, record)
        db.commit()
        statement = select(PostActivationRule)
        if operator:
            statement = statement.where(PostActivationRule.operator == clean(operator).upper())
        if item_type:
            statement = statement.where(PostActivationRule.item_type == clean(item_type).upper())
        if not include_inactive:
            statement = statement.where(PostActivationRule.is_active.is_(True))
        items = db.scalars(statement.order_by(
            PostActivationRule.operator, PostActivationRule.item_type, PostActivationRule.item_name
        )).all()
        return [serialize_post_activation_rule(item) for item in items]


@app.post("/api/v1/post-activation-rules")
def create_post_activation_rule(data: PostActivationRuleRequest):
    operator = clean(data.operator).upper() or "WINDTRE"
    item_type = clean(data.item_type).upper()
    item_name = clean(data.item_name)
    if item_type not in {"OFFER", "SERVICE", "OPTION"}:
        raise HTTPException(422, "Tipologia non valida")
    if not item_name:
        raise HTTPException(422, "Nome obbligatorio")
    item_key = post_activation_rule_key(item_type, item_name)
    with SessionLocal() as db:
        if db.scalar(select(PostActivationRule).where(
            PostActivationRule.operator == operator,
            PostActivationRule.item_type == item_type,
            PostActivationRule.item_key == item_key,
        )):
            raise HTTPException(409, "Elemento già presente")
        rule = PostActivationRule(
            operator=operator, item_type=item_type, item_key=item_key, item_name=item_name,
            can_deactivate=data.can_deactivate,
            default_action_required=data.default_action_required and data.can_deactivate,
            is_active=data.is_active, notes=clean(data.notes) or None,
        )
        db.add(rule)
        db.commit()
        db.refresh(rule)
        return serialize_post_activation_rule(rule)


@app.patch("/api/v1/post-activation-rules/{rule_id}")
def update_post_activation_rule(rule_id: uuid.UUID, data: PostActivationRuleUpdate):
    with SessionLocal() as db:
        rule = db.get(PostActivationRule, rule_id)
        if not rule:
            raise HTTPException(404, "Configurazione non trovata")
        if data.can_deactivate is not None:
            rule.can_deactivate = data.can_deactivate
            if not data.can_deactivate:
                rule.default_action_required = False
        if data.default_action_required is not None:
            if data.default_action_required and not rule.can_deactivate:
                raise HTTPException(422, "Abilita prima la possibilità di disattivazione")
            rule.default_action_required = data.default_action_required
        if data.is_active is not None:
            rule.is_active = data.is_active
        if data.notes is not None:
            rule.notes = clean(data.notes) or None
        if rule.can_deactivate and rule.default_action_required and rule.is_active:
            pending_items = db.scalars(select(PostActivationTask).where(
                PostActivationTask.item_type == rule.item_type,
                PostActivationTask.item_key == rule.item_key,
                PostActivationTask.status == "REVIEW",
            )).all()
            for item in pending_items:
                item.action_required = True
                item.status = "PENDING"
                item.due_date = first_business_day_next_month(item.activation_date)
        db.commit()
        db.refresh(rule)
        return serialize_post_activation_rule(rule)


@app.get("/api/v1/incentives/{competition_id}/report")
def get_incentive_report(competition_id: uuid.UUID):
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        return incentive_report(db, competition)


@app.get("/api/v1/incentives/{competition_id}/report.xlsx")
def export_incentive_report_excel(competition_id: uuid.UUID):
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        report = incentive_report(db, competition)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Soglie"
    sheet.append(["Pista", "Eventi validi", "Punti", "Soglia raggiunta", "Prossimo target", "Mancanti", "Commissioning"])
    for item in report["tracks"]:
        sheet.append([item["label"], item["valid_events"], item["points"], item["reached"], item["next_target"], item["remaining"], item["commission"]])
    sheet.append(["TOTALE", report["valid_events"], "", "", "", "", report["commissioning_total"]])
    detail = workbook.create_sheet("Commissioning")
    detail.append(["Data", "Cliente", "Pista", "Utenza", "Offerta", "Canone", "Punti", "Soglia", "Moltiplicatore", "Gettone", "Commissione", "Esito"])
    for item in report["activations"]:
        detail.append([
            item["activation_date"], item["customer_name"], item["track"], item["asset_number"], item["offer"],
            item["monthly_fee"], item["points"], item["threshold"], item["multiplier"], item["direct_bonus"],
            item["commission"], "Valida" if item["eligible"] else "Non remunerata",
        ])
    for current in (sheet, detail):
        current.freeze_panes = "A2"
        current.auto_filter.ref = current.dimensions
        for cell in current[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="18233A")
        for column in current.columns:
            current.column_dimensions[column[0].column_letter].width = min(36, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    filename = f"Report_Gara_{competition.start_date.strftime('%Y_%m')}.xlsx"
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/v1/incentives/{competition_id}/report.pdf")
def export_incentive_report_pdf(competition_id: uuid.UUID):
    with SessionLocal() as db:
        competition = db.get(IncentiveCompetition, competition_id)
        if not competition:
            raise HTTPException(404, "Gara non trovata")
        report = incentive_report(db, competition)
        operator_brand = db.scalar(select(OperatorBrand).where(OperatorBrand.operator == competition.operator))
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    story = []
    if operator_brand and operator_brand.logo_path:
        logo_path = UPLOAD_DIR / operator_brand.logo_path
        if logo_path.is_file():
            story.extend([Image(str(logo_path), width=48*mm, height=18*mm), Spacer(1, 4*mm)])
    story.extend([
        Paragraph(competition.name, styles["Title"]),
        Paragraph(f"Periodo {competition.start_date.strftime('%d/%m/%Y')} - {competition.end_date.strftime('%d/%m/%Y')} · Dealer {competition.dealer_code or 'non indicato'}", styles["Normal"]),
        Spacer(1, 5 * mm),
    ])
    threshold_data = [["Pista", "Validi", "Punti", "Soglia", "Mancanti", "Commissioning"]]
    threshold_data += [[item["label"], item["valid_events"], item["points"], item["reached"], item["remaining"], f"€ {item['commission']:.2f}"] for item in report["tracks"]]
    threshold_data.append(["TOTALE", report["valid_events"], "", "", "", f"€ {report['commissioning_total']:.2f}"])
    table = Table(threshold_data, repeatRows=1, colWidths=[62*mm, 22*mm, 22*mm, 35*mm, 25*mm, 35*mm])
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#18233A")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#CBD2DD")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-2),[colors.white,colors.HexColor("#F4F6F9")]),("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"),("ALIGN",(1,1),(-1,-1),"RIGHT"),("BOTTOMPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),6)]))
    story.extend([table, Spacer(1, 7*mm), Paragraph("Dettaglio commissioning", styles["Heading2"])])
    detail_data = [["Data", "Cliente", "Pista", "Utenza", "Offerta", "Punti", "Moltip.", "Commissione"]]
    for item in report["activations"]:
        detail_data.append([item["activation_date"], item["customer_name"][:25], item["track"], item["asset_number"], item["offer"][:28], item["points"], item["multiplier"], f"€ {item['commission']:.2f}"])
    detail_table = Table(detail_data, repeatRows=1, colWidths=[24*mm,42*mm,22*mm,30*mm,56*mm,17*mm,20*mm,27*mm])
    detail_table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#18233A")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D6DBE4")),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7F8FA")])]))
    story.append(detail_table)
    document.build(story)
    stream.seek(0)
    filename = f"Report_Gara_{competition.start_date.strftime('%Y_%m')}.pdf"
    return StreamingResponse(stream, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/v1/dashboard/summary")
def summary():
    with SessionLocal() as db:
        customer_count = db.scalar(select(func.count()).select_from(Customer)) or 0
        last_import = db.scalar(select(WindTreImport).order_by(WindTreImport.uploaded_at.desc()).limit(1))
        return {
            "customers": customer_count,
            "active_contracts": 0,
            "open_practices": (last_import.field_changes + last_import.campaign_changes) if last_import else 0,
            "monthly_value": 0,
            "last_import": serialize_import(last_import) if last_import else None,
        }


@app.get("/api/v1/dashboard/portfolio")
def portfolio_summary(segment: str = "BUSINESS_SME"):
    with SessionLocal() as db:
        customers_query = select(func.count()).select_from(Customer).where(
            Customer.segment == segment,
            Customer.portfolio_status == "ACTIVE",
        )
        customer_count = db.scalar(customers_query) or 0
        latest_import = db.scalar(
            select(WindTreImport)
            .order_by(WindTreImport.competence_month.desc(), WindTreImport.uploaded_at.desc())
            .limit(1)
        )
        rows = (
            db.scalars(
                select(WindTreImportRow)
                .join(Customer, Customer.id == WindTreImportRow.customer_id)
                .where(
                    WindTreImportRow.import_id == latest_import.id,
                    Customer.segment == segment,
                )
            ).all()
            if latest_import
            else []
        )
        totals = {
            "MOBILE_VOICE": {"count": 0, "mrr": 0.0},
            "MOBILE_DATA": {"count": 0, "mrr": 0.0},
            "MOBILE_M2M": {"count": 0, "mrr": 0.0},
            "MOBILE_OTHER": {"count": 0, "mrr": 0.0},
            "FIXED_DATA": {"count": 0, "mrr": 0.0},
            "ICT": {"count": 0, "mrr": 0.0},
            "OTHER": {"count": 0, "mrr": 0.0},
        }
        for row in rows:
            category = classify_asset(row)
            totals[category]["count"] += 1
            totals[category]["mrr"] = round(
                totals[category]["mrr"] + parse_monthly_fee(row.monthly_fee),
                2,
            )
        mobile_categories = ("MOBILE_VOICE", "MOBILE_DATA", "MOBILE_M2M", "MOBILE_OTHER")
        mobile = {
            "count": sum(totals[key]["count"] for key in mobile_categories),
            "mrr": round(sum(totals[key]["mrr"] for key in mobile_categories), 2),
            "breakdown": {
                "voice": totals["MOBILE_VOICE"],
                "data": totals["MOBILE_DATA"],
                "m2m": totals["MOBILE_M2M"],
                "other": totals["MOBILE_OTHER"],
            },
        }
        return {
            "segment": segment,
            "competence_month": latest_import.competence_month if latest_import else None,
            "customers": customer_count,
            "mobile": mobile,
            "fixed_data": totals["FIXED_DATA"],
            "ict": totals["ICT"],
            "other_services": totals["OTHER"],
            "total_mrr": round(sum(item["mrr"] for item in totals.values()), 2),
        }


@app.get("/api/v1/customers")
def customers(search: str = "", segment: str | None = None, limit: int = Query(100, ge=1, le=500)):
    with SessionLocal() as db:
        query = select(Customer).order_by(Customer.business_name).limit(limit)
        if segment:
            query = query.where(Customer.segment == segment)
        if search.strip():
            pattern = f"%{search.strip()}%"
            query = query.where(
                or_(
                    Customer.business_name.ilike(pattern),
                    Customer.tax_id.ilike(pattern),
                    Customer.fiscal_code.ilike(pattern),
                    Customer.windtre_customer_code.ilike(pattern),
                    Customer.email.ilike(pattern),
                    Customer.phone.ilike(pattern),
                    Customer.id.in_(
                        select(CustomerAccountCode.customer_id).where(
                            CustomerAccountCode.customer_code.ilike(pattern)
                        )
                    ),
                )
            )
        items = db.scalars(query).all()
        codes_by_customer: dict[uuid.UUID, list[dict[str, Any]]] = defaultdict(list)
        if items:
            for code in db.scalars(
                select(CustomerAccountCode)
                .where(CustomerAccountCode.customer_id.in_([item.id for item in items]))
                .order_by(CustomerAccountCode.operator, CustomerAccountCode.market, CustomerAccountCode.customer_code)
            ).all():
                codes_by_customer[code.customer_id].append(serialize_account_code(code))
        latest_import = db.scalar(
            select(WindTreImport)
            .order_by(WindTreImport.competence_month.desc(), WindTreImport.uploaded_at.desc())
            .limit(1)
        )
        monthly_spend_by_customer: dict[uuid.UUID, float] = {}
        if latest_import and items:
            current_rows = db.scalars(
                select(WindTreImportRow).where(
                    WindTreImportRow.import_id == latest_import.id,
                    WindTreImportRow.customer_id.in_([item.id for item in items]),
                )
            ).all()
            for row in current_rows:
                if row.customer_id:
                    monthly_spend_by_customer[row.customer_id] = round(
                        monthly_spend_by_customer.get(row.customer_id, 0) + parse_monthly_fee(row.monthly_fee),
                        2,
                    )
        return [
            {
                "id": str(item.id),
                "business_name": item.business_name,
                "first_name": item.first_name,
                "last_name": item.last_name,
                "segment": item.segment,
                "tax_id": item.tax_id,
                "fiscal_code": item.fiscal_code,
                "windtre_customer_code": item.windtre_customer_code,
                "account_codes": codes_by_customer.get(item.id, []),
                "portfolio_status": item.portfolio_status,
                "first_seen_month": item.first_seen_month,
                "last_seen_month": item.last_seen_month,
                "monthly_spend": monthly_spend_by_customer.get(item.id, 0),
                "email": item.email,
                "phone": item.phone,
            }
            for item in items
        ]


@app.get("/api/v1/customers/{customer_id}")
def customer_detail(customer_id: uuid.UUID):
    with SessionLocal() as db:
        customer = db.get(Customer, customer_id)
        if not customer:
            raise HTTPException(404, "Cliente non trovato")
        latest_import = db.scalar(
            select(WindTreImport)
            .order_by(WindTreImport.competence_month.desc(), WindTreImport.uploaded_at.desc())
            .limit(1)
        )
        latest_rows = (
            db.scalars(
                select(WindTreImportRow)
                .where(
                    WindTreImportRow.customer_id == customer_id,
                    WindTreImportRow.import_id == latest_import.id,
                )
                .order_by(WindTreImportRow.row_number)
            ).all()
            if latest_import
            else []
        )
        monthly_spend = round(sum(parse_monthly_fee(row.monthly_fee) for row in latest_rows), 2)
        account_codes = db.scalars(
            select(CustomerAccountCode)
            .where(CustomerAccountCode.customer_id == customer_id)
            .order_by(CustomerAccountCode.operator, CustomerAccountCode.market, CustomerAccountCode.customer_code)
        ).all()
        imported_numbers = {clean(row.asset_number) for row in latest_rows if clean(row.asset_number)}
        inventory_sims = db.scalars(
            select(SimInventory).where(SimInventory.customer_id == customer_id).order_by(SimInventory.msisdn, SimInventory.iccid)
        ).all()
        sim_assets = [
            {
                "asset_key": f"SIM:{sim.id}",
                "asset_type": "SIM MOBILE",
                "asset_number": sim.msisdn or sim.iccid,
                "plan": sim.product.name if sim.product else "",
                "status": sim.status,
                "monthly_fee": "",
                "activation_date": sim.updated_at.isoformat() if sim.status == "ATTIVATA" else None,
                "details": [{"key": "ICCID", "label": "ICCID", "value": sim.iccid}],
                "campaigns": {},
            }
            for sim in inventory_sims
            if not sim.msisdn or clean(sim.msisdn) not in imported_numbers
        ]
        return {
            "id": str(customer.id),
            "business_name": customer.business_name,
            "first_name": customer.first_name,
            "last_name": customer.last_name,
            "segment": customer.segment,
            "tax_id": customer.tax_id,
            "fiscal_code": customer.fiscal_code,
            "windtre_customer_code": customer.windtre_customer_code,
            "account_codes": [serialize_account_code(code) for code in account_codes],
            "email": customer.email,
            "phone": customer.phone,
            "address": customer.address,
            "postal_code": customer.postal_code,
            "city": customer.city,
            "province": customer.province,
            "birth_date": customer.birth_date.isoformat() if customer.birth_date else None,
            "birth_place": customer.birth_place,
            "birth_province": customer.birth_province,
            "gender": customer.gender,
            "document_type": customer.document_type,
            "document_number": customer.document_number,
            "document_issue_date": customer.document_issue_date.isoformat() if customer.document_issue_date else None,
            "document_expiry_date": customer.document_expiry_date.isoformat() if customer.document_expiry_date else None,
            "document_issuer": customer.document_issuer,
            "portfolio_status": customer.portfolio_status,
            "first_seen_month": customer.first_seen_month,
            "last_seen_month": customer.last_seen_month,
            "snapshot_month": latest_import.competence_month if latest_import else None,
            "monthly_spend": monthly_spend,
            "assets": [
                {
                    "asset_key": row.asset_key,
                    "asset_type": row.asset_type,
                    "asset_number": row.asset_number,
                    "plan": row.current_plan,
                    "status": row.current_status,
                    "monthly_fee": row.monthly_fee,
                    "activation_date": find_activation_date(row.raw_data),
                    "details": asset_details(row.raw_data),
                    "campaigns": row.campaigns,
                }
                for row in latest_rows[-100:]
            ] + sim_assets,
        }


def configurator_plan_type(row: WindTreImportRow) -> str | None:
    return {
        "MOBILE_VOICE": "VOCE",
        "MOBILE_DATA": "DATI",
        "MOBILE_M2M": "DATI_M2M",
        "FIXED_DATA": "FISSO",
    }.get(classify_asset(row))


def configurator_proposal(db, row: WindTreImportRow) -> dict[str, Any]:
    current_fee = parse_monthly_fee(row.monthly_fee)
    raw = row.raw_data or {}
    renewal = any("RINNOV" in normalize_header(name) and clean(value) for name, value in (row.campaigns or {}).items())
    plan_type = configurator_plan_type(row)
    candidate = None
    if plan_type:
        candidates = db.scalars(
            select(TariffPlan)
            .where(TariffPlan.plan_type == plan_type, TariffPlan.subscribable.is_(True))
            .order_by(TariffPlan.monthly_fee_cents)
        ).all()
        if candidates:
            candidate = min(
                candidates,
                key=lambda item: abs(((item.monthly_fee_cents + item.secure_web_cents) / 100) - current_fee),
            )
    if renewal:
        proposed_plan = candidate.name if candidate else (row.current_plan or "Piano attuale")
        proposed_fee = ((candidate.monthly_fee_cents + candidate.secure_web_cents) / 100) if candidate else current_fee
        offer_type = "Rinnovo"
    elif candidate and normalize_header(candidate.name) != normalize_header(row.current_plan):
        proposed_plan = candidate.name
        proposed_fee = (candidate.monthly_fee_cents + candidate.secure_web_cents) / 100
        offer_type = "Cambio Piano"
    else:
        proposed_plan = row.current_plan or "Piano attuale"
        proposed_fee = current_fee
        offer_type = "Rinnovo Isopiano"
    hardware = next(
        (clean(raw.get(key)) for key in ("DESCRIZIONE_TERMINALE", "TERMINALE", "DEVICE", "MODELLO_TERMINALE") if clean(raw.get(key))),
        "",
    )
    return {
        "asset_key": row.asset_key,
        "msisdn": row.asset_number or row.asset_key,
        "line_type": classify_asset(row),
        "current_plan": row.current_plan or "",
        "current_fee": round(current_fee, 2),
        "proposed_plan": proposed_plan,
        "proposed_fee": round(proposed_fee, 2),
        "offer_type": offer_type,
        "activation_cost": round((candidate.activation_cost_cents / 100) if candidate else 0, 2),
        "hardware": hardware,
        "hardware_monthly": 0,
        "hardware_final": 0,
    }


def configurator_workbook(customer: Customer, proposals: list[dict[str, Any]]) -> Workbook:
    template = UPLOAD_DIR / "templates" / "template_configuratore.xlsx"
    if template.exists():
        workbook = load_workbook(template)
    else:
        workbook = Workbook()
        first = workbook.active
        first.title = "P1 - Proposta e Firma"
        second = workbook.create_sheet("P2 - Condizioni Economiche")
        first["B3"] = "PROPOSTA COMMERCIALE"
        first["B3"].font = Font(size=18, bold=True, color="FFFFFF")
        first["B3"].fill = PatternFill("solid", fgColor="292C3A")
        first.merge_cells("B3:J4")
        for cell, label in (("F20", "P.IVA / Codice fiscale"), ("F21", "Ragione sociale"), ("F22", "Sede commerciale")):
            first[cell] = label
            first[cell].font = Font(bold=True, color="626B7D")
        headers = {
            "B7": "MSISDN", "C7": "Piano attuale", "D7": "Costo attuale", "E7": "Proposta",
            "G7": "Prezzo proposto", "I7": "Tipo offerta", "J7": "Costo attivazione",
            "K7": "Terminale", "N7": "Rata mensile", "O7": "Rata finale",
        }
        for cell, value in headers.items():
            second[cell] = value
            second[cell].font = Font(bold=True, color="FFFFFF")
            second[cell].fill = PatternFill("solid", fgColor="292C3A")
        second.freeze_panes = "B8"
    first = workbook["P1 - Proposta e Firma"]
    second = workbook["P2 - Condizioni Economiche"]
    first["H20"] = customer.tax_id or customer.fiscal_code or ""
    first["H21"] = customer.business_name
    first["H22"] = customer.address or ""
    for index, proposal in enumerate(proposals, 8):
        if index > 8:
            for column in range(1, 16):
                source = second.cell(8, column)
                target = second.cell(index, column)
                if source.has_style:
                    target._style = copy(source._style)
                target.number_format = source.number_format
        values = {
            2: proposal["msisdn"], 3: proposal["current_plan"], 4: proposal["current_fee"],
            5: proposal["proposed_plan"], 7: proposal["proposed_fee"], 9: proposal["offer_type"],
            10: proposal["activation_cost"], 11: proposal["hardware"],
            14: proposal["hardware_monthly"], 15: proposal["hardware_final"],
        }
        for column, value in values.items():
            second.cell(index, column, value)
        for column in (4, 7, 10, 14, 15):
            second.cell(index, column).number_format = '€ #,##0.00'
    for column, width in {"B": 18, "C": 28, "D": 15, "E": 28, "G": 18, "I": 22, "J": 18, "K": 28, "N": 16, "O": 16}.items():
        if second.column_dimensions[column].width is None or second.column_dimensions[column].width < width:
            second.column_dimensions[column].width = width
    return workbook


@app.get("/api/v1/customers/{customer_id}/quotes")
def customer_quotes(customer_id: uuid.UUID):
    with SessionLocal() as db:
        if not db.get(Customer, customer_id):
            raise HTTPException(404, "Cliente non trovato")
        items = db.scalars(select(Quote).where(Quote.customer_id == customer_id).order_by(Quote.created_at.desc())).all()
        return [{
            "id": str(item.id), "quote_type": item.quote_type, "file_name": item.file_name,
            "status": item.status, "line_count": item.line_count,
            "current_mrr": item.current_mrr_cents / 100, "proposed_mrr": item.proposed_mrr_cents / 100,
            "created_at": item.created_at.isoformat(), "payload": item.payload,
        } for item in items]


@app.post("/api/v1/customers/{customer_id}/configurator.xlsx")
def generate_customer_configurator(customer_id: uuid.UUID, data: ConfiguratorRequest):
    if not data.asset_keys:
        raise HTTPException(422, "Seleziona almeno una linea")
    with SessionLocal() as db:
        customer = db.get(Customer, customer_id)
        if not customer:
            raise HTTPException(404, "Cliente non trovato")
        latest_import, rows = latest_customer_snapshot(db, customer_id)
        selected = [row for row in rows if row.asset_key in set(data.asset_keys) and is_active_service(row)]
        if not selected:
            raise HTTPException(422, "Le linee selezionate non risultano attive nell’ultima estrazione")
        proposals = [configurator_proposal(db, row) for row in selected]
        workbook = configurator_workbook(customer, proposals)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        safe_customer = re.sub(r"[^A-Za-z0-9]+", "_", customer.business_name).strip("_") or "Cliente"
        line_label = re.sub(r"[^A-Za-z0-9]+", "_", proposals[0]["msisdn"]).strip("_") if len(proposals) == 1 else "Multi"
        filename = f"Configuratore_{safe_customer}_{line_label}.xlsx"
        quote = Quote(
            customer_id=customer.id,
            file_name=filename,
            line_count=len(proposals),
            current_mrr_cents=round(sum(item["current_fee"] for item in proposals) * 100),
            proposed_mrr_cents=round(sum(item["proposed_fee"] for item in proposals) * 100),
            payload={"snapshot_month": latest_import.competence_month if latest_import else None, "proposals": proposals},
        )
        db.add(quote)
        db.commit()
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "X-Quote-Id": str(quote.id),
            },
        )


@app.get("/api/v1/customers/{customer_id}/services-pivot.xlsx")
def customer_services_excel(customer_id: uuid.UUID):
    with SessionLocal() as db:
        customer = db.get(Customer, customer_id)
        if not customer:
            raise HTTPException(404, "Cliente non trovato")
        latest_import, rows = latest_customer_snapshot(db, customer_id)
        pivot = customer_service_pivot(rows)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Servizi attivi"
        sheet.merge_cells("A1:C1")
        sheet["A1"] = f"Servizi attivi - {customer.business_name}"
        sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="1B2035")
        sheet["A1"].alignment = Alignment(horizontal="left")
        sheet["A2"] = "Mese fotografia"
        sheet["B2"] = latest_import.competence_month if latest_import else "N/D"
        sheet.append([])
        sheet.append(["Piano Tariffario", "Utenze Attive", "Ricavo Mensile (MRR)"])
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="287A51")
        for item in pivot:
            sheet.append([item["plan"], item["count"], item["mrr"]])
            sheet.cell(sheet.max_row, 3).number_format = '€ #,##0.00'
        total_count = sum(int(item["count"]) for item in pivot)
        total_mrr = round(sum(float(item["mrr"]) for item in pivot), 2)
        sheet.append(["TOTALE GENERALE", total_count, total_mrr])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8F3EC")
        sheet.cell(sheet.max_row, 3).number_format = '€ #,##0.00'
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 18
        sheet.column_dimensions["C"].width = 24
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:C{sheet.max_row - 1}"

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f"servizi_attivi_{safe_export_name(customer.business_name)}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )


@app.get("/api/v1/customers/{customer_id}/services-pivot.pdf")
def customer_services_pdf(customer_id: uuid.UUID):
    with SessionLocal() as db:
        customer = db.get(Customer, customer_id)
        if not customer:
            raise HTTPException(404, "Cliente non trovato")
        latest_import, rows = latest_customer_snapshot(db, customer_id)
        pivot = customer_service_pivot(rows)
        total_count = sum(int(item["count"]) for item in pivot)
        total_mrr = round(sum(float(item["mrr"]) for item in pivot), 2)

        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=18 * mm,
            leftMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
            title=f"Servizi attivi - {customer.business_name}",
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CornetTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            alignment=TA_LEFT,
            textColor=colors.HexColor("#1B2035"),
            spaceAfter=6,
        )
        story = [
            Paragraph("Servizi attivi", title_style),
            Paragraph(customer.business_name, styles["Heading2"]),
            Paragraph(
                f"Fotografia portafoglio: {latest_import.competence_month if latest_import else 'N/D'}",
                styles["BodyText"],
            ),
            Spacer(1, 8 * mm),
        ]
        data = [["Piano Tariffario", "Utenze Attive", "Ricavo Mensile (MRR)"]]
        data.extend(
            [[item["plan"], str(item["count"]), italian_currency(float(item["mrr"]))] for item in pivot]
        )
        data.append(["TOTALE GENERALE", str(total_count), italian_currency(total_mrr)])
        table = Table(data, colWidths=[91 * mm, 32 * mm, 51 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#287A51")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F3EC")),
                    ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CED8")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(table)
        document.build(story)
        output.seek(0)
        filename = f"servizi_attivi_{safe_export_name(customer.business_name)}.pdf"
        return StreamingResponse(
            output,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )


@app.get("/api/v1/windtre-imports")
def import_history():
    with SessionLocal() as db:
        return [serialize_import(item) for item in db.scalars(select(WindTreImport).order_by(WindTreImport.competence_month.desc())).all()]


@app.get("/api/v1/windtre-imports/{import_id}")
def import_detail(import_id: uuid.UUID):
    with SessionLocal() as db:
        item = db.get(WindTreImport, import_id)
        if not item:
            raise HTTPException(404, "Importazione non trovata")
        changes = db.scalars(
            select(WindTreChange).where(WindTreChange.import_id == import_id).order_by(WindTreChange.change_type, WindTreChange.customer_key)
        ).all()
        result = serialize_import(item)
        result["changes"] = [
            {
                "id": str(change.id),
                "change_type": change.change_type,
                "customer_id": str(change.customer_id) if change.customer_id else None,
                "customer_key": change.customer_key,
                "asset_key": change.asset_key,
                "field_name": change.field_name,
                "old_value": change.old_value,
                "new_value": change.new_value,
            }
            for change in changes
        ]
        return result


@app.post("/api/v1/windtre-imports/preview")
async def preview_windtre_file(file: UploadFile = File(...)):
    contents = await file.read()
    validate_excel_upload(file.filename, contents)
    try:
        parsed = parse_workbook(contents)
    except Exception as exc:
        raise HTTPException(422, f"Impossibile leggere il file Excel: {exc}") from exc
    if not parsed:
        raise HTTPException(422, "Il file non contiene righe cliente riconoscibili")
    return build_preview(file.filename or "estrazione.xlsx", contents, parsed)


@app.post("/api/v1/windtre-imports")
async def import_windtre_file(competence_month: str = Form(...), file: UploadFile = File(...)):
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", competence_month):
        raise HTTPException(422, "Il mese deve essere nel formato AAAA-MM")
    contents = await file.read()
    validate_excel_upload(file.filename, contents)
    digest = sha256(contents).hexdigest()
    try:
        parsed = parse_workbook(contents)
    except Exception as exc:
        raise HTTPException(422, f"Impossibile leggere il file Excel: {exc}") from exc
    if not parsed:
        raise HTTPException(422, "Il file non contiene righe cliente riconoscibili")

    with SessionLocal() as db:
        duplicate = db.scalar(
            select(WindTreImport).where(
                WindTreImport.competence_month == competence_month,
                WindTreImport.file_sha256 == digest,
            )
        )
        if duplicate:
            raise HTTPException(409, "Questo file è già stato importato per il mese selezionato")

        previous = db.scalar(
            select(WindTreImport)
            .where(WindTreImport.competence_month < competence_month)
            .order_by(WindTreImport.competence_month.desc())
            .limit(1)
        )
        previous_rows = db.scalars(select(WindTreImportRow).where(WindTreImportRow.import_id == previous.id)).all() if previous else []
        previous_by_asset = {row.asset_key: row for row in previous_rows}
        previous_customer_keys = {row.customer_key for row in previous_rows}

        record = WindTreImport(
            competence_month=competence_month,
            file_name=file.filename,
            file_sha256=digest,
            row_count=len(parsed),
        )
        db.add(record)
        db.flush()

        current_by_asset: dict[str, dict[str, Any]] = {}
        customer_ids: dict[str, uuid.UUID] = {}
        new_customer_keys: set[str] = set()
        for row in parsed:
            customer = None
            if row["customer_code"]:
                account_code = db.scalar(
                    select(CustomerAccountCode).where(
                        CustomerAccountCode.operator == "WINDTRE",
                        CustomerAccountCode.market == "BUSINESS_SME",
                        CustomerAccountCode.customer_code == row["customer_code"].upper(),
                    )
                )
                customer = db.get(Customer, account_code.customer_id) if account_code else None
                if not customer:
                    customer = db.scalar(select(Customer).where(Customer.windtre_customer_code == row["customer_code"]))
            if not customer and row["tax_id"]:
                customer = db.scalar(select(Customer).where(Customer.tax_id == row["tax_id"]))
            if not customer and row["fiscal_code"]:
                customer = db.scalar(select(Customer).where(Customer.fiscal_code == row["fiscal_code"]))
            if not customer:
                customer = Customer(
                    segment="BUSINESS_SME",
                    business_name=row["business_name"],
                    tax_id=row["tax_id"] or None,
                    fiscal_code=row["fiscal_code"] or None,
                    windtre_customer_code=row["customer_code"] or None,
                    email=row["email"] or None,
                    phone=row["phone"] or None,
                    first_seen_month=competence_month,
                    last_seen_month=competence_month,
                )
                db.add(customer)
                db.flush()
                new_customer_keys.add(row["customer_key"])
            else:
                customer.business_name = row["business_name"] or customer.business_name
                customer.windtre_customer_code = row["customer_code"] or customer.windtre_customer_code
                customer.email = row["email"] or customer.email
                customer.phone = row["phone"] or customer.phone
                customer.last_seen_month = competence_month
                customer.portfolio_status = "ACTIVE"
            if row["customer_code"]:
                normalized_code = row["customer_code"].upper()
                account_code = db.scalar(
                    select(CustomerAccountCode).where(
                        CustomerAccountCode.operator == "WINDTRE",
                        CustomerAccountCode.market == "BUSINESS_SME",
                        CustomerAccountCode.customer_code == normalized_code,
                    )
                )
                if not account_code:
                    db.add(
                        CustomerAccountCode(
                            customer_id=customer.id,
                            operator="WINDTRE",
                            market="BUSINESS_SME",
                            customer_code=normalized_code,
                            is_primary=True,
                        )
                    )
                    db.flush()
            customer_ids[row["customer_key"]] = customer.id
            current_by_asset[row["asset_key"]] = row
            db.add(
                WindTreImportRow(
                    import_id=record.id,
                    customer_id=customer.id,
                    row_number=row["row_number"],
                    customer_key=row["customer_key"],
                    asset_key=row["asset_key"],
                    business_name=row["business_name"],
                    asset_type=row["asset_type"] or None,
                    asset_number=row["asset_number"] or None,
                    current_plan=row["plan"] or None,
                    current_status=row["status"] or None,
                    monthly_fee=row["monthly_fee"] or None,
                    raw_data=row["raw"],
                    campaigns=row["campaigns"],
                )
            )

        def add_change(kind: str, customer_key: str, asset_key: str | None = None, field: str | None = None, old=None, new=None):
            db.add(
                WindTreChange(
                    import_id=record.id,
                    customer_id=customer_ids.get(customer_key),
                    change_type=kind,
                    customer_key=customer_key,
                    asset_key=asset_key,
                    field_name=field,
                    old_value=text(old),
                    new_value=text(new),
                )
            )

        for key in new_customer_keys:
            add_change("NEW_CUSTOMER", key)
        current_customer_keys = {row["customer_key"] for row in parsed}
        missing_customer_keys = previous_customer_keys - current_customer_keys
        for key in missing_customer_keys:
            previous_row = next((old_row for old_row in previous_rows if old_row.customer_key == key), None)
            missing_customer_id = previous_row.customer_id if previous_row else None
            db.add(
                WindTreChange(
                    import_id=record.id,
                    customer_id=missing_customer_id,
                    change_type="MISSING_CUSTOMER",
                    customer_key=key,
                )
            )
            customer = db.get(Customer, missing_customer_id) if missing_customer_id else None
            if customer:
                customer.portfolio_status = "MISSING_FROM_LATEST_IMPORT"

        new_assets = set(current_by_asset) - set(previous_by_asset)
        removed_assets = set(previous_by_asset) - set(current_by_asset)
        for key in new_assets:
            add_change("NEW_ASSET", current_by_asset[key]["customer_key"], key)
        for key in removed_assets:
            previous_row = previous_by_asset[key]
            add_change("REMOVED_ASSET", previous_row.customer_key, key)

        field_change_count = 0
        campaign_change_count = 0
        for key in set(current_by_asset) & set(previous_by_asset):
            current = current_by_asset[key]
            old = previous_by_asset[key]
            old_raw = old.raw_data or {}
            for field in SIGNIFICANT_FIELDS:
                before, after = clean(old_raw.get(field)), clean(current["raw"].get(field))
                if before != after:
                    add_change("FIELD_CHANGED", current["customer_key"], key, field, before, after)
                    field_change_count += 1
            campaign_names = set(old.campaigns or {}) | set(current["campaigns"])
            for campaign in campaign_names:
                before = clean((old.campaigns or {}).get(campaign))
                after = clean(current["campaigns"].get(campaign))
                if before == after:
                    continue
                kind = "CAMPAIGN_ENTERED" if not before else "CAMPAIGN_EXITED" if not after else "CAMPAIGN_CHANGED"
                add_change(kind, current["customer_key"], key, campaign, before, after)
                campaign_change_count += 1

        record.customer_count = len(current_customer_keys)
        record.new_customers = len(new_customer_keys)
        record.missing_customers = len(missing_customer_keys)
        record.new_assets = len(new_assets)
        record.removed_assets = len(removed_assets)
        record.field_changes = field_change_count
        record.campaign_changes = campaign_change_count
        try:
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(409, "Importazione duplicata o dati fiscali incompatibili") from exc
        db.refresh(record)
        return serialize_import(record)
